"""
Motored -- ERP "Entradas x Compra" template for the pending invoices the
administrative analyst enters (odd/tasks/motored-ingresos-responsable-
plantilla.md, T2).

`construir_libro` is pure: it reproduces sheet "Entrada Compra" of the
business template (cell by cell, widths and number formats included).
`preparar` checks that the invoice may be downloaded and gathers its data.

Lines: one per referencia of the invoice (for the principal store, associated
stores rolled up) whose NET quantity is positive. A credit note lowers the
quantity of the line it credits, and a fully credited line is left out, the
same net rule the pending list uses. Unit price: the invoice's "Vlr.
Unitario"; when the line was loaded before that column existed, total over
quantity rounded to cents.
"""
import io
import uuid
from dataclasses import dataclass
from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from typing import Any, Dict, List, Optional, Tuple, Union

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.motored.models.carga_archivo import CargaArchivo
from app.motored.models.factura_proveedor_linea import FacturaProveedorLinea
from app.motored.models.referencia import Referencia
from app.motored.models.sucursal import Sucursal
from app.motored.services import ingresos_pendientes as ingresos
from app.motored.services import parametros, parametros_claves as claves
from app.motored.services import sucursal_grupo

HOJA = "Entrada Compra"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FILA_ENCABEZADOS = 7
PRIMERA_LINEA = 8
ENCABEZADOS = (
    "Consecutivo Entrada", "Referencia", "Bodega", "Precio Unit", "Cantidad",
    "Vr Bruto", "Dcto Item", "Notas del Mvto", "Centro Operación", "U.N.",
)
ANCHOS = {"A": 19.7109375, "B": 41.28515625, "C": 16.85546875,
          "D": 17.85546875, "E": 12.5703125, "F": 15.140625,
          "G": 13.7109375, "H": 11.5703125, "I": 14.0}
FORMATO_MONEDA = '"$"#,##0.00;"$"\\-#,##0.00'
FORMATO_PORCENTAJE = '#,##0.000_ ;\\-#,##0.000\\ '
FORMATO_FECHA = "mm-dd-yy"
FORMATO_VALOR = "#,##0.00"

MSG_NO_PENDIENTE = (
    "La factura no está pendiente de ingreso (no existe o ya fue ingresada).")
MSG_ES_DEL_ASESOR = (
    "Esta factura la ingresa el asesor de la tienda: la plantilla es solo "
    "para las facturas del analista administrativo.")
MSG_SIN_LLEGAR = (
    "Primero hay que confirmar que la factura llegó a la tienda para "
    "descargar la plantilla.")
MSG_SIN_BODEGA = (
    "La tienda no tiene bodega principal. Cárgala en Configuración → "
    "Tiendas y vuelve a descargar la plantilla.")
MSG_SIN_LINEAS = "La factura no tiene referencias con cantidad por ingresar."

_AJUSTES_POR_CLAVE = (
    ("tipo_documento", claves.CLAVE_PLANTILLA_TIPO_DOC),
    ("descuento_global", claves.CLAVE_PLANTILLA_DESC_GLOBAL),
    ("proveedor", claves.CLAVE_PLANTILLA_PROVEEDOR),
    ("sucursal_proveedor", claves.CLAVE_PLANTILLA_SUC_PROVEEDOR),
    ("comprador", claves.CLAVE_PLANTILLA_COMPRADOR),
    ("descuento_item", claves.CLAVE_PLANTILLA_DESC_ITEM),
    ("unidad_negocio", claves.CLAVE_PLANTILLA_UNIDAD_NEGOCIO),
)


class PlantillaError(Exception):
    def __init__(self, status_code: int, detail: str):
        super().__init__(detail)
        self.status_code = status_code
        self.detail = detail


@dataclass(frozen=True)
class LineaPlantilla:
    referencia: str
    cantidad: Decimal
    valor_unitario: Optional[Decimal]
    valor_total: Optional[Decimal] = None


@dataclass(frozen=True)
class DatosPlantilla:
    prefijo_rh: str
    numero_rh: int
    codigo_co: str
    bodega: str
    fecha: date
    ajustes: Dict[str, Any]
    lineas: List[LineaPlantilla]


def _numero(valor: Any) -> Union[int, float]:
    """Whole numbers as `int` (the template shows 3962, not 3962.0)."""
    decimal = Decimal(str(valor))
    if decimal == decimal.to_integral_value():
        return int(decimal)
    return float(decimal)


def _precio(linea: LineaPlantilla) -> Optional[Union[int, float]]:
    if linea.valor_unitario is not None:
        return _numero(linea.valor_unitario)
    if linea.valor_total is not None and linea.cantidad:
        unitario = (linea.valor_total / linea.cantidad).quantize(
            Decimal("0.01"), rounding=ROUND_HALF_UP)
        return _numero(unitario)
    return None


def construir_libro(datos: DatosPlantilla) -> bytes:
    """The workbook with the single sheet "Entrada Compra"."""
    libro = Workbook()
    ws = libro.active
    ws.title = HOJA
    negrita = Font(bold=True)
    documento = f"{datos.prefijo_rh}{datos.numero_rh}"
    ajustes = datos.ajustes

    ws["A1"], ws["B1"], ws["C1"] = "Centro de Operación", datos.codigo_co, "B"
    ws["D1"], ws["E1"] = datos.fecha, "AAAAMMDD"
    ws["H1"], ws["I1"] = "MODO", 2
    ws["D1"].number_format = FORMATO_FECHA
    ws["A2"], ws["B2"] = "Tipo Documento:", _numero(ajustes["tipo_documento"])
    ws["C2"], ws["D2"] = "Descuento Global", _numero(ajustes["descuento_global"])
    ws["A3"], ws["B3"] = "Proveedor", _numero(ajustes["proveedor"])
    ws["C3"], ws["D3"] = "Suc. Proveedor", str(ajustes["sucursal_proveedor"])
    ws["A4"], ws["B4"] = "Comprador", _numero(ajustes["comprador"])
    ws["C4"], ws["D4"] = "Descuento x Item", _numero(ajustes["descuento_item"])
    ws["A5"], ws["B5"] = "Notas Documento", f"ENTRADA POR COMPRA {documento}"
    ws["C5"], ws["D5"] = "Doc. Referencia", documento
    ws["D2"].number_format = ws["D4"].number_format = FORMATO_PORCENTAJE
    for celda in ("A1", "C1", "H1", "I1", "A2", "C2", "A3", "C3", "A4", "C4",
                  "A5", "C5"):
        ws[celda].font = negrita

    for columna, titulo in enumerate(ENCABEZADOS, start=1):
        celda = ws.cell(row=FILA_ENCABEZADOS, column=columna, value=titulo)
        celda.font = negrita

    for orden, linea in enumerate(datos.lineas, start=1):
        fila = PRIMERA_LINEA + orden - 1
        ws.cell(fila, 1, 1)
        ws.cell(fila, 2, linea.referencia)
        ws.cell(fila, 3, datos.bodega).number_format = FORMATO_MONEDA
        ws.cell(fila, 4, _precio(linea)).number_format = FORMATO_MONEDA
        ws.cell(fila, 5, _numero(linea.cantidad))
        ws.cell(fila, 6, f"=D{fila}*E{fila}").number_format = FORMATO_VALOR
        ws.cell(fila, 7, f"=+F{fila}*$D$4%").number_format = FORMATO_VALOR
        ws.cell(fila, 8, f"Mvto {orden}")
        ws.cell(fila, 9, "=+B1" if orden == 1 else datos.codigo_co)
        ws.cell(fila, 10, str(ajustes["unidad_negocio"]))

    for columna, ancho in ANCHOS.items():
        ws.column_dimensions[columna].width = ancho
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


# ---------------------------------------------------------------------------
# Database
# ---------------------------------------------------------------------------

async def _ajustes(db: AsyncSession, hoy: date) -> Dict[str, Any]:
    valores = await parametros.leer_valores(
        db, hoy, claves.RESPALDOS_INGRESOS)
    return {nombre: valores[clave] for nombre, clave in _AJUSTES_POR_CLAVE}


async def _sucursal(db: AsyncSession, sucursal_id: uuid.UUID) -> Optional[Any]:
    return (await db.execute(
        select(Sucursal.codigo_co, Sucursal.bodega_principal)
        .where(Sucursal.id == sucursal_id))).first()


async def _lineas(
    db: AsyncSession, factura: Tuple[str, int], tienda: uuid.UUID,
) -> List[LineaPlantilla]:
    """Net lines of the invoice for the principal store (associated stores
    rolled up), positive quantities only, ordered by referencia code."""
    principal = await sucursal_grupo.principal_de(db)
    grupo = [sid for sid, p in principal.items() if p == tienda] or [tienda]
    neta = func.sum(FacturaProveedorLinea.cantidad)
    stmt = (
        select(
            Referencia.codigo, neta.label("cantidad"),
            func.max(FacturaProveedorLinea.valor_unitario).label("unitario"),
            func.sum(FacturaProveedorLinea.valor_total).label("total"))
        .join(Referencia, Referencia.id == FacturaProveedorLinea.referencia_id)
        .join(CargaArchivo, CargaArchivo.id == FacturaProveedorLinea.carga_id)
        .where(
            CargaArchivo.estado != "ANULADO",
            FacturaProveedorLinea.prefijo_rh == factura[0],
            FacturaProveedorLinea.numero_rh == factura[1],
            FacturaProveedorLinea.sucursal_id.in_(grupo))
        .group_by(Referencia.codigo)
        .having(neta > 0)
        .order_by(Referencia.codigo))
    return [
        LineaPlantilla(
            codigo, Decimal(cantidad),
            None if unitario is None else Decimal(unitario), Decimal(total))
        for codigo, cantidad, unitario, total in (await db.execute(stmt)).all()]


async def preparar(
    db: AsyncSession, factura: Tuple[str, int], sucursal_id: uuid.UUID,
    hoy: date,
) -> DatosPlantilla:
    """Checks the invoice can be downloaded and gathers the workbook data.
    404 not pending, 409 asesor invoice / not confirmed LLEGO / store without
    bodega principal / nothing to enter."""
    principal = await sucursal_grupo.principal_de(db)
    tienda = principal.get(sucursal_id, sucursal_id)
    item = next(
        (i for i in await ingresos.pendientes(db, [tienda], hoy)
         if (i["prefijo_rh"], i["numero_rh"]) == factura), None)
    if item is None:
        raise PlantillaError(404, MSG_NO_PENDIENTE)
    if item["responsable"] != ingresos.RESPONSABLE_ANALISTA:
        raise PlantillaError(409, MSG_ES_DEL_ASESOR)
    if item["estado"] != ingresos.LLEGO:
        raise PlantillaError(409, MSG_SIN_LLEGAR)

    sucursal = await _sucursal(db, tienda)
    bodega = (getattr(sucursal, "bodega_principal", None) or "").strip()
    if not bodega:
        raise PlantillaError(409, MSG_SIN_BODEGA)
    lineas = await _lineas(db, factura, tienda)
    if not lineas:
        raise PlantillaError(409, MSG_SIN_LINEAS)
    return DatosPlantilla(
        prefijo_rh=factura[0], numero_rh=factura[1],
        codigo_co=sucursal.codigo_co, bodega=bodega, fecha=hoy,
        ajustes=await _ajustes(db, hoy), lineas=lineas)
