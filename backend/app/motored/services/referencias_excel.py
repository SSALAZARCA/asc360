"""
Referencias -- the master as an Excel download
(odd/tasks/motored-referencias-descarga-excel.md, T1).

The columns are the bulk-upload template's (`carga_excel.column_labels`,
same text and order) plus a trailing "Estado" (Activa/Inactiva). The
upload ignores headers it does not know, so the owner can edit the file
and re-upload it unchanged: codes are text cells (a leading zero
survives), prices and the pack size are real numbers, and blank values
stay blank (a blank cell is "not provided" on upload, never a delete).
"""
import io
from datetime import date
from decimal import Decimal
from typing import Any, Iterable, List, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.motored.services.carga_excel import column_labels

XLSX = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
TEXTO = "@"
FORMATO_PRECIO = "#,##0.00"
COLUMNAS = (*column_labels("referencia"), "Estado")
ESTADO = {True: "Activa", False: "Inactiva"}
# Positions in a `filas_exportacion` row (same order as COLUMNAS).
_TEXTO = (0, 1, 7)
_PRECIOS = (5, 6)
_HOMOLOGADOS = 8
_ACTIVA = 9
ANCHOS = {
    "Código": 18, "Código del proveedor": 14, "Nombre": 40,
    "Línea comercial": 18, "Código de referencia sustituta": 18,
    "Homologados otras marcas": 30,
}


def nombre_archivo(dia: date) -> str:
    """`referencias_<AAAA-MM-DD>.xlsx`; the caller passes the Bogotá day
    (`reloj.hoy_bogota`)."""
    return f"referencias_{dia.isoformat()}.xlsx"


def _celdas(fila: Sequence[Any]) -> List[Any]:
    valores = list(fila)
    homologados = valores[_HOMOLOGADOS] or []
    valores[_HOMOLOGADOS] = ", ".join(homologados) or None
    valores[_ACTIVA] = ESTADO[bool(valores[_ACTIVA])]
    for indice in _PRECIOS:
        if isinstance(valores[indice], Decimal):
            valores[indice] = _numero(valores[indice])
    return valores


def _numero(valor: Decimal) -> Any:
    """An int when whole, else a float (2-decimal prices are exact)."""
    if valor == valor.to_integral_value():
        return int(valor)
    return float(valor)


def _formatear(hoja, total_filas: int) -> None:
    hoja.freeze_panes = "A2"
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for indice, titulo in enumerate(COLUMNAS, start=1):
        letra = get_column_letter(indice)
        hoja.column_dimensions[letra].width = ANCHOS.get(titulo, 16)
    for fila in hoja.iter_rows(min_row=2, max_row=total_filas + 1):
        for indice in _TEXTO:
            fila[indice].number_format = TEXTO
        for indice in _PRECIOS:
            fila[indice].number_format = FORMATO_PRECIO


def libro(filas: Iterable[Sequence[Any]]) -> bytes:
    """The workbook bytes: the header row, then one row per referencia."""
    workbook = Workbook()
    hoja = workbook.active
    hoja.title = "Referencias"
    hoja.append(list(COLUMNAS))
    total = 0
    for fila in filas:
        hoja.append(_celdas(fila))
        total += 1
    _formatear(hoja, total)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
