"""
KPI's, Inventario: the Excel of the tab (`GET /kpis/inventario/excel`), built from the payload of
`calcular_kpis_inventario` plus the COMPLETE lists (every idle pair and every stockout, not only the top).

Sheets "Tiendas", "Líneas", "Sin movimiento" and "Agotadas". Each starts with header rows (the corte, the
cost-of-sales window and the store filter) and then a bold header row and the table. Amounts are numbers
(`#,##0`), shares use the percent format and days and rotation one decimal.
"""
import io
from typing import Any, Dict, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DINERO, PORCENTAJE, DECIMAL = "#,##0", "0.0%", "0.0"
NEGRITA = Font(bold=True)


def nombre_de_archivo(datos: Dict[str, Any]) -> str:
    return f"inventario_{datos['corte'] or 'sin_corte'}.xlsx"


def _encabezados(hoja: Worksheet, titulo: str, datos: Dict[str, Any], tiendas: Sequence[str]) -> None:
    hoja.append([titulo])
    hoja["A1"].font = Font(bold=True, size=14)
    hoja.append(["Corte de inventario", datos["corte"] or "Sin inventario cargado"])
    hoja.append(["Ventana del costo de venta", f"{datos['costo_desde']} a {datos['costo_hasta']}"])
    hoja.append(["Tiendas", ", ".join(tiendas) if tiendas else "Todas las tiendas"])


def _tabla(hoja: Worksheet, cabecera: List[str], filas: List[List[Any]], formatos: Dict[int, str],
           anchos: Sequence[int]) -> None:
    """Blank row, bold header and the rows; `formatos` maps a 1-based column to its number format."""
    hoja.append([])
    hoja.append(cabecera)
    for celda in hoja[hoja.max_row]:
        celda.font = NEGRITA
    for fila in filas:
        hoja.append(fila)
        for columna, formato in formatos.items():
            hoja.cell(row=hoja.max_row, column=columna).number_format = formato
    for posicion, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(posicion)].width = ancho


def _hoja_tiendas(hoja: Worksheet, datos: Dict[str, Any], tiendas: Sequence[str]) -> None:
    _encabezados(hoja, "Inventario por tienda", datos, tiendas)
    filas = [
        [x["nombre"], x["valor"], x["dias"], x["rotacion"], x["sin_movimiento_valor"], x["disponibilidad_pct"],
         x["agotadas"]] for x in datos["tiendas"]]
    _tabla(hoja, ["Tienda", "Valor", "Días de inventario", "Rotación", "Sin movimiento", "Disponibilidad", "Agotadas"],
           filas, {2: DINERO, 3: DECIMAL, 4: DECIMAL, 5: DINERO, 6: PORCENTAJE}, (30, 16, 18, 12, 16, 16, 12))


def _hoja_lineas(hoja: Worksheet, datos: Dict[str, Any], tiendas: Sequence[str]) -> None:
    _encabezados(hoja, "Inventario por línea", datos, tiendas)
    filas = [
        [x["linea"], x["valor"], x["pct"], x["dias"], x["rotacion"], x["disponibilidad_pct"]]
        for x in datos["lineas"]]
    _tabla(hoja, ["Línea", "Valor", "% del inventario", "Días de inventario", "Rotación", "Disponibilidad"],
           filas, {2: DINERO, 3: PORCENTAJE, 4: DECIMAL, 5: DECIMAL, 6: PORCENTAJE}, (24, 16, 18, 18, 12, 16))


def _hoja_sin_movimiento(
    hoja: Worksheet, datos: Dict[str, Any], tiendas: Sequence[str], pares: List[Dict[str, Any]],
) -> None:
    _encabezados(hoja, "Referencias sin movimiento", datos, tiendas)
    hoja.append(["Umbral (días sin venta)", f"más de {datos['sin_movimiento_umbral_dias']} días"])
    filas = [[x["referencia"], x["nombre"], x["tienda"], x["existencia"], x["valor"], x["dias_sin_venta"]]
             for x in pares]
    _tabla(hoja, ["Referencia", "Nombre", "Tienda", "Existencia", "Valor", "Días sin venta"],
           filas, {5: DINERO}, (22, 36, 28, 12, 16, 16))


def _hoja_agotadas(
    hoja: Worksheet, datos: Dict[str, Any], tiendas: Sequence[str], agotadas: List[Dict[str, Any]],
) -> None:
    _encabezados(hoja, "Agotadas con demanda", datos, tiendas)
    hoja.append(["Demanda", "Vendidas + ventas perdidas de los 3 meses de la ventana"])
    filas = [
        [x["referencia"], x["nombre"], x["tienda"], x["vendidas"], x["perdidas"], x["demanda"], x["transito"],
         x.get("cobertura")] for x in agotadas]
    _tabla(hoja, ["Referencia", "Nombre", "Tienda", "Vendidas", "Perdidas", "Demanda", "En tránsito",
                  "Cobertura del tránsito"], filas, {8: PORCENTAJE}, (22, 36, 28, 12, 12, 12, 14, 22))


def construir_libro(
    datos: Dict[str, Any], sin_movimiento: List[Dict[str, Any]], agotadas: List[Dict[str, Any]],
    tiendas: Optional[Sequence[str]] = (),
) -> bytes:
    """The workbook (bytes). `tiendas` are the NAMES of the store filter; empty means every store."""
    tiendas = list(tiendas or ())
    libro = Workbook()
    _hoja_tiendas(libro.active, datos, tiendas)
    libro.active.title = "Tiendas"
    _hoja_lineas(libro.create_sheet("Líneas"), datos, tiendas)
    _hoja_sin_movimiento(libro.create_sheet("Sin movimiento"), datos, tiendas, sin_movimiento)
    _hoja_agotadas(libro.create_sheet("Agotadas"), datos, tiendas, agotadas)
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
