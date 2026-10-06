"""
KPI's, Comisiones: the Excel of the settled month, built from the payload of `calcular_kpis_comisiones`.

Sheet "Comisiones": header rows (month, stores, rules in force), one row per asesor and a footer with the
totals. After the commission come one column per configured bonus line (a check or a cross with the bonus; the
cell is TEXT), "Bono total" and "Total a pagar" (commission + bonuses). Sheet "Sin presupuesto": who sold without a budget. Amounts are numbers (`#,##0`, shown with the
reader's separators), percentages use the percent format and the cédula is TEXT (keeps leading zeros).
"""
import io
from typing import Any, Dict, List, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MESES = ("enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
         "noviembre", "diciembre")
DINERO, PORCENTAJE, TEXTO = "#,##0", "0.0%", "@"
BASES = {"con_hmcl": "con HMCL", "sin_hmcl": "sin HMCL"}
NEGRITA = Font(bold=True)


def nombre_de_archivo(datos: Dict[str, Any]) -> str:
    return f"comisiones_{datos['mes_liquidado']}.xlsx"


def _mes_largo(mes: str) -> str:
    return f"{MESES[int(mes[5:7]) - 1]} {mes[:4]}"


def _numero(valor: float) -> str:
    """`1.5` -> `1,5`, `90.0` -> `90` (Spanish decimals, as in the tab)."""
    texto = f"{valor:g}"
    return texto.replace(".", ",")


def _tramos_texto(reglas: Dict[str, Any]) -> List[str]:
    return [f"{x['nombre']} desde {_numero(x['desde_pct'])}% · {_numero(x['tasa_pct'])}%" for x in reglas["comision_tramos"]]


def _pesos(valor: float) -> str:
    """`30000` -> `$30.000` (Spanish thousands separator, as the tab shows pesos)."""
    return f"${valor:,.0f}".replace(",", ".")


def _lineas_texto(reglas: Dict[str, Any]) -> List[str]:
    filas = []
    for x in reglas.get("comision_lineas", []):
        apagada = "" if x["activo"] else " · apagada"
        filas.append(f"{x['etiqueta']} desde {_numero(x['pct_meta'])}% de la venta total · {_pesos(x['bono'])}{apagada}")
    return filas


def _cedula(hoja: Worksheet, fila: int, valor: Any) -> None:
    celda = hoja.cell(row=fila, column=1, value=str(valor))
    celda.number_format = TEXTO
    celda.data_type = "s"


def _encabezados(hoja: Worksheet, datos: Dict[str, Any], tiendas: Sequence[str]) -> None:
    reglas = datos["reglas"]
    hoja.append([f"Comisiones · {_mes_largo(datos['mes_liquidado'])}"])
    hoja["A1"].font = Font(bold=True, size=14)
    hoja.append(["Mes liquidado", datos["mes_liquidado"]])
    hoja.append(["Tiendas", ", ".join(tiendas) if tiendas else "Todas las tiendas"])
    hoja.append(["Base de cumplimiento", f"Venta {BASES[reglas['cumplimiento_base']]}"])
    hoja.append(["Base de pago", f"Venta {BASES[reglas['comision_base_pago']]}"])
    hoja.append(["Cargos con comisión", ", ".join(reglas["comision_cargos_asesor"])])
    for i, texto in enumerate(_tramos_texto(reglas)):
        hoja.append(["Tramos" if i == 0 else None, texto])
    umbral = reglas.get("comision_bono_umbral_pct")
    if umbral is not None:
        hoja.append(["Bonos por línea", f"Se activan con un cumplimiento de al menos {_numero(umbral)}%"])
        for texto in _lineas_texto(reglas):
            hoja.append([None, texto])
    hoja.append([])


def _cabecera(reglas: Dict[str, Any]) -> List[str]:
    return [
        "Cédula", "Asesor", "Tienda", "Cargo", "Presupuesto",
        f"Venta {BASES[reglas['cumplimiento_base']]} (base de cumplimiento)", "Cumplimiento %", "Tramo", "% comisión",
        f"Venta {BASES[reglas['comision_base_pago']]} (base de pago)", "Comisión",
        *[f"Bono {x['etiqueta']} (≥ {_numero(x['pct_meta'])}%)" for x in reglas.get("comision_lineas", [])],
        "Bono total", "Total a pagar"]


def _formato_fila(hoja: Worksheet, fila: int, n_lineas: int = 0) -> None:
    for columna in (5, 6, 10, 11, *range(12 + n_lineas, 14 + n_lineas)):
        hoja.cell(row=fila, column=columna).number_format = DINERO
    for columna in (7, 9):
        hoja.cell(row=fila, column=columna).number_format = PORCENTAJE


def _celda_bono(bono: Dict[str, Any], gate: Dict[str, Any]) -> str:
    if bono["paga"]:
        return f"✓ {_pesos(bono['bono_pagado'])}"
    if not bono["cumple"]:
        return "✗"
    motivo = "línea apagada" if not bono["activo"] else f"sin cumplir el {_numero(gate['umbral'])}%"
    return f"✓ no se paga ({motivo})"


def _fila_asesor(hoja: Worksheet, a: Dict[str, Any], n_lineas: int) -> None:
    hoja.append([
        None, a["nombre"], a["tienda"], a.get("cargo"), a["presupuesto"], a["venta_cumplimiento"],
        a["cumplimiento_pct"], a["tramo"], a["tasa_pct"] / 100, a["venta_comision"], a["comision"],
        *[_celda_bono(b, a["gate"]) for b in a["bonos"]], a["bono_total"], a["total_a_pagar"]])
    _cedula(hoja, hoja.max_row, a["cedula"])
    _formato_fila(hoja, hoja.max_row, n_lineas)


def _totales(hoja: Worksheet, asesores: List[Dict[str, Any]], lineas: List[Dict[str, Any]]) -> None:
    suma = {k: sum(a[k] for a in asesores) for k in (
        "presupuesto", "venta_cumplimiento", "venta_comision", "comision", "bono_total", "total_a_pagar")}
    por_linea = [sum(a["bonos"][i]["bono_pagado"] for a in asesores) for i in range(len(lineas))]
    hoja.append([
        "Total", None, None, None, suma["presupuesto"], suma["venta_cumplimiento"], None, None, None,
        suma["venta_comision"], round(suma["comision"], 2), *por_linea, suma["bono_total"],
        round(suma["total_a_pagar"], 2)])
    _formato_fila(hoja, hoja.max_row, len(lineas))
    for columna in range(12, 12 + len(lineas)):
        hoja.cell(row=hoja.max_row, column=columna).number_format = DINERO
    for celda in hoja[hoja.max_row]:
        celda.font = NEGRITA


def _hoja_comisiones(hoja: Worksheet, datos: Dict[str, Any], tiendas: Sequence[str]) -> None:
    _encabezados(hoja, datos, tiendas)
    hoja.append(_cabecera(datos["reglas"]))
    for celda in hoja[hoja.max_row]:
        celda.font = NEGRITA
    lineas = datos["reglas"].get("comision_lineas", [])
    for a in datos["asesores"]:
        _fila_asesor(hoja, a, len(lineas))
    _totales(hoja, datos["asesores"], lineas)
    anchos = (16, 30, 24, 28, 16, 24, 14, 10, 12, 24, 16, *[26] * len(lineas), 16, 16)
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho


def _hoja_sin_presupuesto(hoja: Worksheet, datos: Dict[str, Any]) -> None:
    hoja.append(["Cédula", "Nombre", "Venta"])
    for celda in hoja[1]:
        celda.font = NEGRITA
    for x in datos["advertencias"]["sin_presupuesto"]:
        hoja.append([None, x["nombre"], x["venta"]])
        _cedula(hoja, hoja.max_row, x["cedula"])
        hoja.cell(row=hoja.max_row, column=3).number_format = DINERO
    hoja.column_dimensions["A"].width, hoja.column_dimensions["B"].width = 16, 30


def construir_libro(datos: Dict[str, Any], tiendas: Sequence[str] = ()) -> bytes:
    """The workbook (bytes). `tiendas` are the NAMES of the store filter; empty means every store."""
    libro = Workbook()
    _hoja_comisiones(libro.active, datos, tiendas)
    libro.active.title = "Comisiones"
    _hoja_sin_presupuesto(libro.create_sheet("Sin presupuesto"), datos)
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
