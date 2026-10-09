"""
Motored satisfaction survey -- Excel of survey results (per carga or by date range).

Cédula and teléfono are TEXT cells (keep leading zeros); instants are shown in
Bogota time as `AAAA-MM-DD HH:MM`.
"""
import io
import re
from datetime import date, datetime, timezone
from typing import Any, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app.motored.services.encuesta_resultados import clasificar, fila_de
from app.motored.services.reloj import BOGOTA_OFFSET

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
COLUMNAS = ["Cliente", "Cédula", "Teléfono", "Tienda", "Estado", "Nota", "Categoría",
            "Comentario", "Fecha de respuesta", "Archivo de carga", "Fecha de envío"]
TEXTO = "@"
ESTADOS = {"RESPONDIDA": "Respondida", "SIN_RESPONDER": "Sin responder"}
CATEGORIAS = {"DETRACTOR": "Detractor", "SATISFECHO": "Satisfecho"}
POR_TEXTO = {"envio": "fecha de envío de la encuesta", "respuesta": "fecha de respuesta del cliente"}


def _bogota(valor: Optional[datetime]) -> Optional[datetime]:
    if valor is None:
        return None
    if valor.tzinfo is None:
        valor = valor.replace(tzinfo=timezone.utc)
    return valor.astimezone(BOGOTA_OFFSET)


def _fecha_hora(valor: Optional[datetime]) -> Optional[str]:
    local = _bogota(valor)
    return local.strftime("%Y-%m-%d %H:%M") if local else None


def _slug(nombre_archivo: str) -> str:
    base = re.sub(r"\.xlsx?$", "", nombre_archivo, flags=re.IGNORECASE)
    return re.sub(r"[^A-Za-z0-9]+", "_", base).strip("_") or "carga"


def nombre_por_carga(rows: List[Any]) -> str:
    dia = _bogota(rows[0].carga_created_at).date().isoformat()
    return f"encuesta_{_slug(rows[0].nombre_archivo)}_{dia}.xlsx"


def nombre_por_rango(desde: date, hasta: date, por: str) -> str:
    return f"encuestas_{por}_{desde.isoformat()}_{hasta.isoformat()}.xlsx"


def _texto(hoja: Worksheet, fila: int, columna: int, valor: Any) -> None:
    if valor is None:
        return
    celda = hoja.cell(row=fila, column=columna, value=str(valor))
    celda.number_format = TEXTO
    celda.data_type = "s"


def _tabla(hoja: Worksheet, rows: List[Any]) -> None:
    hoja.append(COLUMNAS)
    for celda in hoja[hoja.max_row]:
        celda.font = Font(bold=True)
    for row in rows:
        d = fila_de(row)
        hoja.append([
            d["cliente"], None, None, d["tienda"], ESTADOS[d["estado"]], d["nota"],
            CATEGORIAS.get(clasificar(d["nota"]) or ""), d["comentario"],
            _fecha_hora(row.respuesta_created_at), row.nombre_archivo, _fecha_hora(row.carga_created_at),
        ])
        _texto(hoja, hoja.max_row, 2, d["cedula"])
        _texto(hoja, hoja.max_row, 3, d["telefono"])
    for letra, ancho in zip("ABCDEFGHIJK", (28, 14, 14, 22, 14, 8, 12, 40, 18, 28, 18)):
        hoja.column_dimensions[letra].width = ancho


def _bytes(libro: Workbook) -> bytes:
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def construir_por_carga(rows: List[Any]) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Encuesta"
    _tabla(hoja, rows)
    return _bytes(libro)


def construir_por_rango(rows: List[Any], desde: date, hasta: date, por: str, generado: datetime) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Encuestas"
    hoja.append([f"Resultados de encuestas · {desde.isoformat()} a {hasta.isoformat()}"])
    hoja["A1"].font = Font(bold=True, size=14)
    hoja.append(["Rango", f"{desde.isoformat()} a {hasta.isoformat()}"])
    hoja.append(["Filtrado por", POR_TEXTO[por]])
    hoja.append(["Generado", _fecha_hora(generado)])
    hoja.append([])
    _tabla(hoja, rows)
    return _bytes(libro)
