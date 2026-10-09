"""
Inventory counts -- the ERP adjustment Excel (odd/motored-conteos-
inventario, WU10; design §1.10, §4.8; owner decision 2026-10-08: no ERP
template yet).

`libro` builds the workbook in memory and returns its bytes:

- "Ajustes": one row per known referencia whose difference is not 0,
  largest |value| first, all assigned to the store's principal bodega.
  Unknown codes (not in the master) cannot be adjusted in the ERP, so
  they only appear in "Sin costo";
- "Resumen": store, dates, leader, accuracy KPI, money totals and the
  forced-close reason;
- "Sin costo" (only when there is one): differences with no cost, which
  stay out of the money totals, including unknown codes.

Codes are text cells; money is whole pesos (`#,##0`); quantities are
numbers; instants are Bogotá wall-clock. The same layout serves the live
"avance" download of an open conteo.
"""
import io
from datetime import date, datetime, timezone
from decimal import ROUND_HALF_UP, Decimal
from typing import Any, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.motored.services.conteos.cierre import Encabezado, Linea
from app.motored.services.reloj import BOGOTA_OFFSET

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PESOS = "#,##0"
TEXTO = "@"
FECHA_HORA = "yyyy-mm-dd hh:mm"
FECHA = "yyyy-mm-dd"
SI_NO = {True: "Sí", False: "No"}
COLUMNAS_AJUSTES = (
    "Referencia", "Descripción", "Bodega", "Cantidad sistema",
    "Cantidad contada", "Diferencia", "Costo unitario", "Valor diferencia",
    "Ubicaciones", "Con reconteo", "Crítica")
COLUMNAS_SIN_COSTO = (
    "Referencia", "Descripción", "Cantidad sistema", "Cantidad contada",
    "Diferencia", "Ubicaciones", "Observación")
ANCHOS = {"Descripción": 40, "Ubicaciones": 30, "Observación": 45}
DINERO = ("Costo unitario", "Valor diferencia")
NO_EN_MAESTRO = (
    "El código no existe en el maestro de referencias: revíselo antes "
    "de ajustar.")
SIN_COSTO = "La referencia no tiene costo: valórela en el ERP."
ESTADOS = {
    "EN_CONTEO": "En conteo", "EN_RECONTEO": "En reconteo",
    "CERRADO": "Cerrado"}


def nombre_archivo(prefijo: str, codigo_co: str, fecha: date) -> str:
    """`<prefijo>_conteo_<C.O.>_<AAAA-MM-DD>.xlsx`."""
    return f"{prefijo}_conteo_{codigo_co}_{fecha.isoformat()}.xlsx"


def _pesos(valor: Optional[Decimal]) -> Optional[int]:
    if valor is None:
        return None
    return int(Decimal(valor).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _numero(valor: Optional[Decimal]) -> Any:
    """A quantity as an int when whole, else a float."""
    if valor is None:
        return None
    if valor == valor.to_integral_value():
        return int(valor)
    return float(valor)


def _bogota(valor: Optional[datetime]) -> Optional[datetime]:
    """Bogotá wall-clock, naive (openpyxl rejects aware values)."""
    if valor is None:
        return None
    if valor.tzinfo is None:
        valor = valor.replace(tzinfo=timezone.utc)
    return valor.astimezone(BOGOTA_OFFSET).replace(tzinfo=None)


def _cabecera(hoja: Worksheet, columnas: Sequence[str]) -> None:
    hoja.append(list(columnas))
    for indice, nombre in enumerate(columnas, start=1):
        hoja.cell(row=1, column=indice).font = Font(bold=True)
        letra = get_column_letter(indice)
        hoja.column_dimensions[letra].width = ANCHOS.get(nombre, 16)
    hoja.freeze_panes = "A2"


def _formatear(hoja: Worksheet, columnas: Sequence[str]) -> None:
    """Codes as text, money as whole pesos, for every data row."""
    for indice, nombre in enumerate(columnas, start=1):
        formato = TEXTO if nombre == "Referencia" else (
            PESOS if nombre in DINERO else None)
        if formato is None:
            continue
        for (celda,) in hoja.iter_rows(
                min_row=2, min_col=indice, max_col=indice):
            celda.number_format = formato
            if formato == TEXTO and celda.value is not None:
                celda.data_type = "s"


def _hoja_ajustes(hoja: Worksheet, lineas: Sequence[Linea]) -> None:
    _cabecera(hoja, COLUMNAS_AJUSTES)
    for f in lineas:
        if f.diferencia == 0 or f.referencia_id is None:
            continue
        hoja.append([
            f.codigo, f.descripcion, f.bodega, _numero(f.sistema),
            _numero(f.contado), _numero(f.diferencia),
            _pesos(f.costo_unitario), _pesos(f.valor),
            ", ".join(f.ubicaciones), SI_NO[f.con_reconteo],
            SI_NO[f.critico]])
    _formatear(hoja, COLUMNAS_AJUSTES)


def _hoja_sin_costo(hoja: Worksheet, lineas: Sequence[Linea]) -> None:
    _cabecera(hoja, COLUMNAS_SIN_COSTO)
    for f in lineas:
        hoja.append([
            f.codigo, f.descripcion, _numero(f.sistema), _numero(f.contado),
            _numero(f.diferencia), ", ".join(f.ubicaciones),
            NO_EN_MAESTRO if f.referencia_id is None else SIN_COSTO])
    _formatear(hoja, COLUMNAS_SIN_COSTO)


def _filas_resumen(e: Encabezado) -> List[tuple]:
    kpi = e.kpi
    pct = None if kpi.exactitud_pct is None else float(kpi.exactitud_pct)
    return [
        ("Tienda", e.tienda, None), ("C.O.", e.codigo_co, TEXTO),
        ("Estado", ESTADOS.get(e.estado, e.estado), None),
        ("Líder", e.lider, None),
        ("Fecha programada", e.fecha_programada, FECHA),
        ("Inventario del sistema (corte)", e.fecha_corte, FECHA),
        ("Iniciado", _bogota(e.iniciado_en), FECHA_HORA),
        ("Ronda de conteo terminada", _bogota(e.ronda_terminada_en),
         FECHA_HORA),
        ("Cerrado", _bogota(e.cerrado_en), FECHA_HORA),
        ("Bodega de ajuste", e.bodega, TEXTO),
        ("Referencias evaluadas", kpi.refs_universo, None),
        ("Referencias exactas", kpi.refs_exactas, None),
        ("Exactitud (%)", pct, "0.00"),
        ("Valor del inventario del sistema", _pesos(kpi.valor_sistema),
         PESOS),
        ("Diferencia neta", _pesos(kpi.valor_diferencia_neta), PESOS),
        ("Diferencia absoluta", _pesos(kpi.valor_diferencia_abs), PESOS),
        ("Motivo del cierre forzado", e.motivo_cierre_forzado, None),
    ]


def _hoja_resumen(hoja: Worksheet, encabezado: Encabezado) -> None:
    hoja.column_dimensions["A"].width = 34
    hoja.column_dimensions["B"].width = 40
    for fila, (etiqueta, valor, formato) in enumerate(
            _filas_resumen(encabezado), start=1):
        hoja.cell(row=fila, column=1, value=etiqueta).font = Font(
            bold=True)
        celda = hoja.cell(row=fila, column=2, value=valor)
        if formato is not None:
            celda.number_format = formato


def libro(encabezado: Encabezado, lineas: Sequence[Linea]) -> bytes:
    """The workbook's bytes; `lineas` already sorted."""
    salida = Workbook()
    _hoja_ajustes(salida.active, lineas)
    salida.active.title = "Ajustes"
    _hoja_resumen(salida.create_sheet("Resumen"), encabezado)
    sin_costo = [
        f for f in lineas if f.valor is None and f.diferencia != 0]
    if sin_costo:
        _hoja_sin_costo(salida.create_sheet("Sin costo"), sin_costo)
    archivo = io.BytesIO()
    salida.save(archivo)
    return archivo.getvalue()
