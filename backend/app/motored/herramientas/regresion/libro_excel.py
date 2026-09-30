"""
Motored Pedidos F3 "Motor" (S8b, ADR-10) — primitivas para escribir libros.

`LibroSalida` envuelve un `openpyxl.Workbook` en modo `write_only`: las filas
se van escribiendo una a una y un libro grande (47 sucursales, decenas de
miles de filas) nunca vive entero en memoria. Los números exactos del motor
(`Fraction`) salen cuantizados a seis decimales con ROUND_HALF_UP, igual que
al persistir; nada vuelve del Excel a un cálculo.

La ruta de destino se valida al crear el libro, con `ruta_salida_segura`,
ANTES de generar ninguna fila: un informe con datos reales jamás cae dentro de
un árbol git y una ruta rechazada no deja nada a medio escribir.
"""
import re
from fractions import Fraction
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.motored.herramientas.regresion.rutas import ruta_salida_segura
from app.motored.services.motor.aritmetica import cuantizar

MAX_TITULO = 31
ESCALA_INFORME = 6
ANCHO_MINIMO = 10
ANCHO_MAXIMO = 45
TITULO_POR_DEFECTO = "Hoja"
PROHIBIDOS = re.compile(r"[\[\]:*?/\\]")
NEGRITA = Font(bold=True)


def titulo_hoja(texto: str) -> str:
    """Título válido de hoja: sin `[]:*?/\\`, hasta 31 caracteres."""
    limpio = PROHIBIDOS.sub("-", texto).strip().strip("'")
    limpio = limpio[:MAX_TITULO].strip()
    return limpio or TITULO_POR_DEFECTO


def valor_celda(valor: Any) -> Any:
    """Valor escribible: las `Fraction` pasan a `Decimal` cuantizado."""
    if isinstance(valor, Fraction):
        return cuantizar(valor, ESCALA_INFORME)
    return valor


class LibroSalida:
    """Libro de Excel en modo `write_only`, con títulos de hoja únicos.

    Levanta `RutaInsegura` si `destino` está dentro de un árbol git.
    """

    def __init__(self, destino) -> None:
        self.ruta = ruta_salida_segura(destino)
        self.libro = Workbook(write_only=True)
        self._usados: set[str] = set()

    def _titulo_unico(self, titulo: str) -> str:
        base = titulo_hoja(titulo)
        candidato, numero = base, 1
        while candidato in self._usados:
            numero += 1
            sufijo = f" ({numero})"
            candidato = base[:MAX_TITULO - len(sufijo)] + sufijo
        self._usados.add(candidato)
        return candidato

    def hoja(self, titulo: str, encabezado: Sequence[str],
             filas: Iterable[Sequence[Any]], *,
             previas: Sequence[Sequence[Any]] = ()) -> str:
        """Agrega una hoja: `previas`, el encabezado en negrita y las filas.

        Devuelve el título final (único) de la hoja.
        """
        nombre = self._titulo_unico(titulo)
        hoja = self.libro.create_sheet(nombre)
        for posicion, texto in enumerate(encabezado, start=1):
            ancho = min(max(len(texto) + 2, ANCHO_MINIMO), ANCHO_MAXIMO)
            hoja.column_dimensions[get_column_letter(posicion)].width = ancho
        for previa in previas:
            hoja.append([valor_celda(v) for v in previa])
        hoja.append([self._titulo_columna(hoja, t) for t in encabezado])
        for fila in filas:
            hoja.append([valor_celda(v) for v in fila])
        return nombre

    @staticmethod
    def _titulo_columna(hoja, texto: str) -> WriteOnlyCell:
        celda = WriteOnlyCell(hoja, value=texto)
        celda.font = NEGRITA
        return celda

    def guardar(self) -> Path:
        """Guarda el libro en su destino y devuelve la ruta."""
        self.ruta.parent.mkdir(parents=True, exist_ok=True)
        self.libro.save(self.ruta)
        return self.ruta
