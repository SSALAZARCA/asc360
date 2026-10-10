"""
Inventory counts, test counts (odd/tasks/motored-conteo-prueba.md).

ADMIN may schedule a count as a test (`es_prueba`). It behaves like a real
one, but only ADMIN sees it (any other role gets 404, lists leave it out),
the ADMIN list hides it unless asked, its Excel says it must not reach the
ERP, it never blocks a real count and ADMIN may hard-delete it.

Same seams as `test_conteos_api.py`: the leading `[]` of every execute
queue is the DB probe, the conteo comes from `get_queue`.
"""
import datetime
import hashlib
import io
import uuid
from decimal import Decimal
from unittest.mock import AsyncMock

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import settings
from app.main import app
from app.motored.models.conteo import Conteo
from app.motored.models.conteo_sesion import ConteoIntegrante, ConteoSesion
from app.motored.models.sucursal import Sucursal
from app.motored.models.usuario import MotoredRole, Usuario
from app.motored.schemas import conteos as esquemas
from app.motored.services.auth import MotoredUser
from app.motored.services.conteos import (
    cierre, consultas, errores, excel_ajustes, panel, sesiones, snapshot,
)
from tests.motored.conftest import (
    FakeAsyncSession, override_motored_db, override_motored_user,
)

BASE = "/api/motored/conteos"
PUBLICA = "/api/motored/publico/conteos"
UTC = datetime.timezone.utc
AHORA = datetime.datetime(2026, 10, 10, 13, 0, tzinfo=UTC)
ADMIN_ID = uuid.uuid4()
LIDER_ID = uuid.uuid4()
NO_ADMIN = ("LIDER_INVENTARIOS", "COORDINADOR_REPUESTOS", "GERENCIA")
TOKEN = "token-de-dispositivo-" + "x" * 30
SLUG = "SlugDePrueba1234"


@pytest.fixture(autouse=True)
def _listo(monkeypatch):
    monkeypatch.setattr(settings, "MOTORED_ENABLED", True)
    monkeypatch.setattr(settings, "MOTORED_SECRET_KEY", "conteos-prueba")
    monkeypatch.setattr(settings, "SECRET_KEY", "conteos-prueba-asc360")
    monkeypatch.setattr(
        settings, "MOTORED_PUBLIC_URL", "https://motored.test/")
    yield
    app.dependency_overrides.clear()


@pytest.fixture
def datos(monkeypatch):
    monkeypatch.setattr(consultas, "datos_conteo", AsyncMock(
        return_value=consultas.DatosConteo("Quilichao", "Lina", None)))


def _usuario(rol):
    ids = {"ADMIN": ADMIN_ID}
    return MotoredUser(
        user_id=str(ids.get(rol, LIDER_ID)), role=rol)


def _conteo(estado="EN_CONTEO", es_prueba=True):
    abierto = estado in ("EN_CONTEO", "EN_RECONTEO")
    return Conteo(
        id=uuid.uuid4(), tipo="TOTAL", estado=estado, origen="MANUAL",
        sucursal_id=uuid.uuid4(), lider_id=LIDER_ID, es_prueba=es_prueba,
        fecha_programada=AHORA.date(), created_at=AHORA,
        enlace_slug=SLUG if abierto else None,
        codigo_hash="f" * 64 if abierto else None,
        snapshot_tomado_en=AHORA if abierto else None,
        umbral_reconteo_pesos=Decimal("100000.00"),
        umbral_critico_pesos=Decimal("500000.00"))


def _llamar(rol, metodo, ruta, conteo=None, json=None, cola=None):
    override_motored_user(_usuario(rol))
    sesion = FakeAsyncSession(
        execute_queue=[[]] + list(cola or []),
        get_queue=[conteo] if conteo is not None else [None])
    override_motored_db(sesion)
    respuesta = TestClient(app).request(metodo, BASE + ruta, json=json)
    return respuesta, sesion


# --- scheduling -------------------------------------------------------------


def test_the_schedule_body_defaults_to_a_real_count():
    cuerpo = esquemas.ProgramarEntrada(
        sucursal_id=uuid.uuid4(), lider_id=uuid.uuid4(),
        fecha_programada=AHORA.date())

    assert cuerpo.es_prueba is False


async def test_the_service_stores_the_test_flag():
    tienda = Sucursal(id=uuid.uuid4(), nombre="S", activa=True)
    lider = Usuario(
        id=uuid.uuid4(), nombre="A", email="a@x.com", hashed_password="h",
        role=MotoredRole.ADMIN, activo=True, status="approved")
    db = FakeAsyncSession(get_queue=[tienda, lider])

    conteo = await snapshot.programar_conteo(
        db, tienda.id, lider.id, AHORA.date(), uuid.uuid4(),
        es_prueba=True)

    assert conteo.es_prueba is True


def test_admin_schedules_a_test_count(monkeypatch, datos):
    conteo = _conteo(estado="PROGRAMADO")
    programar = AsyncMock(return_value=conteo)
    monkeypatch.setattr(snapshot, "programar_conteo", programar)
    cuerpo = {"sucursal_id": str(conteo.sucursal_id),
              "lider_id": str(LIDER_ID), "fecha_programada": "2026-10-12",
              "es_prueba": True}

    r, _ = _llamar("ADMIN", "POST", "", json=cuerpo)

    assert r.status_code == 201, r.text
    assert r.json()["es_prueba"] is True
    assert programar.await_args.kwargs["es_prueba"] is True


@pytest.mark.parametrize("rol", NO_ADMIN)
def test_no_other_role_may_set_the_test_flag(monkeypatch, rol):
    programar = AsyncMock()
    monkeypatch.setattr(snapshot, "programar_conteo", programar)
    cuerpo = {"sucursal_id": str(uuid.uuid4()),
              "lider_id": str(LIDER_ID), "fecha_programada": "2026-10-12",
              "es_prueba": True}

    r, _ = _llamar(rol, "POST", "", json=cuerpo)

    assert r.status_code == 403, r.text
    programar.assert_not_awaited()


# --- visibility -------------------------------------------------------------


def test_admin_sees_a_test_count_and_its_flag(datos):
    conteo = _conteo()

    r, _ = _llamar("ADMIN", "GET", f"/{conteo.id}", conteo)

    assert r.status_code == 200, r.text
    assert r.json()["es_prueba"] is True


def test_a_real_count_says_it_is_not_a_test(datos):
    conteo = _conteo(es_prueba=False)

    r, _ = _llamar("LIDER_INVENTARIOS", "GET", f"/{conteo.id}", conteo)

    assert r.status_code == 200, r.text
    assert r.json()["es_prueba"] is False


@pytest.mark.parametrize("rol", NO_ADMIN)
def test_a_test_count_is_a_404_for_every_other_role(datos, rol):
    conteo = _conteo()

    r, _ = _llamar(rol, "GET", f"/{conteo.id}", conteo)

    assert r.status_code == 404, r.text


def test_an_unsaved_flag_reads_as_a_real_count():
    conteo = _conteo(es_prueba=None)

    assert consultas.ve_conteo(_usuario("GERENCIA"), conteo)


@pytest.mark.parametrize("rol", NO_ADMIN)
def test_the_panel_of_a_test_count_is_a_404(monkeypatch, rol):
    monkeypatch.setattr(panel, "huella", AsyncMock(
        return_value=panel.Huella(LIDER_ID, 42, True)))

    r, _ = _llamar(rol, "GET", f"/{uuid.uuid4()}/panel?version=42")

    assert r.status_code == 404, r.text


def test_admin_polls_the_panel_of_a_test_count(monkeypatch):
    monkeypatch.setattr(panel, "huella", AsyncMock(
        return_value=panel.Huella(LIDER_ID, 42, True)))

    r, _ = _llamar("ADMIN", "GET", f"/{uuid.uuid4()}/panel?version=42")

    assert r.status_code == 200, r.text
    assert r.json() == {"version": 42, "sin_cambios": True}


@pytest.mark.parametrize("rol,consulta,incluye", [
    ("ADMIN", "", False),
    ("ADMIN", "?incluir_pruebas=true", True),
    ("GERENCIA", "?incluir_pruebas=true", False),
    ("LIDER_INVENTARIOS", "?incluir_pruebas=true", False),
])
def test_the_list_hides_test_counts_unless_admin_asks(
        monkeypatch, rol, consulta, incluye):
    conteo = _conteo(estado="CERRADO")
    listar = AsyncMock(return_value=[(conteo, "Quilichao", "Lina")])
    monkeypatch.setattr(consultas, "listar", listar)
    monkeypatch.setattr(panel, "progreso", AsyncMock(return_value={}))

    r, _ = _llamar(rol, "GET", consulta)

    assert r.status_code == 200, r.text
    assert listar.await_args.kwargs["incluir_pruebas"] is incluye
    assert r.json()[0]["es_prueba"] is True


# --- the open-count rule ----------------------------------------------------


class _Escalar:
    """Records the one `db.scalar` query of the open-count check."""

    def __init__(self):
        self.consultas = []

    async def scalar(self, consulta):
        self.consultas.append(consulta)


async def test_a_test_count_skips_the_open_count_check():
    db = _Escalar()

    await snapshot._exigir_sin_total_abierto(db, _conteo(estado="PROGRAMADO"))

    assert db.consultas == []


async def test_a_real_count_ignores_open_test_counts():
    db = _Escalar()

    await snapshot._exigir_sin_total_abierto(
        db, _conteo(estado="PROGRAMADO", es_prueba=False))

    sql = str(db.consultas[0].compile(
        compile_kwargs={"literal_binds": True}))
    assert "es_prueba" in sql


# --- Excel ------------------------------------------------------------------


def test_a_test_count_excel_name_starts_with_prueba():
    nombre = excel_ajustes.nombre_archivo(
        "ajustes", "045", AHORA.date(), prueba=True)

    assert nombre == "PRUEBA_ajustes_conteo_045_2026-10-10.xlsx"
    assert excel_ajustes.nombre_archivo(
        "ajustes", "045", AHORA.date()).startswith("ajustes_")


def _encabezado():
    kpi = cierre.Kpi(
        refs_universo=1, refs_exactas=0, exactitud_pct=Decimal("0"),
        valor_sistema=Decimal("1000"), valor_diferencia_neta=Decimal("-1"),
        valor_diferencia_abs=Decimal("1"))
    return cierre.Encabezado(
        tienda="Quilichao", codigo_co="045", lider="Lina",
        estado="CERRADO", fecha_programada=AHORA.date(),
        fecha_corte=AHORA.date(), iniciado_en=AHORA,
        ronda_terminada_en=AHORA, cerrado_en=AHORA, bodega="BX001",
        kpi=kpi, motivo_cierre_forzado=None)


def _linea():
    return cierre.Linea(
        referencia_id=uuid.uuid4(), codigo="R-1", descripcion="Tornillo",
        bodega="BX001", sistema=Decimal("5"), contado=Decimal("3"),
        diferencia=Decimal("-2"), costo_unitario=Decimal("1000"),
        costo_fuente="BODEGA", valor=Decimal("-2000"), ubicaciones=["A1"],
        con_reconteo=False, critico=False, confirmada=None)


def test_a_test_count_excel_says_do_not_load_it():
    libro = load_workbook(io.BytesIO(excel_ajustes.libro(
        _encabezado(), [_linea()], prueba=True)))

    for hoja in libro.worksheets:
        assert hoja["A1"].value == excel_ajustes.AVISO_PRUEBA
    assert excel_ajustes.AVISO_PRUEBA == "PRUEBA – NO CARGAR AL ERP"
    assert libro["Ajustes"]["A2"].value == "Referencia"
    assert libro["Ajustes"]["A3"].value == "R-1"


def test_a_real_count_excel_keeps_its_layout():
    libro = load_workbook(io.BytesIO(excel_ajustes.libro(
        _encabezado(), [_linea()])))

    assert libro["Ajustes"]["A1"].value == "Referencia"
    assert libro["Resumen"]["A1"].value == "Tienda"


# --- delete -----------------------------------------------------------------


def test_admin_deletes_a_test_count(monkeypatch):
    conteo = _conteo()
    borrar = AsyncMock()
    monkeypatch.setattr(snapshot, "borrar_conteo_prueba", borrar)

    r, sesion = _llamar("ADMIN", "DELETE", f"/{conteo.id}", conteo)

    assert r.status_code == 204, r.text
    assert borrar.await_args.args[1] == conteo.id
    assert sesion.committed


def test_a_real_count_can_never_be_deleted(monkeypatch):
    conteo = _conteo(es_prueba=False)
    monkeypatch.setattr(snapshot, "borrar_conteo_prueba", AsyncMock(
        side_effect=errores.NoEsPrueba()))

    r, sesion = _llamar("ADMIN", "DELETE", f"/{conteo.id}", conteo)

    assert r.status_code == 409, r.text
    assert r.json()["detail"]["code"] == "NO_ES_PRUEBA"
    assert not sesion.committed


@pytest.mark.parametrize("rol", NO_ADMIN)
def test_only_admin_deletes(monkeypatch, rol):
    borrar = AsyncMock()
    monkeypatch.setattr(snapshot, "borrar_conteo_prueba", borrar)

    r, _ = _llamar(rol, "DELETE", f"/{uuid.uuid4()}", _conteo())

    assert r.status_code == 403, r.text
    borrar.assert_not_awaited()


async def test_the_service_refuses_a_real_count(monkeypatch):
    conteo = _conteo(estado="CERRADO", es_prueba=False)
    db = FakeAsyncSession()
    monkeypatch.setattr(
        snapshot.acceso, "bloquear_conteo", AsyncMock(return_value=conteo))

    with pytest.raises(errores.NoEsPrueba):
        await snapshot.borrar_conteo_prueba(db, conteo.id)

    assert db.executed_statements == []


# --- the pair device --------------------------------------------------------


def _sesion_pareja(conteo):
    return ConteoSesion(
        id=uuid.uuid4(), conteo_id=conteo.id, tipo="PAREJA",
        token_hash=hashlib.sha256(TOKEN.encode()).hexdigest(),
        estado="CONECTADA", dispositivo="ESCRITORIO", conectada_en=AHORA,
        ultima_actividad_en=datetime.datetime.now(UTC))


def test_the_pair_session_says_it_is_a_test(monkeypatch):
    conteo = _conteo()
    sesion = _sesion_pareja(conteo)

    async def _describir(db, s):
        uno = ConteoIntegrante(orden=1, nombre="Ana Ruiz", cedula="1")
        return sesiones.FilaSesion(s, 1, [uno], None)
    monkeypatch.setattr(sesiones, "describir", _describir)
    override_motored_db(FakeAsyncSession(
        execute_queue=[[], [(sesion, conteo, "Quilichao")]]))

    r = TestClient(app).get(
        f"{PUBLICA}/{SLUG}/sesion",
        headers={"Authorization": f"Bearer {TOKEN}"})

    assert r.status_code == 200, r.text
    assert r.json()["es_prueba"] is True
