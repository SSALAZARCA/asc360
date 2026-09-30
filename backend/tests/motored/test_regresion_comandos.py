"""
Motored Pedidos F3 "Motor" (S8b-1, ADR-10) — subcomandos de
`scripts/motored_regresion.py`: `nivel-a`, `nivel-b`, `delta` y
`exportar-corrida`.

Con libros sintéticos en `tmp_path`. Las salidas van a una carpeta (fuera de
cualquier árbol git); una carpeta dentro del repositorio se rechaza sin
escribir nada. Los modos con base de datos (nivel B, delta y exportación de una
corrida guardada) se prueban en `pg_real/test_regresion_nivel_b_pg.py`.
"""
import io
import json
import os
import subprocess
import sys
from pathlib import Path

import openpyxl

from app.motored.herramientas.regresion import comandos
from app.motored.herramientas.regresion.comparador import (
    atributos_de,
    entradas_de,
)
from app.motored.herramientas.regresion.extractor_excel import leer_libro
from app.motored.herramientas.regresion.rutas import (
    ENV_ACEPTACIONES,
    ENV_DB_URL,
    ENV_EXCEL,
    ENV_SALIDA,
)
from app.motored.services.motor.motor import calcular_sucursal
from app.motored.services.motor.tipos import ParametrosMotor
from tests.motored.fixtures.regresion.conjuntos import filas_base
from tests.motored.fixtures.regresion.libro_sintetico import escribir_libro
from tests.motored.fixtures.regresion.oraculo_excel import CLASES

BACKEND = Path(__file__).resolve().parents[2]


def _libro(tmp_path, **opciones) -> Path:
    ruta = tmp_path / "datos" / "libro.xlsx"
    ruta.parent.mkdir(exist_ok=True)
    escribir_libro(ruta, filas_base(), etiquetas=CLASES, **opciones)
    return ruta


def _correr(argumentos, entorno=None):
    salida = io.StringIO()
    codigo = comandos.main(argumentos, entorno or {}, salida)
    return codigo, salida.getvalue()


def _hojas(ruta: Path) -> list:
    libro = openpyxl.load_workbook(ruta, read_only=True)
    nombres = list(libro.sheetnames)
    libro.close()
    return nombres


def _filas(ruta: Path, hoja: str) -> list:
    libro = openpyxl.load_workbook(ruta, read_only=True)
    filas = list(libro[hoja].iter_rows(values_only=True))
    libro.close()
    ancho = max((len(fila) for fila in filas), default=0)
    return [fila + (None,) * (ancho - len(fila)) for fila in filas]


class TestNivelA:
    def test_escribe_el_json_y_el_libro_en_la_carpeta_de_salida(
        self, tmp_path
    ):
        salida = tmp_path / "reportes"
        codigo, texto = _correr([
            "nivel-a", "--excel", str(_libro(tmp_path)),
            "--salida", str(salida),
        ])
        assert codigo == 0
        informe = json.loads((salida / "nivel_a.json").read_text("utf-8"))
        assert informe["filas"] == 6
        assert _hojas(salida / "comparacion_nivel_a.xlsx")[:5] == [
            "Resumen", "Antigüedad de datos", "Filas", "A1", "A2",
        ]
        assert "A1: 6 de 6 filas exactas" in texto
        assert "comparacion_nivel_a.xlsx" in texto

    def test_un_nivel_con_sin_categoria_sale_con_codigo_1_pero_escribe(
        self, tmp_path
    ):
        ruta = _libro(tmp_path, sobrescribe={"A-100": {"AB": 1}})
        salida = tmp_path / "reportes"
        codigo, texto = _correr([
            "nivel-a", "--excel", str(ruta), "--salida", str(salida),
        ])
        assert codigo == 1
        assert "FALLA" in texto
        assert (salida / "comparacion_nivel_a.xlsx").exists()

    def test_las_rutas_pueden_venir_del_entorno(self, tmp_path):
        salida = tmp_path / "reportes"
        codigo, _ = _correr(["nivel-a"], {
            ENV_EXCEL: str(_libro(tmp_path)), ENV_SALIDA: str(salida),
        })
        assert codigo == 0 and (salida / "nivel_a.json").exists()

    def test_las_aceptaciones_se_aplican(self, tmp_path):
        ruta = _libro(tmp_path, sobrescribe={"A-100": {"AB": 1}})
        aceptaciones = tmp_path / "aceptaciones.json"
        aceptaciones.write_text('{"A-100": "T7"}', encoding="utf-8")
        codigo, _ = _correr([
            "nivel-a", "--excel", str(ruta), "--salida",
            str(tmp_path / "r"), "--aceptaciones", str(aceptaciones),
        ])
        assert codigo == 0
        assert ENV_ACEPTACIONES == "MOTORED_REGRESION_ACEPTACIONES"

    def test_una_salida_dentro_del_repositorio_se_rechaza_sin_escribir(
        self, tmp_path
    ):
        dentro = Path(__file__).parent / "salida_prohibida"
        codigo, texto = _correr([
            "nivel-a", "--excel", str(_libro(tmp_path)),
            "--salida", str(dentro),
        ])
        assert codigo == 2
        assert "repositorio" in texto
        assert not dentro.exists()

    def test_sin_rutas_explica_cuales_faltan(self):
        codigo, texto = _correr(["nivel-a"])
        assert codigo == 2
        assert ENV_EXCEL in texto and ENV_SALIDA in texto

    def test_un_libro_inexistente_es_un_error_de_uso(self, tmp_path):
        codigo, texto = _correr([
            "nivel-a", "--excel", str(tmp_path / "no_existe.xlsx"),
            "--salida", str(tmp_path / "r"),
        ])
        assert codigo == 2 and "no_existe.xlsx" in texto


class TestDelta:
    def test_con_el_libro_mide_dias_y_explica_lo_que_no_puede_medir(
        self, tmp_path
    ):
        salida = tmp_path / "reportes"
        codigo, texto = _correr([
            "delta", "--excel", str(_libro(tmp_path)),
            "--salida", str(salida),
        ])
        ruta = salida / "efecto_switches.xlsx"
        hojas = _hojas(ruta)
        assert codigo == 0
        assert "Delta dias_entre_pedidos=7" in hojas
        assert "Delta dias_entre_pedidos=15" in hojas
        assert "Delta perdida x1" not in hojas
        assert "dias_entre_pedidos=7: MEDIDO" in texto
        assert "perdida x1: NO MEDIDO" in texto
        assert "demanda perdida" in texto

    def test_acepta_otros_factores_y_otros_periodos_de_revision(
        self, tmp_path
    ):
        salida = tmp_path / "reportes"
        _correr([
            "delta", "--excel", str(_libro(tmp_path)),
            "--salida", str(salida), "--factores", "1,2", "--dias", "10",
        ])
        filas = _filas(salida / "efecto_switches.xlsx", "Resumen switches")
        nombres = [f[0] for f in filas[1:]]
        assert "perdida x2" in nombres
        assert "dias_entre_pedidos=10" in nombres
        assert "dias_entre_pedidos=7" not in nombres

    def test_no_mide_ningun_switch_si_el_nivel_a_no_pasa(self, tmp_path):
        ruta = _libro(tmp_path, sobrescribe={"A-100": {"AB": 1}})
        salida = tmp_path / "reportes"
        codigo, texto = _correr([
            "delta", "--excel", str(ruta), "--salida", str(salida),
        ])
        hojas = _hojas(salida / "efecto_switches.xlsx")
        assert codigo == 1
        assert "nivel A" in texto
        assert not any(h.startswith("Delta ") for h in hojas)
        assert "Resumen switches" not in hojas

    def test_sin_libro_ni_base_de_datos_no_hay_fuente(self, tmp_path):
        codigo, texto = _correr(["delta", "--salida", str(tmp_path / "r")])
        assert codigo == 2 and ENV_EXCEL in texto

    def test_la_medicion_con_base_exige_la_url_y_el_corte(self, tmp_path):
        codigo, texto = _correr([
            "delta", "--sucursal-id", "00000000-0000-0000-0000-000000000001",
            "--salida", str(tmp_path / "r"),
        ])
        assert codigo == 2 and ENV_DB_URL in texto


class TestExportarCorrida:
    def test_arma_el_libro_de_corrida_con_los_insumos_del_libro(
        self, tmp_path
    ):
        ruta = _libro(tmp_path)
        salida = tmp_path / "reportes"
        codigo, texto = _correr([
            "exportar-corrida", "--excel", str(ruta),
            "--salida", str(salida),
        ])
        libro = salida / "corrida_SUCURSAL_SINTETICA.xlsx"
        assert codigo == 0 and libro.exists()
        assert str(libro) in texto
        hojas = _hojas(libro)
        assert hojas[:3] == [
            "Resumen clases", "Antigüedad de datos", "SUCURSAL SINTETICA",
        ]
        antiguedad = _filas(libro, "Antigüedad de datos")
        assert any(f[0] == "Inventario" and f[4] for f in antiguedad)

    def test_las_unidades_son_las_del_motor_con_divisor_21(self, tmp_path):
        ruta = _libro(tmp_path)
        salida = tmp_path / "reportes"
        _correr([
            "exportar-corrida", "--excel", str(ruta),
            "--salida", str(salida),
        ])
        lectura = leer_libro(ruta)
        esperado = calcular_sucursal(
            entradas_de(lectura), atributos_de(lectura), ParametrosMotor()
        ).resumen.total.unidades
        filas = _filas(
            salida / "corrida_SUCURSAL_SINTETICA.xlsx", "Resumen clases"
        )
        total = next(f for f in filas if f[1] == "TOTAL")
        assert total[2] == float(esperado)

    def test_sin_fuente_es_un_error_de_uso(self, tmp_path):
        codigo, texto = _correr(
            ["exportar-corrida", "--salida", str(tmp_path / "r")]
        )
        assert codigo == 2 and ENV_EXCEL in texto


class TestNivelB:
    def test_exige_la_url_de_la_base_de_pruebas(self, tmp_path):
        codigo, texto = _correr([
            "nivel-b", "--excel", str(_libro(tmp_path)),
            "--salida", str(tmp_path / "r"),
            "--sucursal-id", "00000000-0000-0000-0000-000000000001",
            "--fecha-corte", "2026-09-21",
        ])
        assert codigo == 2 and ENV_DB_URL in texto

    def test_exige_el_corte_si_no_se_indica_una_corrida(self, tmp_path):
        codigo, texto = _correr([
            "nivel-b", "--excel", str(_libro(tmp_path)),
            "--salida", str(tmp_path / "r"),
            "--sucursal-id", "00000000-0000-0000-0000-000000000001",
            "--db-url", "postgresql+asyncpg://x@127.0.0.1:1/x",
        ])
        assert codigo == 2 and "--fecha-corte" in texto

    def test_una_fecha_mal_escrita_es_un_error_de_uso(self, tmp_path):
        codigo, texto = _correr([
            "nivel-b", "--excel", str(_libro(tmp_path)),
            "--salida", str(tmp_path / "r"),
            "--sucursal-id", "00000000-0000-0000-0000-000000000001",
            "--db-url", "postgresql+asyncpg://x@127.0.0.1:1/x",
            "--fecha-corte", "21/09/2026",
        ])
        assert codigo == 2 and "--fecha-corte" in texto


class TestSubcomandos:
    def test_sin_subcomando_es_un_error_de_uso(self):
        codigo, _ = _correr([])
        assert codigo == 2

    def test_un_subcomando_desconocido_es_un_error_de_uso(self):
        codigo, _ = _correr(["inventado"])
        assert codigo == 2

    def test_el_script_lista_los_cuatro_subcomandos(self):
        proceso = subprocess.run(
            [sys.executable, "scripts/motored_regresion.py", "--help"],
            cwd=BACKEND, capture_output=True, text=True, timeout=120,
        )
        assert proceso.returncode == 0
        for nombre in ("nivel-a", "nivel-b", "delta", "exportar-corrida"):
            assert nombre in proceso.stdout

    def test_los_modos_con_libro_no_necesitan_la_configuracion_de_la_app(
        self, tmp_path
    ):
        """Sin DATABASE_URL, SECRET_KEY, etc. (el entorno de producción):
        el libro de Excel no debe arrastrar `app.config`."""
        salida = tmp_path / "reportes"
        entorno = {"PATH": os.environ.get("PATH", "")}
        for subcomando in ("nivel-a", "delta", "exportar-corrida"):
            proceso = subprocess.run(
                [sys.executable, "scripts/motored_regresion.py", subcomando,
                 "--excel", str(_libro(tmp_path)), "--salida", str(salida)],
                cwd=BACKEND, capture_output=True, text=True, timeout=300,
                env=entorno,
            )
            assert proceso.returncode == 0, proceso.stderr[-400:]

    def test_el_script_sin_argumentos_sale_con_codigo_2(self):
        proceso = subprocess.run(
            [sys.executable, "scripts/motored_regresion.py"],
            cwd=BACKEND, capture_output=True, text=True, timeout=120,
        )
        assert proceso.returncode == 2
