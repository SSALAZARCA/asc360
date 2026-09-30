"""
Motored Pedidos F3 "Motor" (S8a-1, ADR-10) — extractor de la hoja `Pedido`
del Excel de referencia, probado con un libro sintético generado en
`tmp_path` (jamás con datos reales).

Dos pasadas `read_only`: texto de fórmula (divisor /18 vs /21, S estática) y
valores cacheados (E:J, K, T..Z, N..AD, AA1:AE11, `Meses de cobertura`).
"""
from decimal import Decimal

import openpyxl
import pytest

from app.motored.herramientas.regresion import extractor_excel
from app.motored.herramientas.regresion.extractor_excel import (
    ErrorExtractor,
    informe_empates,
    leer_libro,
)
from tests.motored.fixtures.regresion.conjuntos import (
    CODIGO_18,
    filas_base,
    filas_con_empate,
)
from tests.motored.fixtures.regresion.libro_sintetico import escribir_libro
from tests.motored.fixtures.regresion.oraculo_excel import (
    CLASES,
    ETIQUETAS_EXCEL,
    FilaSintetica,
)


def _dec(valor) -> Decimal:
    return Decimal(repr(valor))


@pytest.fixture
def base(tmp_path):
    ruta = tmp_path / "libro.xlsx"
    calculo = escribir_libro(ruta, filas_base())
    return leer_libro(ruta), calculo


class TestFilas:
    def test_lee_las_filas_con_dato_en_orden_fisico(self, base):
        lectura, _ = base
        assert [f.codigo for f in lectura.filas] == [
            "A-100", "B-200", "C-300", "D-400", "E-500", "F-600",
        ]
        assert [f.fila for f in lectura.filas] == [14, 15, 16, 17, 18, 19]

    def test_insumos_e_a_j_k_t_u_v_w_x_z(self, base):
        lectura, _ = base
        segunda = lectura.filas[1]
        assert segunda.ventas == (Decimal(22),) * 6
        assert segunda.unidad_empaque == 12
        assert segunda.precio == Decimal("55.5")
        assert (segunda.inventario, segunda.transito, segunda.backorder) == (
            Decimal(40), Decimal(0), Decimal(3),
        )
        assert lectura.filas[4].ajuste == Decimal(15)
        assert lectura.filas[0].perdida == Decimal(0)

    def test_precio_vacio_es_none(self, base):
        lectura, _ = base
        assert lectura.filas[3].codigo == "D-400"
        assert lectura.filas[3].precio is None
        assert lectura.filas[0].precio == Decimal(100)

    def test_valores_cacheados_de_n_a_ad_son_los_del_excel(self, base):
        lectura, calculo = base
        for fila, esperado in zip(lectura.filas, calculo.filas):
            assert fila.n == _dec(esperado["N"])
            assert fila.peso == _dec(esperado["O"])
            assert fila.acumulado == _dec(esperado["P"])
            assert (fila.clase_abc, fila.clase_fms, fila.clase_estatica) == (
                esperado["Q"], esperado["R"], esperado["S"],
            )
            assert fila.inventario_final == _dec(esperado["Y"])
            assert fila.stock_objetivo == _dec(esperado["AA"])
            assert fila.pedido == _dec(esperado["AB"])
            assert fila.valor_pedido == _dec(esperado["AC"])
            assert fila.cobertura_final == _dec(esperado["AD"])

    def test_pedido_de_la_fila_con_ajuste_positivo_es_z(self, base):
        lectura, _ = base
        assert lectura.filas[4].pedido == Decimal(15)

    def test_divisor_18_sale_del_texto_de_la_formula(self, base):
        lectura, _ = base
        por_codigo = {f.codigo: f for f in lectura.filas}
        assert por_codigo[CODIGO_18].divisor == 18
        assert {f.divisor for f in lectura.filas if f.codigo != CODIGO_18} == {
            21
        }

    def test_s_estatica_y_s_con_formula(self, tmp_path):
        ruta = tmp_path / "s.xlsx"
        escribir_libro(ruta, filas_base()[:2], s_formula=True)
        assert all(f.s_es_formula for f in leer_libro(ruta).filas)
        escribir_libro(ruta, filas_base()[:2])
        assert not any(f.s_es_formula for f in leer_libro(ruta).filas)

    def test_error_cacheado_en_ad_queda_como_none(self, tmp_path):
        ruta = tmp_path / "cero.xlsx"
        cero = FilaSintetica("Z-000", (0,) * 6, precio=5.0, inventario=3)
        escribir_libro(ruta, [filas_base()[0], cero])
        fila = leer_libro(ruta).filas[1]
        assert fila.cobertura_final is None
        assert fila.n == Decimal(0) and fila.clase_abc == "D"


class TestSucursalYCoberturas:
    def test_atributos_de_la_sucursal(self, base):
        lectura, _ = base
        sucursal = lectura.sucursal
        assert sucursal.nombre == "SUCURSAL SINTETICA"
        assert (sucursal.bodega, sucursal.sic) == ("BE999", 1999)
        assert sucursal.dias_empaque == Decimal(3)
        assert sucursal.dias_transito == Decimal(2)
        assert sucursal.dias_seguridad == Decimal("2.5")

    def test_factores_de_cobertura_e3_n3(self, base):
        lectura, calculo = base
        assert list(lectura.coberturas) == list(CLASES)
        for clase in CLASES:
            assert lectura.coberturas[clase] == _dec(calculo.coberturas[clase])
        assert lectura.coberturas["AF"] == Decimal("1.75")

    def test_resumen_aa1_ae11(self, base):
        lectura, calculo = base
        assert [r.etiqueta for r in lectura.resumen] == list(ETIQUETAS_EXCEL)
        primera = lectura.resumen[0]
        assert primera.unidades == _dec(calculo.resumen[0][1])
        assert primera.referencias == calculo.resumen[0][2]
        assert primera.valor == _dec(calculo.resumen[0][3])
        assert lectura.total.etiqueta == "Total general"
        assert lectura.total.unidades == _dec(calculo.total[1])


class TestErrores:
    def test_hoja_pedido_ausente(self, tmp_path):
        ruta = tmp_path / "vacio.xlsx"
        libro = openpyxl.Workbook()
        libro.active.title = "Otra"
        libro.save(ruta)
        with pytest.raises(ErrorExtractor, match="Pedido"):
            leer_libro(ruta)

    def test_sucursal_ausente_de_meses_de_cobertura(self, tmp_path):
        ruta = tmp_path / "s.xlsx"
        escribir_libro(ruta, filas_base()[:1])
        libro = openpyxl.load_workbook(ruta)
        libro["Pedido"]["D4"] = "NO EXISTE"
        libro.save(ruta)
        with pytest.raises(ErrorExtractor, match="NO EXISTE"):
            leer_libro(ruta)

    def test_referencia_repetida(self, tmp_path):
        ruta = tmp_path / "repetida.xlsx"
        escribir_libro(ruta, [filas_base()[0], filas_base()[0]])
        with pytest.raises(ErrorExtractor, match="A-100"):
            leer_libro(ruta)

    def test_hoja_pedido_sin_encabezado_no_tiene_sucursal(self, tmp_path):
        ruta = tmp_path / "corto.xlsx"
        libro = openpyxl.Workbook()
        libro.active.title = "Pedido"
        libro.create_sheet("Meses de cobertura")
        libro.save(ruta)
        with pytest.raises(ErrorExtractor, match="sucursal"):
            leer_libro(ruta)

    def test_unidad_de_empaque_no_entera(self, tmp_path):
        ruta = tmp_path / "u.xlsx"
        escribir_libro(
            ruta, filas_base()[:1],
            sobrescribe={"A-100": {"U": 2.5}},
        )
        with pytest.raises(ErrorExtractor, match="A-100"):
            leer_libro(ruta)


class TestModoDeLectura:
    def test_dos_pasadas_read_only_con_y_sin_data_only(
        self, tmp_path, monkeypatch
    ):
        ruta = tmp_path / "libro.xlsx"
        escribir_libro(ruta, filas_base()[:2])
        llamadas = []
        original = openpyxl.load_workbook

        def espia(*args, **kwargs):
            llamadas.append(kwargs)
            return original(*args, **kwargs)

        monkeypatch.setattr(extractor_excel.openpyxl, "load_workbook", espia)
        leer_libro(ruta)
        assert llamadas and all(c.get("read_only") is True for c in llamadas)
        assert {c.get("data_only") for c in llamadas} == {True, False}


class TestInformeEmpates:
    def test_compara_el_orden_fisico_con_codigo_ascendente(self, tmp_path):
        ruta = tmp_path / "empates.xlsx"
        escribir_libro(ruta, filas_con_empate())
        grupos = informe_empates(leer_libro(ruta).filas)
        assert len(grupos) == 1
        assert grupos[0].codigos_excel == ("G-2", "G-1")
        assert grupos[0].codigos_asc == ("G-1", "G-2")
        assert grupos[0].coincide is False

    def test_empate_en_orden_ascendente_coincide(self, tmp_path):
        ruta = tmp_path / "empates.xlsx"
        filas = filas_con_empate()
        escribir_libro(ruta, [filas[0], filas[2], filas[1], filas[3]])
        grupos = informe_empates(leer_libro(ruta).filas)
        assert [g.coincide for g in grupos] == [True]

    def test_las_filas_con_n_cero_no_forman_grupos(self, tmp_path):
        ruta = tmp_path / "ceros.xlsx"
        ceros = [
            FilaSintetica("Z-0", (0,) * 6, inventario=1),
            FilaSintetica("Z-1", (0,) * 6, inventario=2),
        ]
        escribir_libro(ruta, [filas_base()[0]] + ceros)
        lectura = leer_libro(ruta)
        assert [f.n for f in lectura.filas[1:]] == [Decimal(0), Decimal(0)]
        assert informe_empates(lectura.filas) == ()

    def test_sin_empates_no_hay_grupos(self, base):
        lectura, _ = base
        assert informe_empates(lectura.filas) == ()
