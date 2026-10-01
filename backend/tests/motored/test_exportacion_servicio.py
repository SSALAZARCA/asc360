"""
Fase 4 "Pantallas del pedido" (sdd/motored-pedidos-ui, B4, ADR-7, spec
EX-01..EX-15, decisiones F4-1 y A5): qué tiendas se exportan, en qué orden se
chequea y cómo se arma el archivo.

Con una sesión de juguete que graba el SQL literal se prueba lo que no
depende del motor de Postgres: el ORDEN de los chequeos (404, 042, 040,
041, 065, 055, 056, 057), el orden de los bloqueos (corrida SHARE -> tiendas
SHARE por `sucursal_id`), que las líneas se piden SÓLO de las tiendas que se
exportan y sólo con cantidad mayor que 0 y no excluidas, y que un rechazo
no consulta nada más. La construcción de los archivos (memoria temporal) se
abre con openpyxl y con zipfile. Lo mismo contra Postgres real corre en
`pg_real/test_exportacion_pg.py`.
"""
import datetime
import io
import uuid
import zipfile
from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from sqlalchemy.dialects import postgresql

from app.motored.services.corridas import codigos, exportacion
from app.motored.services.corridas import exportacion_hmcl as hmcl
from app.motored.services.corridas.codigos import ErrorCorrida
from tests.motored.conftest import FakeAsyncSession
from tests.motored.fixtures import corridas_api as fx

CORRIDA = fx.CORRIDA_ID
SUC_A, SUC_B = fx.SUC_A, fx.SUC_B
SUC_C, SUC_D, SUC_E = (uuid.UUID(int=603), uuid.UUID(int=604),
                       uuid.UUID(int=605))
D = Decimal


def _sql(sentencia) -> str:
    return str(sentencia.compile(
        dialect=postgresql.dialect(),
        compile_kwargs={"literal_binds": True}))


def _corrida(**campos):
    base = dict(
        id=CORRIDA, codigo="PED-2026-S39-001", estado="BORRADOR",
        fecha_corte=fx.CORTE, es_escenario=False, invalidada=False)
    return SimpleNamespace(**{**base, **campos})


def _fila(sucursal_id=SUC_A, nombre="Manizales", sic="1234",
          estado_pedido="CERRADO", estado="OK"):
    tienda = SimpleNamespace(
        corrida_id=CORRIDA, sucursal_id=sucursal_id, estado=estado,
        estado_pedido=estado_pedido)
    return (tienda, nombre, sic)


# --- Una tienda --------------------------------------------------------------


async def _preparar(cola, sucursal_id=SUC_A):
    db = FakeAsyncSession(execute_queue=cola)
    return db, await exportacion.preparar_tienda(db, CORRIDA, sucursal_id)


async def _rechazada(cola, codigo, sucursal_id=SUC_A):
    db = FakeAsyncSession(execute_queue=cola)
    with pytest.raises(ErrorCorrida) as excepcion:
        await exportacion.preparar_tienda(db, CORRIDA, sucursal_id)
    assert excepcion.value.codigo == codigo
    return db, excepcion.value


async def test_a_closed_tienda_gives_its_header_data_and_lines_ex_01():
    db, datos = await _preparar([
        [_corrida()], [_fila(nombre=" Manizales ")],
        [(SUC_A, "94109-12000S", D("50.00")),
         (SUC_A, "00123-AB", D("12.00"))]])

    assert datos == hmcl.DatosTienda(
        "Manizales", "1234", fx.CORTE,
        [("94109-12000S", D("50.00")), ("00123-AB", D("12.00"))])
    assert len(db.executed_statements) == 3


async def test_a_sent_tienda_can_be_exported_again_ex_12_a5():
    _, datos = await _preparar([
        [_corrida()], [_fila(estado_pedido="ENVIADO")],
        [(SUC_A, "A-1", D("3.00"))]])

    assert datos.lineas == [("A-1", D("3.00"))]


async def test_a_legacy_cerrada_corrida_still_exports():
    _, datos = await _preparar([
        [_corrida(estado="CERRADA")], [_fila()],
        [(SUC_A, "A-1", D("3.00"))]])

    assert datos.nombre == "Manizales"


async def test_locks_are_share_corrida_then_tienda_and_lines_are_filtered():
    db, _ = await _preparar([
        [_corrida()], [_fila()], [(SUC_A, "A-1", D("3.00"))]])

    corrida, tienda, lineas = (_sql(s) for s in db.executed_statements)

    assert "FROM corrida" in corrida and "FOR SHARE" in corrida
    assert "FOR UPDATE" not in corrida
    assert "FOR SHARE OF corrida_sucursal" in tienda
    assert f"corrida_sucursal.sucursal_id IN ('{SUC_A}')" in tienda
    assert "FROM corrida_linea" in lineas
    assert f"corrida_linea.sucursal_id IN ('{SUC_A}')" in lineas
    assert "corrida_linea.motivo_exclusion IS NULL" in lineas
    assert "corrida_linea.pedido_final > 0" in lineas
    assert "FOR UPDATE" not in lineas


async def test_an_unknown_corrida_is_a_lookup_error_with_no_more_queries():
    db = FakeAsyncSession(execute_queue=[[]])

    with pytest.raises(LookupError):
        await exportacion.preparar_tienda(db, CORRIDA, SUC_A)

    assert len(db.executed_statements) == 1


async def test_a_tienda_not_in_the_corrida_is_a_lookup_error_ex_15():
    db = FakeAsyncSession(execute_queue=[[_corrida()], []])

    with pytest.raises(LookupError):
        await exportacion.preparar_tienda(db, CORRIDA, SUC_B)

    assert len(db.executed_statements) == 2


async def test_a_scenario_is_042_naming_the_action_ex_07():
    db, error = await _rechazada(
        [[_corrida(es_escenario=True)], [_fila()]],
        codigos.E_CORRIDA_ESCENARIO_NO_SE_CIERRA)

    assert "no se puede exportar" in error.mensaje
    assert len(db.executed_statements) == 2


@pytest.mark.parametrize("estado", [
    "PENDIENTE", "CALCULANDO", "FALLIDA", "ANULADA"])
async def test_a_corrida_not_calculated_is_040(estado):
    await _rechazada(
        [[_corrida(estado=estado)], [_fila()]],
        codigos.E_CORRIDA_ESTADO_NO_ADMITE)


async def test_an_invalidated_corrida_is_041():
    await _rechazada(
        [[_corrida(invalidada=True)], [_fila()]],
        codigos.E_CORRIDA_INVALIDADA)


async def test_a_failed_tienda_has_no_pedido_065_naming_it_ex_15():
    db, error = await _rechazada(
        [[_corrida()], [_fila(estado="FALLIDA", estado_pedido=None)]],
        codigos.E_CORRIDA_SIN_PEDIDO)

    assert "Manizales" in error.mensaje
    assert len(db.executed_statements) == 2


async def test_a_borrador_tienda_is_055_with_its_state_ex_06():
    db, error = await _rechazada(
        [[_corrida()], [_fila(estado_pedido="BORRADOR")]],
        codigos.E_CORRIDA_EXPORTAR_NO_CERRADO)

    assert "Manizales" in error.mensaje and "BORRADOR" in error.mensaje
    assert error.detalle["estado_pedido"] == "BORRADOR"
    assert len(db.executed_statements) == 2


async def test_a_pedido_with_every_line_at_zero_is_056_ex_08():
    db, error = await _rechazada(
        [[_corrida()], [_fila()], []], codigos.E_CORRIDA_NADA_QUE_ENVIAR)

    assert "Manizales" in error.mensaje
    assert "exportar" in error.mensaje and "enviar" not in error.mensaje
    assert len(db.executed_statements) == 3


@pytest.mark.parametrize("sic", [None, "", "   "])
async def test_a_tienda_without_sic_is_057_naming_it_ex_09(sic):
    _, error = await _rechazada(
        [[_corrida()], [_fila(sic=sic)], [(SUC_A, "A-1", D("3.00"))]],
        codigos.E_CORRIDA_EXPORTAR_SIN_SIC)

    assert "Manizales" in error.mensaje


async def test_nothing_to_export_beats_the_missing_sic():
    await _rechazada(
        [[_corrida()], [_fila(sic=None)], []],
        codigos.E_CORRIDA_NADA_QUE_ENVIAR)


async def test_the_pedido_state_check_comes_before_the_missing_sic():
    await _rechazada(
        [[_corrida()], [_fila(sic=None, estado_pedido="BORRADOR")]],
        codigos.E_CORRIDA_EXPORTAR_NO_CERRADO)


# --- La corrida entera -------------------------------------------------------


def _tiendas():
    """A y B exportables con filas, C en BORRADOR, D fallida y E cerrada sin
    nada para pedir."""
    return [
        _fila(SUC_A, "Manizales", "1234"),
        _fila(SUC_B, "Medellín Poblado", "77", estado_pedido="ENVIADO"),
        _fila(SUC_C, "Pereira", "88", estado_pedido="BORRADOR"),
        _fila(SUC_D, "Cali", "99", estado="FALLIDA", estado_pedido=None),
        _fila(SUC_E, "Armenia", "55")]


def _lineas():
    return [
        (SUC_A, "94109-12000S", D("50.00")), (SUC_A, "00123-AB", D("12.00")),
        (SUC_B, "55512-A", D("7.00"))]


async def _zip(cola, ids=None):
    db = FakeAsyncSession(execute_queue=cola)
    return db, await exportacion.preparar_corrida(db, CORRIDA, ids)


async def _zip_rechazado(cola, codigo, ids=None):
    db = FakeAsyncSession(execute_queue=cola)
    with pytest.raises(ErrorCorrida) as excepcion:
        await exportacion.preparar_corrida(db, CORRIDA, ids)
    assert excepcion.value.codigo == codigo
    return db, excepcion.value


async def test_the_zip_has_only_closed_or_sent_tiendas_with_rows_ex_02():
    db, seleccion = await _zip([[_corrida()], _tiendas(), _lineas()])

    assert seleccion.corrida.codigo == "PED-2026-S39-001"
    assert [t.nombre for t in seleccion.tiendas] == [
        "Manizales", "Medellín Poblado"]
    assert seleccion.tiendas[0] == hmcl.DatosTienda(
        "Manizales", "1234", fx.CORTE,
        [("94109-12000S", D("50.00")), ("00123-AB", D("12.00"))])
    assert seleccion.tiendas[1].lineas == [("55512-A", D("7.00"))]
    assert len(db.executed_statements) == 3


async def test_every_skipped_tienda_comes_with_its_reason_ex_02_ex_10():
    _, seleccion = await _zip([[_corrida()], _tiendas(), _lineas()])

    assert seleccion.omitidas == [
        {"sucursal_id": str(SUC_C), "nombre": "Pereira",
         "codigo": "BORRADOR", "motivo": "El pedido sigue en BORRADOR"},
        {"sucursal_id": str(SUC_D), "nombre": "Cali",
         "codigo": "SIN_PEDIDO", "motivo": "No tiene pedido (su cálculo "
         "no terminó bien)"},
        {"sucursal_id": str(SUC_E), "nombre": "Armenia",
         "codigo": "SIN_CANTIDAD",
         "motivo": "Todas sus cantidades son 0"}]


async def test_locks_and_the_lines_query_only_ask_for_what_is_exported():
    db, _ = await _zip([[_corrida()], _tiendas(), _lineas()])

    corrida, tiendas, lineas = (_sql(s) for s in db.executed_statements)

    assert "FOR SHARE" in corrida and "FOR UPDATE" not in corrida
    assert "FOR SHARE OF corrida_sucursal" in tiendas
    assert "ORDER BY corrida_sucursal.sucursal_id" in tiendas
    assert "corrida_sucursal.sucursal_id IN" not in tiendas
    assert f"'{SUC_A}'" in lineas and f"'{SUC_E}'" in lineas
    assert f"'{SUC_C}'" not in lineas and f"'{SUC_D}'" not in lineas
    assert "corrida_linea.motivo_exclusion IS NULL" in lineas
    assert "corrida_linea.pedido_final > 0" in lineas


async def test_a_listed_request_reads_only_the_listed_tiendas():
    db, seleccion = await _zip(
        [[_corrida()], [_tiendas()[0]], _lineas()[:2]], [SUC_A])

    tiendas = _sql(db.executed_statements[1])

    assert f"corrida_sucursal.sucursal_id IN ('{SUC_A}')" in tiendas
    assert [t.nombre for t in seleccion.tiendas] == ["Manizales"]
    assert seleccion.omitidas == []


async def test_a_listed_tienda_with_nothing_to_order_is_skipped_not_fatal():
    _, seleccion = await _zip(
        [[_corrida()], [_tiendas()[0], _tiendas()[4]], _lineas()[:2]],
        [SUC_A, SUC_E])

    assert [t.nombre for t in seleccion.tiendas] == ["Manizales"]
    assert [o["codigo"] for o in seleccion.omitidas] == ["SIN_CANTIDAD"]


async def test_a_listed_borrador_tienda_is_055_naming_it():
    db, error = await _zip_rechazado(
        [[_corrida()], [_tiendas()[0], _tiendas()[2]]],
        codigos.E_CORRIDA_EXPORTAR_NO_CERRADO, [SUC_A, SUC_C])

    assert "Pereira" in error.mensaje
    assert len(db.executed_statements) == 2


async def test_a_listed_tienda_without_pedido_is_065_naming_it():
    _, error = await _zip_rechazado(
        [[_corrida()], [_tiendas()[0], _tiendas()[3]]],
        codigos.E_CORRIDA_SIN_PEDIDO, [SUC_A, SUC_D])

    assert "Cali" in error.mensaje


async def test_a_listed_tienda_that_is_not_in_the_corrida_is_a_lookup_error():
    db = FakeAsyncSession(execute_queue=[[_corrida()], [_tiendas()[0]]])

    with pytest.raises(LookupError):
        await exportacion.preparar_corrida(db, CORRIDA, [SUC_A, SUC_B])

    assert len(db.executed_statements) == 2


async def test_a_repeated_listed_id_counts_once():
    _, seleccion = await _zip(
        [[_corrida()], [_tiendas()[0]], _lineas()[:2]], [SUC_A, SUC_A])

    assert [t.nombre for t in seleccion.tiendas] == ["Manizales"]


async def test_no_closed_or_sent_tienda_at_all_is_055_ex_06():
    db, error = await _zip_rechazado(
        [[_corrida()], [_tiendas()[2], _tiendas()[3]]],
        codigos.E_CORRIDA_EXPORTAR_NO_CERRADO)

    assert "ninguna" in error.mensaje.lower()
    assert len(db.executed_statements) == 2


async def test_closed_tiendas_that_all_have_nothing_to_order_are_056():
    db, error = await _zip_rechazado(
        [[_corrida()], [_tiendas()[0], _tiendas()[1]], []],
        codigos.E_CORRIDA_NADA_QUE_ENVIAR)

    assert "exportar" in error.mensaje
    assert len(db.executed_statements) == 3


async def test_a_tienda_with_rows_and_no_sic_is_057_naming_it():
    tiendas = [_fila(SUC_A, "Manizales", None), _fila(SUC_B, "Cali", "9")]

    _, error = await _zip_rechazado(
        [[_corrida()], tiendas, _lineas()], codigos.E_CORRIDA_EXPORTAR_SIN_SIC)

    assert "Manizales" in error.mensaje
    assert error.detalle["sucursal_id"] == str(SUC_A)


async def test_a_tienda_with_no_rows_and_no_sic_is_just_skipped():
    tiendas = [_fila(SUC_A, "Manizales", "1234"),
               _fila(SUC_E, "Armenia", None)]

    _, seleccion = await _zip(
        [[_corrida()], tiendas, _lineas()[:2]])

    assert [t.nombre for t in seleccion.tiendas] == ["Manizales"]
    assert seleccion.omitidas[0]["codigo"] == "SIN_CANTIDAD"


async def test_a_scenario_or_an_uncalculated_corrida_reads_no_lines():
    db, _ = await _zip_rechazado(
        [[_corrida(es_escenario=True)], _tiendas()],
        codigos.E_CORRIDA_ESCENARIO_NO_SE_CIERRA)
    assert len(db.executed_statements) == 2
    db, _ = await _zip_rechazado(
        [[_corrida(estado="FALLIDA")], _tiendas()],
        codigos.E_CORRIDA_ESTADO_NO_ADMITE)
    assert len(db.executed_statements) == 2
    db, _ = await _zip_rechazado(
        [[_corrida(invalidada=True)], _tiendas()],
        codigos.E_CORRIDA_INVALIDADA)
    assert len(db.executed_statements) == 2


async def test_an_unknown_corrida_is_a_lookup_error_for_the_zip():
    db = FakeAsyncSession(execute_queue=[[]])

    with pytest.raises(LookupError):
        await exportacion.preparar_corrida(db, CORRIDA, None)

    assert len(db.executed_statements) == 1


# --- Construir los archivos --------------------------------------------------


def _datos(nombre="Manizales", sic="1234", lineas=(("A-1", D("5.00")),)):
    return hmcl.DatosTienda(nombre, sic, fx.CORTE, list(lineas))


def test_the_xlsx_is_built_in_a_rewound_temp_file_with_its_size():
    archivo, tamano = exportacion.construir_xlsx(_datos())

    contenido = archivo.read()
    archivo.close()

    assert len(contenido) == tamano > 0
    hoja = load_workbook(io.BytesIO(contenido))["Pedido"]
    assert [c.value for c in hoja[6]] == ["A-1", 5]


def test_the_zip_names_each_file_with_the_f4_1_name():
    tiendas = [
        _datos("Manizales", "1234"),
        _datos("Medellín Poblado", "77", [("B-2", D("9.00"))])]

    archivo, tamano = exportacion.construir_zip(tiendas)

    contenido = archivo.read()
    archivo.close()
    paquete = zipfile.ZipFile(io.BytesIO(contenido))
    assert len(contenido) == tamano
    assert paquete.namelist() == [
        "Pedido_SIC1234_Manizales_2026-09-21.xlsx",
        "Pedido_SIC77_Medellin_Poblado_2026-09-21.xlsx"]
    hoja = load_workbook(io.BytesIO(paquete.read(paquete.namelist()[1])))
    assert hoja["Pedido"]["B2"].value == "77"
    assert hoja["Pedido"]["A6"].value == "B-2"


def test_the_spool_limit_is_eight_megabytes():
    assert exportacion.MAX_EN_MEMORIA == 8 * 1024 * 1024


def test_the_xlsx_pipe_matches_the_pure_builder_byte_for_content():
    datos = _datos(lineas=[("000123", D("5.00")), ("=1+1", D("2.00"))])

    archivo, _ = exportacion.construir_xlsx(datos)

    hoja = load_workbook(io.BytesIO(archivo.read()))["Pedido"]
    archivo.close()
    assert [(c.value, c.data_type) for c in hoja["A"][5:]] == [
        ("000123", "s"), ("=1+1", "s")]
    assert datetime.date(2026, 9, 21) == hoja["B3"].value.date()
