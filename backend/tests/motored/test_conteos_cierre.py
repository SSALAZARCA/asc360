"""
Inventory counts, close (odd/motored-conteos-inventario, WU10; design
§4.8, §5.1, §6.1, §7).

Pure rules first (the close guards, the result line of a code, the
accuracy KPI, SIN_COSTO out of the money totals, a surplus outside the
snapshot), then the adjustment Excel opened back with openpyxl, then the
HTTP layer with the services stubbed: scoping (another leader's conteo
is a 404), GERENCIA reads the result and the Excel and gets 403 on the
close. The SQL (materialization, the principal bodega, the forced close,
the double close, the public link dying) runs for real in
`pg_real/test_conteos_cierre_pg.py`.
"""
import datetime
import io
import uuid
from decimal import Decimal
from unittest.mock import AsyncMock

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import settings
from app.main import app
from app.motored.models.conteo import Conteo
from app.motored.services.auth import MotoredUser
from app.motored.services.conteos import (
    cierre, diferencias, errores, excel_ajustes,
)
from tests.motored.conftest import (
    FakeAsyncSession, override_motored_db, override_motored_user,
)

BASE = "/api/motored/conteos"
UTC = datetime.timezone.utc
AHORA = datetime.datetime(2026, 10, 9, 13, 0, tzinfo=UTC)
LIDER_ID = uuid.uuid4()
OTRO_LIDER_ID = uuid.uuid4()
CRITICO = Decimal("500000")
D = Decimal


@pytest.fixture(autouse=True)
def _listo(monkeypatch):
    monkeypatch.setattr(settings, "MOTORED_ENABLED", True)
    monkeypatch.setattr(settings, "MOTORED_SECRET_KEY", "conteos-cierre")
    monkeypatch.setattr(settings, "SECRET_KEY", "conteos-cierre-asc360")
    yield
    app.dependency_overrides.clear()


# --- guards ------------------------------------------------------------------


def _conteo(estado="EN_RECONTEO", lider_id=LIDER_ID):
    return Conteo(
        id=uuid.uuid4(), tipo="TOTAL", estado=estado, origen="MANUAL",
        sucursal_id=uuid.uuid4(), lider_id=lider_id,
        fecha_programada=AHORA.date(), created_at=AHORA,
        snapshot_tomado_en=AHORA, umbral_reconteo_pesos=D("100000"),
        umbral_critico_pesos=CRITICO)


@pytest.mark.parametrize("estado", [
    "PROGRAMADO", "EN_CONTEO", "CERRADO", "ANULADO"])
def test_only_the_reconteo_phase_can_close(estado):
    with pytest.raises(errores.EstadoInvalido):
        cierre.exigir_fase(_conteo(estado))


def test_the_reconteo_phase_can_close():
    cierre.exigir_fase(_conteo("EN_RECONTEO"))


def test_open_reconteos_refuse_the_close_with_their_counts():
    with pytest.raises(errores.ReconteosAbiertos) as error:
        cierre.decidir_cierre(
            {"PENDIENTE": 3, "ASIGNADO": 1}, forzar=False, motivo=None)

    assert error.value.datos == {"pendientes": 3, "asignados": 1}


def test_no_open_reconteo_closes_without_a_reason():
    assert cierre.decidir_cierre({}, forzar=False, motivo="x") is None


@pytest.mark.parametrize("motivo", [None, "", "   "])
def test_a_forced_close_needs_a_reason(motivo):
    with pytest.raises(errores.MotivoRequerido):
        cierre.decidir_cierre(
            {"PENDIENTE": 2}, forzar=True, motivo=motivo)


def test_a_forced_close_keeps_the_trimmed_reason():
    assert cierre.decidir_cierre(
        {"ASIGNADO": 2}, forzar=True, motivo="  Cierre de tienda ") == (
        "Cierre de tienda")


# --- one result line ---------------------------------------------------------


def _cruda(sistema=None, ronda1=None, costo=None, fuente="BODEGA",
           reconteo_estado=None, ronda2=None, codigo="R-1",
           referencia_id="nueva"):
    if referencia_id == "nueva":
        referencia_id = uuid.uuid4()
    return diferencias.FilaCruda(
        referencia_id=referencia_id, codigo=codigo, nombre="Pastilla",
        sistema=sistema, ronda1=ronda1, costo_unitario=costo,
        costo_fuente=fuente, ubicaciones=["Estante A3"],
        reconteo_id=None if reconteo_estado is None else uuid.uuid4(),
        reconteo_estado=reconteo_estado, reconteo_origen="UMBRAL",
        reconteo_sesion_id=None, misma_pareja_autorizada=False,
        ronda2=ronda2)


def test_a_finished_reconteo_with_the_same_sign_is_confirmed():
    linea = cierre.linea_de(_cruda(
        D("10"), D("1"), D("50000"), reconteo_estado="TERMINADO",
        ronda2=D("4")), CRITICO)

    assert (linea.sistema, linea.contado, linea.diferencia) == (
        D("10"), D("4"), D("-6"))
    assert linea.valor == D("-300000.00")
    assert linea.con_reconteo and linea.confirmada is True
    assert not linea.critico


def test_a_reconteo_that_flips_the_sign_is_not_confirmed():
    linea = cierre.linea_de(_cruda(
        D("10"), D("1"), D("50000"), reconteo_estado="TERMINADO",
        ronda2=D("12")), CRITICO)

    assert linea.diferencia == D("2") and linea.confirmada is False


def test_a_line_without_a_finished_reconteo_keeps_round_one():
    linea = cierre.linea_de(_cruda(
        D("10"), D("0"), D("60000"), reconteo_estado="ASIGNADO",
        ronda2=D("10")), CRITICO)

    assert (linea.contado, linea.diferencia) == (D("0"), D("-10"))
    assert not linea.con_reconteo and linea.confirmada is None
    assert linea.critico


def test_a_surplus_outside_the_snapshot_takes_the_fallback_cost():
    cruda = _cruda(None, D("2"), None, fuente=None)

    linea = cierre.linea_de(
        cierre.con_costo(cruda, D("7000"), "PRECIO"), CRITICO)

    assert (linea.sistema, linea.contado, linea.diferencia) == (
        D("0"), D("2"), D("2"))
    assert (linea.costo_unitario, linea.costo_fuente) == (
        D("7000"), "PRECIO")
    assert linea.valor == D("14000.00")


def test_a_surplus_without_any_cost_is_sin_costo():
    linea = cierre.linea_de(_cruda(None, D("2"), None, fuente=None), CRITICO)

    assert linea.costo_fuente == "SIN_COSTO"
    assert linea.costo_unitario is None and linea.valor is None


def test_an_unknown_code_keeps_no_referencia():
    linea = cierre.linea_de(_cruda(
        None, D("1"), None, fuente=None, codigo="ZZ-1",
        referencia_id=None), CRITICO)

    assert linea.referencia_id is None and linea.codigo == "ZZ-1"
    assert linea.costo_fuente == "SIN_COSTO"


# --- the accuracy KPI --------------------------------------------------------


def _linea(sistema, contado, costo, codigo, fuente="BODEGA", ref=True):
    cruda = _cruda(
        D(sistema), D(contado), None if costo is None else D(costo),
        fuente=fuente, codigo=codigo,
        referencia_id=uuid.uuid4() if ref else None)
    return cierre.linea_de(cruda, CRITICO)


def _fixture():
    """A -2 x 20000; B -6 x 50000; C +2 x 7000; D exact; E -5 x 20000;
    S sin costo +1; Z = 0 vs 0 (outside the universe)."""
    return [
        _linea("10", "8", "20000", "A"), _linea("10", "4", "50000", "B"),
        _linea("0", "2", "7000", "C", fuente="PRECIO"),
        _linea("3", "3", "1000", "D"), _linea("5", "0", "20000", "E"),
        _linea("0", "1", None, "S", fuente="SIN_COSTO", ref=False),
        _linea("0", "0", "900", "Z"),
    ]


def test_the_kpi_matches_hand_computed_values():
    kpi = cierre.calcular_kpi(_fixture())

    assert (kpi.refs_universo, kpi.refs_exactas) == (6, 1)
    assert kpi.exactitud_pct == D("16.67")
    assert kpi.valor_sistema == D("803000.00")
    assert kpi.valor_diferencia_neta == D("-426000.00")
    assert kpi.valor_diferencia_abs == D("454000.00")


def test_sin_costo_lines_stay_out_of_the_money_totals():
    kpi = cierre.calcular_kpi([
        _linea("5", "1", None, "S", fuente="SIN_COSTO"),
        _linea("1", "2", "100", "A")])

    assert kpi.valor_diferencia_neta == D("100.00")
    assert kpi.valor_diferencia_abs == D("100.00")
    assert kpi.valor_sistema == D("100.00")
    assert (kpi.refs_universo, kpi.refs_exactas) == (2, 0)


def test_negative_stock_adds_no_stock_value():
    kpi = cierre.calcular_kpi([_linea("-2", "0", "1000", "N")])

    assert kpi.valor_sistema == D("0.00")
    assert kpi.valor_diferencia_neta == D("2000.00")


def test_an_empty_count_has_no_accuracy():
    kpi = cierre.calcular_kpi([])

    assert (kpi.refs_universo, kpi.exactitud_pct) == (0, None)


def test_lines_sort_by_absolute_value_with_unvalued_last():
    orden = [f.codigo for f in cierre.ordenar(_fixture())]

    assert orden == ["B", "E", "A", "C", "D", "Z", "S"]


# --- the adjustment Excel ----------------------------------------------------


def _encabezado(**extra):
    datos = dict(
        tienda="Tienda Norte", codigo_co="E05", lider="Ana Líder",
        estado="CERRADO", fecha_programada=AHORA.date(),
        iniciado_en=AHORA, ronda_terminada_en=AHORA, cerrado_en=AHORA,
        fecha_corte=AHORA.date(), bodega="BA061",
        kpi=cierre.calcular_kpi(_fixture()),
        motivo_cierre_forzado="Se fue la luz")
    datos.update(extra)
    return cierre.Encabezado(**datos)


def _libro(lineas=None, **extra):
    lineas = cierre.ordenar(_fixture() if lineas is None else lineas)
    lineas = [f._replace(bodega="BA061") for f in lineas]
    datos = excel_ajustes.libro(_encabezado(**extra), lineas)
    return load_workbook(io.BytesIO(datos))


def test_the_excel_has_the_three_sheets_and_the_headers():
    libro = _libro()

    assert libro.sheetnames == ["Ajustes", "Resumen", "Sin costo"]
    cabecera = [c.value for c in libro["Ajustes"][1]]
    assert cabecera == [
        "Referencia", "Descripción", "Bodega", "Cantidad sistema",
        "Cantidad contada", "Diferencia", "Costo unitario",
        "Valor diferencia", "Ubicaciones", "Con reconteo", "Crítica"]


def test_the_adjustments_sheet_has_only_non_zero_rows_by_value():
    hoja = _libro()["Ajustes"]

    filas = [[c.value for c in fila] for fila in hoja.iter_rows(min_row=2)]
    assert [f[0] for f in filas] == ["B", "E", "A", "C"]
    b = filas[0]
    assert b[2:8] == ["BA061", 10, 4, -6, 50000, -300000]
    assert b[8:] == ["Estante A3", "No", "No"]
    assert hoja["H2"].number_format == "#,##0"
    assert hoja["A2"].data_type == "s"


def test_the_sin_costo_sheet_lists_unvalued_differences():
    hoja = _libro()["Sin costo"]

    filas = [[c.value for c in fila] for fila in hoja.iter_rows(min_row=2)]
    assert [f[0] for f in filas] == ["S"]
    assert "maestro" in filas[0][-1]


def test_no_sin_costo_sheet_when_every_difference_is_valued():
    libro = _libro([_linea("10", "8", "20000", "A")])

    assert libro.sheetnames == ["Ajustes", "Resumen"]


def test_the_summary_sheet_carries_the_kpi_and_the_forced_reason():
    hoja = _libro()["Resumen"]

    datos = {fila[0].value: fila[1].value for fila in hoja.iter_rows()}
    assert datos["Tienda"] == "Tienda Norte"
    assert datos["Líder"] == "Ana Líder"
    assert datos["Exactitud (%)"] == 16.67
    assert datos["Referencias evaluadas"] == 6
    assert datos["Diferencia neta"] == -426000
    assert datos["Motivo del cierre forzado"] == "Se fue la luz"


def test_the_file_name_uses_the_store_co_and_the_date():
    assert excel_ajustes.nombre_archivo(
        "ajustes", "E05", AHORA.date()) == (
        "ajustes_conteo_E05_2026-10-09.xlsx")


# --- HTTP: scoping and RBAC --------------------------------------------------


def _usuario(rol):
    ids = {"LIDER_INVENTARIOS": LIDER_ID}
    return MotoredUser(user_id=str(ids.get(rol, uuid.uuid4())), role=rol)


def _llamar(rol, metodo, ruta, conteo, json=None):
    override_motored_user(_usuario(rol))
    db = FakeAsyncSession(execute_queue=[[]], get_queue=[conteo])
    override_motored_db(db)
    return TestClient(app).request(metodo, BASE + ruta, json=json), db


def test_gerencia_gets_403_on_the_close():
    conteo = _conteo()

    r, _ = _llamar("GERENCIA", "POST", f"/{conteo.id}/cerrar", conteo,
                   json={})

    assert r.status_code == 403, r.text


@pytest.mark.parametrize("metodo, ruta", [
    ("POST", "/cerrar"), ("GET", "/resultado"), ("GET", "/ajustes.xlsx"),
    ("GET", "/avance.xlsx")])
def test_another_leaders_conteo_is_a_404(metodo, ruta):
    conteo = _conteo("CERRADO", lider_id=OTRO_LIDER_ID)

    r, _ = _llamar("LIDER_INVENTARIOS", metodo, f"/{conteo.id}{ruta}",
                   conteo, json={} if metodo == "POST" else None)

    assert r.status_code == 404, r.text
    assert r.json()["detail"]["code"] == "CONTEO_NO_ENCONTRADO"


def test_the_leader_closes_and_gets_the_kpi(monkeypatch):
    conteo = _conteo()
    kpi = cierre.calcular_kpi(_fixture())
    cerrar = AsyncMock(return_value=cierre.Cierre(conteo, kpi, 2))
    monkeypatch.setattr(cierre, "cerrar", cerrar)

    r, db = _llamar(
        "LIDER_INVENTARIOS", "POST", f"/{conteo.id}/cerrar", conteo,
        json={"forzar": True, "motivo": "Se fue la luz"})

    assert r.status_code == 200, r.text
    cuerpo = r.json()
    assert cuerpo["reconteos_cancelados"] == 2
    assert cuerpo["kpi"]["exactitud_pct"] == "16.67"
    assert cuerpo["kpi"]["refs_universo"] == 6
    kwargs = cerrar.await_args.kwargs
    assert (kwargs["forzar"], kwargs["motivo"]) == (True, "Se fue la luz")
    assert db.committed


@pytest.mark.parametrize("error, estado, codigo", [
    (errores.ReconteosAbiertos(pendientes=3, asignados=1), 409,
     "RECONTEOS_ABIERTOS"),
    (errores.SinBodegaPrincipal(), 409, "SIN_BODEGA_PRINCIPAL"),
    (errores.MotivoRequerido(), 422, "MOTIVO_REQUERIDO"),
    (errores.EstadoInvalido(), 409, "ESTADO_INVALIDO"),
])
def test_close_errors_map_to_http(monkeypatch, error, estado, codigo):
    conteo = _conteo()
    monkeypatch.setattr(cierre, "cerrar", AsyncMock(side_effect=error))

    r, db = _llamar("ADMIN", "POST", f"/{conteo.id}/cerrar", conteo,
                    json={"forzar": False})

    assert r.status_code == estado, r.text
    assert r.json()["detail"]["code"] == codigo
    assert not db.committed


def test_open_reconteos_answer_their_counts(monkeypatch):
    conteo = _conteo()
    monkeypatch.setattr(cierre, "cerrar", AsyncMock(
        side_effect=errores.ReconteosAbiertos(pendientes=3, asignados=1)))

    r, _ = _llamar("ADMIN", "POST", f"/{conteo.id}/cerrar", conteo)

    detalle = r.json()["detail"]
    assert (detalle["pendientes"], detalle["asignados"]) == (3, 1)


def _stub_resultado(monkeypatch):
    lineas = [f._replace(bodega="BA061") for f in cierre.ordenar(
        _fixture())]
    monkeypatch.setattr(cierre, "resultado", AsyncMock(return_value=lineas))
    monkeypatch.setattr(cierre, "encabezado", AsyncMock(
        return_value=_encabezado()))


@pytest.mark.parametrize("rol", ["ADMIN", "LIDER_INVENTARIOS", "GERENCIA"])
def test_readers_get_the_result_sorted_by_value(monkeypatch, rol):
    conteo = _conteo("CERRADO")
    _stub_resultado(monkeypatch)

    r, _ = _llamar(rol, "GET", f"/{conteo.id}/resultado", conteo)

    assert r.status_code == 200, r.text
    cuerpo = r.json()
    assert cuerpo["kpi"]["valor_diferencia_neta"] == "-426000.00"
    assert cuerpo["total"] == 7
    assert [i["codigo"] for i in cuerpo["items"]][:2] == ["B", "E"]
    assert cuerpo["items"][0]["bodega"] == "BA061"
    assert cuerpo["motivo_cierre_forzado"] == "Se fue la luz"


def test_the_result_of_an_open_conteo_is_a_409():
    conteo = _conteo("EN_RECONTEO")

    r, _ = _llamar("ADMIN", "GET", f"/{conteo.id}/resultado", conteo)

    assert r.status_code == 409
    assert r.json()["detail"]["code"] == "ESTADO_INVALIDO"


@pytest.mark.parametrize("rol", ["ADMIN", "LIDER_INVENTARIOS", "GERENCIA"])
def test_readers_download_the_adjustment_excel(monkeypatch, rol):
    conteo = _conteo("CERRADO")
    _stub_resultado(monkeypatch)

    r, _ = _llamar(rol, "GET", f"/{conteo.id}/ajustes.xlsx", conteo)

    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == excel_ajustes.XLSX
    assert "ajustes_conteo_E05_2026-10-09.xlsx" in (
        r.headers["content-disposition"])
    assert r.headers["cache-control"] == "no-store"
    libro = load_workbook(io.BytesIO(r.content))
    assert libro.sheetnames[0] == "Ajustes"


def test_the_progress_excel_is_only_for_an_open_conteo(monkeypatch):
    cerrado = _conteo("CERRADO")
    abierto = _conteo("EN_CONTEO")
    monkeypatch.setattr(cierre, "avance", AsyncMock(
        return_value=cierre.ordenar(_fixture())))
    monkeypatch.setattr(cierre, "encabezado", AsyncMock(
        return_value=_encabezado(estado="EN_CONTEO", cerrado_en=None)))

    r_cerrado, _ = _llamar(
        "ADMIN", "GET", f"/{cerrado.id}/avance.xlsx", cerrado)
    r_abierto, _ = _llamar(
        "GERENCIA", "GET", f"/{abierto.id}/avance.xlsx", abierto)

    assert r_cerrado.status_code == 409
    assert r_abierto.status_code == 200, r_abierto.text
    assert "avance_conteo_E05_" in r_abierto.headers["content-disposition"]
