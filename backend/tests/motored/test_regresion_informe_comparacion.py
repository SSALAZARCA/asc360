"""
Motored Pedidos F3 "Motor" (S8b-1, ADR-10, decisiones #11 y #16) — libro de
comparación/efecto para el dueño: Resumen, Antigüedad de datos, Filas,
A1, A2, B-entradas, B-salidas, el resumen de los switches y una hoja de delta
por cada switch medido.

Con libros sintéticos en `tmp_path`; los libros se leen de vuelta con openpyxl.
"""
from dataclasses import replace
from datetime import date
from decimal import Decimal
from pathlib import Path
from uuid import UUID

import openpyxl
import pytest

from app.motored.herramientas.regresion.antiguedad import sin_datos
from app.motored.herramientas.regresion.comparador import (
    atributos_de,
    entradas_de,
)
from app.motored.herramientas.regresion.delta import (
    InsumosMemoria,
    escenarios_estandar,
    medir_en_memoria,
)
from app.motored.herramientas.regresion.informe_comparacion import (
    HOJA_SWITCHES,
    ContenidoComparacion,
    escribir_libro_comparacion,
)
from app.motored.herramientas.regresion.extractor_excel import leer_libro
from app.motored.herramientas.regresion.informe_excel import HOJA_ANTIGUEDAD
from app.motored.herramientas.regresion.nivel_a import ejecutar_nivel_a
from app.motored.herramientas.regresion.nivel_b import ejecutar_nivel_b
from app.motored.herramientas.regresion.rutas import RutaInsegura
from app.motored.services.motor.sustitucion import resolver_cadenas
from app.motored.services.motor.tipos import NodoMaestro, ParametrosMotor
from tests.motored.fixtures.regresion.conjuntos import filas_base
from tests.motored.fixtures.regresion.entradas_app import entradas_app
from tests.motored.fixtures.regresion.libro_sintetico import escribir_libro
from tests.motored.fixtures.regresion.oraculo_excel import CLASES

CORTE = date(2026, 8, 15)
ID_VIEJA = UUID(int=21)
ID_NUEVA = UUID(int=22)


def _rellenar(filas: list) -> list:
    ancho = max((len(fila) for fila in filas), default=0)
    return [fila + (None,) * (ancho - len(fila)) for fila in filas]


def _leer(ruta: Path) -> dict:
    libro = openpyxl.load_workbook(ruta, read_only=True)
    hojas = {
        nombre: _rellenar(list(libro[nombre].iter_rows(values_only=True)))
        for nombre in libro.sheetnames
    }
    libro.close()
    return hojas


@pytest.fixture
def lectura(tmp_path):
    ruta = tmp_path / "libro.xlsx"
    escribir_libro(ruta, filas_base(), etiquetas=CLASES)
    return leer_libro(ruta)


def _insumos(lectura, *, consolidar=False):
    entradas = tuple(entradas_de(lectura))
    resoluciones = {}
    if consolidar:
        viejo, nuevo, *resto = entradas
        entradas = (
            replace(viejo, referencia_id=ID_VIEJA),
            replace(nuevo, referencia_id=ID_NUEVA), *resto,
        )
        resoluciones = resolver_cadenas(
            {ID_VIEJA: NodoMaestro(ID_VIEJA, False, ID_NUEVA)}
        )
    return InsumosMemoria(
        entradas=entradas, atributos=atributos_de(lectura),
        params=ParametrosMotor(), resoluciones=resoluciones,
    )


def _contenido(tmp_path, lectura, *, con_b=False, consolidar=False):
    libro = tmp_path / "libro.xlsx"
    deltas = medir_en_memoria(
        _insumos(lectura, consolidar=consolidar), escenarios_estandar()
    )
    nivel_b = None
    if con_b:
        app = entradas_app(
            lectura, cambios={"A-100": {"inventario": Decimal(11)}}
        )
        nivel_b = ejecutar_nivel_b(lectura, app, {"A-100": "T7"})
    return ContenidoComparacion(
        titulo="libro sintético",
        fecha_corte=CORTE,
        sucursal=lectura.sucursal.nombre,
        antiguedad=sin_datos("libro de Excel de referencia"),
        nivel_a=ejecutar_nivel_a(libro),
        nivel_b=nivel_b,
        deltas=deltas,
    )


def _escribir(tmp_path, lectura, **opciones):
    contenido = _contenido(tmp_path, lectura, **opciones)
    return _leer(escribir_libro_comparacion(contenido, tmp_path / "s.xlsx"))


class TestHojas:
    def test_nivel_a_y_switches_van_en_el_orden_del_informe(
        self, tmp_path, lectura
    ):
        hojas = list(_escribir(tmp_path, lectura))
        assert hojas[:5] == [
            "Resumen", HOJA_ANTIGUEDAD, "Filas", "A1", "A2",
        ]
        assert hojas[5] == HOJA_SWITCHES
        assert hojas[-1] == "Taxonomía"
        assert "B-entradas" not in hojas

    def test_el_nivel_b_pone_las_entradas_antes_que_las_salidas(
        self, tmp_path, lectura
    ):
        hojas = list(_escribir(tmp_path, lectura, con_b=True))
        assert hojas.index("B-entradas") + 1 == hojas.index("B-salidas")
        assert hojas.index("A2") < hojas.index("B-entradas")

    def test_hay_una_hoja_de_delta_por_switch_medido_y_ninguna_por_los_demas(
        self, tmp_path, lectura
    ):
        hojas = list(_escribir(tmp_path, lectura))
        deltas = [h for h in hojas if h.startswith("Delta ")]
        assert deltas == [
            "Delta dias_entre_pedidos=7", "Delta dias_entre_pedidos=15",
        ]


class TestResumen:
    def test_lista_cada_nivel_con_sus_conteos_y_resultado(
        self, tmp_path, lectura
    ):
        filas = _escribir(tmp_path, lectura, con_b=True)["Resumen"]
        niveles = {f[0]: f for f in filas if f[0] in (
            "A1", "A2", "B-entradas", "B-salidas"
        )}
        assert niveles["A1"][1:4] == (6, 6, 0)
        assert niveles["A1"][5] == "PASA"
        assert niveles["A2"][3] > 0 and "T1" in niveles["A2"][4]
        assert niveles["B-entradas"][3] == 1
        assert niveles["B-entradas"][5] == "PASA"
        assert niveles["B-salidas"][5] == "PASA"

    def test_encabeza_con_la_sucursal_la_fecha_y_la_fuente(
        self, tmp_path, lectura
    ):
        filas = _escribir(tmp_path, lectura)["Resumen"]
        assert filas[0][:2] == ("Sucursal", "SUCURSAL SINTETICA")
        assert filas[2][:2] == ("Fuente", "libro sintético")


class TestResultado:
    def test_un_nivel_con_sin_categoria_dice_falla(self, tmp_path):
        ruta = tmp_path / "libro.xlsx"
        escribir_libro(
            ruta, filas_base(), etiquetas=CLASES,
            sobrescribe={"A-100": {"AB": 1}},
        )
        lectura = leer_libro(ruta)
        filas = _escribir(tmp_path, lectura)["Resumen"]
        niveles = {f[0]: f for f in filas if f[0] in ("A1", "A2")}
        assert niveles["A1"][5] == "FALLA"
        assert niveles["A2"][5] == "FALLA"


class TestAntiguedad:
    def test_la_seccion_existe_y_explica_que_el_libro_no_trae_fechas(
        self, tmp_path, lectura
    ):
        filas = _escribir(tmp_path, lectura)[HOJA_ANTIGUEDAD]
        tipos = [f[0] for f in filas if f[0] in (
            "Inventario", "Backorder", "Facturas de pedidos",
            "Ingresos de facturas",
        )]
        assert len(tipos) == 4
        assert {f[4] for f in filas if f[0] == "Inventario"} == {
            "libro de Excel de referencia"
        }


class TestNiveles:
    def test_la_hoja_a2_trae_valor_excel_valor_app_diferencia_y_categoria(
        self, tmp_path, lectura
    ):
        filas = _escribir(tmp_path, lectura)["A2"]
        assert filas[0][:7] == (
            "Referencia / clase", "Fila Excel", "Columna", "Valor Excel",
            "Valor aplicación", "Diferencia", "Categoría",
        )
        c300 = [f for f in filas if f[0] == "C-300" and f[2] == "N"]
        assert len(c300) == 1
        assert c300[0][6] == "T1"
        assert c300[0][7].startswith("Divisor /18")
        assert isinstance(c300[0][3], (int, float))
        assert c300[0][5] < 0

    def test_a1_sin_diferencias_queda_solo_con_el_encabezado(
        self, tmp_path, lectura
    ):
        assert len(_escribir(tmp_path, lectura)["A1"]) == 1

    def test_filas_resume_por_referencia_las_diferencias_de_cada_nivel(
        self, tmp_path, lectura
    ):
        filas = _escribir(tmp_path, lectura, con_b=True)["Filas"]
        por_codigo = {f[1]: f for f in filas[1:]}
        assert filas[0][:6] == (
            "Fila Excel", "Referencia", "Dif. A1", "Dif. A2",
            "Dif. B-entradas", "Dif. B-salidas",
        )
        assert por_codigo["C-300"][2] == 0 and por_codigo["C-300"][3] > 0
        assert por_codigo["A-100"][4] == 1
        assert "T1" in por_codigo["C-300"][6]

    def test_b_entradas_lista_la_divergencia_con_su_categoria(
        self, tmp_path, lectura
    ):
        filas = _escribir(tmp_path, lectura, con_b=True)["B-entradas"]
        assert filas[1][0] == "A-100" and filas[1][2] == "V"
        assert filas[1][6] == "T7"


class TestSwitches:
    def test_el_resumen_distingue_medido_y_no_medido_con_su_motivo(
        self, tmp_path, lectura
    ):
        filas = _escribir(tmp_path, lectura)[HOJA_SWITCHES]
        por_nombre = {f[0]: f for f in filas[1:]}
        assert por_nombre["Línea base (todo apagado)"][1] == "BASE"
        dias = por_nombre["dias_entre_pedidos=7"]
        assert dias[1] == "MEDIDO"
        assert dias[4] < dias[3]
        assert dias[5] == dias[4] - dias[3]
        perdida = por_nombre["perdida x1"]
        assert perdida[1] == "NO MEDIDO"
        assert "demanda perdida" in perdida[-1]

    def test_la_hoja_de_delta_trae_totales_y_el_detalle_por_linea(
        self, tmp_path, lectura
    ):
        filas = _escribir(tmp_path, lectura)["Delta dias_entre_pedidos=7"]
        clave = {f[0]: f[1] for f in filas if f[0] and f[1] is not None}
        assert clave["Escenario"] == "dias_entre_pedidos=7"
        assert clave["Unidades base"] > clave["Unidades escenario"]
        assert clave["Líneas cambiadas"] > 0
        encabezado = next(f for f in filas if f[0] == "Referencia")
        assert encabezado[:4] == (
            "Referencia", "Tipo", "Clase base", "Clase escenario"
        )
        detalle = [f for f in filas if f[1] == "CAMBIO"]
        assert detalle and detalle[0][8] is not None

    def test_la_consolidacion_lista_las_transferidas(self, tmp_path, lectura):
        hojas = _escribir(tmp_path, lectura, consolidar=True)
        filas = hojas["Delta consolidar_sustituidas"]
        transferidas = [f for f in filas if f[1] == "TRANSFERIDA"]
        assert len(transferidas) == 1
        assert transferidas[0][0] == "A-100"
        assert transferidas[0][-1] == "Demanda transferida a B-200"


def test_rechaza_escribir_dentro_del_repositorio(tmp_path, lectura):
    contenido = _contenido(tmp_path, lectura)
    dentro = Path(__file__).parent / "comparacion_prohibida.xlsx"
    with pytest.raises(RutaInsegura):
        escribir_libro_comparacion(contenido, dentro)
    assert not dentro.exists()
