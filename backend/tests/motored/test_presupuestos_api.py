"""
Motored budgets (odd/motored-presupuestos-gerencia, T2): the HTTP layer.

The service is swapped for stubs (its behavior is covered against a real
Postgres in `pg_real/test_presupuestos_pg.py`); what is asserted here is RBAC
(ADMIN and GERENCIA only, through the REAL confinement in
`get_current_motored_user`), the upload guards, the template and the mapping of
service errors to HTTP statuses.
"""
import io
import uuid

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.config import settings
from app.main import app
from app.motored.auth import create_motored_token
from app.motored.deps import get_motored_user_lookup
from app.motored.services import presupuestos
from app.motored.services.auth import MotoredUser
from tests.motored.conftest import FakeAsyncSession, override_motored_db

RUTA = "/api/motored/presupuestos"
ID = uuid.uuid4()
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(autouse=True)
def _listo(monkeypatch):
    monkeypatch.setattr(settings, "MOTORED_ENABLED", True)
    monkeypatch.setattr(settings, "MOTORED_SECRET_KEY", "presupuestos-api-secret")
    monkeypatch.setattr(settings, "SECRET_KEY", "presupuestos-asc360-secret")
    yield
    app.dependency_overrides.clear()


def _detalle(mes, id_=str(ID)):
    return {
        "id": id_, "mes": mes, "version": 1, "origen": "EXCEL", "archivo_nombre": None, "nota": None,
        "created_at": "2031-03-01T00:00:00", "asesores": 0, "total": 0, "lineas": [], "por_tienda": [],
    }


@pytest.fixture(autouse=True)
def _servicio(monkeypatch):
    """Stubs: every read returns an empty-ish shape, writes echo a version."""
    llamadas = []

    async def validar(db, filas):
        llamadas.append(("validar", len(filas)))
        return {"valido": True, "filas": len(filas), "meses": [], "errores": [], "warnings": []}

    async def aplicar(db, filas, nombre, usuario_id):
        llamadas.append(("aplicar", nombre, usuario_id))
        return {"meses": [{"mes": "2031-03", "version": 1, "asesores": len(filas), "total": 1}]}

    async def editar(db, mes, cedula, sucursal_id, monto, nota, usuario_id):
        llamadas.append(("editar", mes, cedula, sucursal_id, monto, nota))
        return {"mes": "2031-03", "version": 2}

    async def quitar(db, mes, cedula, nota, usuario_id):
        llamadas.append(("quitar", mes, cedula, nota))
        return {"mes": "2031-03", "version": 3}

    async def meses(db):
        return []

    async def detalle_mes(db, mes):
        return _detalle(f"{mes:%Y-%m}")

    async def historial(db, mes):
        return []

    async def tiendas(db):
        return [{"id": str(ID), "nombre": "Cali"}]

    async def detalle_version(db, version_id):
        return _detalle("2031-03", str(version_id))

    for nombre, funcion in {
        "validar_archivo": validar, "aplicar_archivo": aplicar, "editar_asesor": editar,
        "quitar_asesor": quitar, "listar_meses": meses, "detalle_mes": detalle_mes,
        "historial_mes": historial, "listar_tiendas": tiendas, "detalle_version": detalle_version,
    }.items():
        monkeypatch.setattr(presupuestos, nombre, funcion)
    return llamadas


def _archivo(filas=(), encabezados=("Cédula", "Mes", "Tienda", "Presupuesto"), nombre="p.xlsx"):
    libro = openpyxl.Workbook()
    libro.active.append(list(encabezados))
    for fila in filas:
        libro.active.append(list(fila))
    buffer = io.BytesIO()
    libro.save(buffer)
    return {"file": (nombre, buffer.getvalue(), XLSX)}


def _llamar(rol, metodo, ruta, **kwargs):
    usuario = str(uuid.uuid4())

    async def _lookup(user_id: str):
        return MotoredUser(user_id=user_id, role=rol)

    app.dependency_overrides[get_motored_user_lookup] = lambda: _lookup
    override_motored_db(FakeAsyncSession(execute_queue=[[], [], []]))
    token = create_motored_token(sub=usuario, role=rol)
    with TestClient(app) as cliente:
        return cliente.request(metodo, RUTA + ruta, headers={"Authorization": f"Bearer {token}"}, **kwargs)


ENDPOINTS = [
    ("GET", "/plantilla.xlsx", {}),
    ("POST", "/validar", {"files": _archivo()}),
    ("POST", "/aplicar", {"files": _archivo([("1", "2031-03", "Cali", 10)])}),
    ("GET", "/meses", {}),
    ("GET", "/tiendas", {}),
    ("GET", "/meses/2031-03", {}),
    ("GET", "/meses/2031-03/versiones", {}),
    ("GET", f"/versiones/{ID}", {}),
    ("PUT", "/meses/2031-03/asesores/123", {"json": {"sucursal_id": str(ID), "monto": 5}}),
    ("DELETE", "/meses/2031-03/asesores/123", {}),
]


@pytest.mark.parametrize("rol", ["ADMIN", "GERENCIA"])
@pytest.mark.parametrize("metodo, ruta, kwargs", ENDPOINTS)
def test_admin_and_gerencia_reach_every_endpoint(rol, metodo, ruta, kwargs):
    respuesta = _llamar(rol, metodo, ruta, **kwargs)

    assert respuesta.status_code == 200, (rol, ruta, respuesta.text)


@pytest.mark.parametrize(
    "rol", ["COMPRAS", "CONSULTA", "SUCURSAL", "SERVICIO_CLIENTE"])
@pytest.mark.parametrize("metodo, ruta, kwargs", ENDPOINTS)
def test_every_other_role_is_forbidden(rol, metodo, ruta, kwargs):
    respuesta = _llamar(rol, metodo, ruta, **kwargs)

    assert respuesta.status_code == 403, (rol, ruta, respuesta.text)


@pytest.mark.parametrize("metodo, ruta, kwargs", ENDPOINTS)
def test_unauthenticated_requests_are_401(metodo, ruta, kwargs):
    override_motored_db(FakeAsyncSession(execute_queue=[[], [], []]))
    with TestClient(app) as cliente:
        respuesta = cliente.request(metodo, RUTA + ruta, **kwargs)

    assert respuesta.status_code == 401


def test_template_has_anio_mes_headers_and_example_rows():
    respuesta = _llamar("GERENCIA", "GET", "/plantilla.xlsx")

    assert respuesta.headers["content-type"].startswith(XLSX)
    assert respuesta.headers["content-disposition"] == 'attachment; filename="plantilla_presupuestos.xlsx"'
    libro = openpyxl.load_workbook(io.BytesIO(respuesta.content))
    filas = list(libro.active.iter_rows(values_only=True))
    assert filas[0] == ("Año", "Mes", "Cédula", "Tienda", "Presupuesto")
    assert len(filas) >= 2 and all(celda is not None for celda in filas[1])
    assert filas[1][:2] == (2026, 10)


def test_apply_passes_the_file_name_and_the_user_to_the_service(_servicio):
    respuesta = _llamar("GERENCIA", "POST", "/aplicar", files=_archivo([("1", "2031-03", "Cali", 10)], nombre="marzo.xlsx"))

    assert respuesta.json() == {"meses": [{"mes": "2031-03", "version": 1, "asesores": 1, "total": 1}]}
    [(_, nombre, usuario_id)] = _servicio
    assert nombre == "marzo.xlsx" and isinstance(usuario_id, uuid.UUID)


def test_a_file_that_is_not_xlsx_is_a_400():
    respuesta = _llamar("ADMIN", "POST", "/validar", files={"file": ("p.csv", b"a,b", "text/csv")})

    assert respuesta.status_code == 400


def test_a_file_missing_a_column_is_a_400():
    respuesta = _llamar("ADMIN", "POST", "/validar", files=_archivo(encabezados=("Cédula", "Mes", "Tienda")))

    assert respuesta.status_code == 400 and "Presupuesto" in respuesta.json()["detail"]


def test_a_file_over_the_row_limit_is_a_422(monkeypatch):
    monkeypatch.setattr(settings, "MOTORED_MAX_UPLOAD_ROWS", 2)

    respuesta = _llamar("ADMIN", "POST", "/validar", files=_archivo([("1", "2031-03", "Cali", 10)] * 3))

    assert respuesta.status_code == 422


def test_a_file_over_the_size_limit_is_a_413(monkeypatch):
    monkeypatch.setattr(settings, "MOTORED_MAX_UPLOAD_MB", 0)

    respuesta = _llamar("ADMIN", "POST", "/validar", files=_archivo([("1", "2031-03", "Cali", 10)]))

    assert respuesta.status_code == 413


def test_service_errors_map_to_http_statuses(monkeypatch):
    async def invalido(db, filas, nombre, usuario_id):
        raise presupuestos.PresupuestoInvalido([{"fila": 1, "columna": "Cédula", "mensaje": "mal"}])

    async def conflicto(db, mes):
        raise presupuestos.PresupuestoConflicto("otra persona")

    async def no_encontrado(db, mes):
        raise presupuestos.PresupuestoNoEncontrado("no hay")

    monkeypatch.setattr(presupuestos, "aplicar_archivo", invalido)
    monkeypatch.setattr(presupuestos, "detalle_mes", no_encontrado)
    assert _llamar("ADMIN", "POST", "/aplicar", files=_archivo([("1", "2031-03", "Cali", 10)])).status_code == 422
    assert _llamar("ADMIN", "GET", "/meses/2031-03").status_code == 404
    monkeypatch.setattr(presupuestos, "detalle_mes", conflicto)
    assert _llamar("ADMIN", "GET", "/meses/2031-03").status_code == 409


def test_the_422_of_a_rejected_apply_carries_the_row_errors(monkeypatch):
    async def invalido(db, filas, nombre, usuario_id):
        raise presupuestos.PresupuestoInvalido([{"fila": 1, "columna": "Cédula", "mensaje": "mal"}])

    monkeypatch.setattr(presupuestos, "aplicar_archivo", invalido)

    respuesta = _llamar("ADMIN", "POST", "/aplicar", files=_archivo([("1", "2031-03", "Cali", 10)]))

    assert respuesta.json()["detail"]["errores"] == [{"fila": 1, "columna": "Cédula", "mensaje": "mal"}]


@pytest.mark.parametrize("mes", ["2031-3", "2031-13", "marzo", "2031-03-01", "03-2031"])
def test_a_malformed_month_path_is_a_422(mes):
    assert _llamar("ADMIN", "GET", f"/meses/{mes}").status_code == 422


def test_manual_edit_validates_the_body_and_forwards_the_note(_servicio):
    malo = _llamar("ADMIN", "PUT", "/meses/2031-03/asesores/123", json={"sucursal_id": str(ID), "monto": 0})
    bueno = _llamar(
        "ADMIN", "PUT", "/meses/2031-03/asesores/123", json={"sucursal_id": str(ID), "monto": 5, "nota": "ajuste"})

    assert malo.status_code == 422 and bueno.json() == {"mes": "2031-03", "version": 2}
    assert _servicio[-1][-1] == "ajuste"


def test_manual_removal_forwards_the_optional_note(_servicio):
    respuesta = _llamar("ADMIN", "DELETE", "/meses/2031-03/asesores/123", params={"nota": "baja"})

    assert respuesta.json()["version"] == 3 and _servicio[-1][-1] == "baja"


def test_presupuesto_is_not_a_generic_bulk_upload_maestro():
    from app.motored.services.validators import _SCHEMA_BY_ENTIDAD

    assert "presupuesto" not in _SCHEMA_BY_ENTIDAD
    assert _llamar("ADMIN", "POST", "/../maestros/presupuesto/carga/validar", json={"filas": []}).status_code == 404


def test_tiendas_lists_active_sucursales_for_gerencia():
    respuesta = _llamar("GERENCIA", "GET", "/tiendas")

    assert respuesta.json() == [{"id": str(ID), "nombre": "Cali"}]


def test_a_long_file_name_is_truncated_to_the_column_keeping_the_extension(_servicio):
    nombre = "a" * 400 + ".xlsx"

    _llamar("ADMIN", "POST", "/aplicar", files=_archivo([("1", "2031-03", "Cali", 10)], nombre=nombre))

    guardado = _servicio[-1][1]
    assert len(guardado) == 255 and guardado.endswith(".xlsx")


def test_manual_edit_rejects_an_amount_that_would_overflow():
    cuerpo = {"sucursal_id": str(ID), "monto": 10 ** 11 + 1}

    respuesta = _llamar("ADMIN", "PUT", "/meses/2031-03/asesores/123", json=cuerpo)

    assert respuesta.status_code == 422


def test_manual_edit_accepts_the_maximum_amount():
    cuerpo = {"sucursal_id": str(ID), "monto": 10 ** 11}

    assert _llamar("ADMIN", "PUT", "/meses/2031-03/asesores/123", json=cuerpo).status_code == 200


def test_a_database_integrity_error_on_a_write_is_a_409(monkeypatch):
    from sqlalchemy.exc import IntegrityError

    async def carrera(db, *args):
        raise IntegrityError("insert", {}, Exception("duplicate key"))

    monkeypatch.setattr(presupuestos, "editar_asesor", carrera)

    respuesta = _llamar("ADMIN", "PUT", "/meses/2031-03/asesores/123", json={"sucursal_id": str(ID), "monto": 5})

    assert respuesta.status_code == 409 and "Intente de nuevo" in respuesta.json()["detail"]
