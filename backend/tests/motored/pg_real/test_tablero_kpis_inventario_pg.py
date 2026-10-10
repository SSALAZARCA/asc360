"""
KPI's, pestana Inventario, against a real Postgres (opt-in): the queries and the payload end to end.

World (year 2096 so it never crosses other data; reference month 2096-10, whose 3-month window Aug..Oct has 92
days). Stores: S1 (principal), S2, S3 (associated to S1), S4 (sells, no inventory). Cortes: Aug 31, Sep 30,
Oct 7 and an annulled one on Oct 20. Lines: R1/R2/R5 REPUESTOS, R3 ACCESORIOS, R4 no line, R6 LUBRICANTES,
R7 BATERIAS, R8 GPS.

Inventory at Oct 7 (value): S1+S3 R1 3000 (S3's 500 rolls up), S1 R2 200, R3 200, R4 10, R5 120 (precio_normal),
S2 R1 4000, R7 50, R8 200, R2 with 0 stock. Total 7780. Aug 31 = 2800, Sep 30 = 3900.
Last month with sales: S1 R1 2096-10, R2 2096-03, R4 2096-09, R5 2096-02, R3 never; S2 R1 2096-08,
R7 2095-06, R8 2096-05. Idle over 180 days at Oct 7: R2 200, R3 200 (never sold since 2095-06), R5 120, R7 50 = 570.
Demand Aug..Oct: S1 R1 34 (30 + 4 from S3), S1 R4 3, S2 R1 12, S1 R6 20 sold + 6 lost (BOT) = 26 with no
stock, S2 R2 7 sold + 3 lost = 10 with no stock. Transit: S1 R6 70 (60 + 10 from S3), two invoices worth 700.
Cost of sales Aug..Oct: S1 1200 (S3's 300 rolls up), S2 800, S4 1000 (no inventory: stays out).
"""
import datetime
import uuid
from decimal import Decimal

import pytest

from app.motored.models.carga_archivo import CargaArchivo
from app.motored.models.demanda_perdida import DemandaPerdida
from app.motored.models.factura_proveedor_linea import FacturaProveedorLinea
from app.motored.models.ingreso_factura import IngresoFactura
from app.motored.models.inventario_detalle import InventarioDetalle
from app.motored.models.proveedor import Proveedor
from app.motored.models.referencia import Referencia
from app.motored.models.sucursal import Sucursal
from app.motored.models.venta_detalle import VentaDetalle
from app.motored.models.venta_mensual import VentaMensual
from app.motored.services import tablero_asesores_consultas as q
from app.motored.services import tablero_kpis as k
from app.motored.services import tablero_kpis_inventario as inv
from tests.motored.pg_real.codigos_co import codigo_co_unico
from tests.motored.pg_real.test_tablero_asesores_pg import _carga, pytestmark, sesion  # noqa: F401

D = Decimal
F = datetime.date
C_AGO, C_SEP, C_OCT, C_ANUL = F(2096, 8, 31), F(2096, 9, 30), F(2096, 10, 7), F(2096, 10, 20)


class Mundo:
    pass


def _bot(estado):
    carga = _carga("DEMANDA_PERDIDA", estado)
    carga.origen = "BOT"
    return carga


async def _mundo(db, con_inventario=True):
    w, sfx = Mundo(), uuid.uuid4().hex[:6].upper()
    prov = Proveedor(id=uuid.uuid4(), codigo=f"P-{sfx}", nombre="P", es_principal=True,
                     dias_empaque_default=4, dias_transito_default=5, dias_seguridad_default=D("3"))
    s1, s2, s4 = (Sucursal(id=uuid.uuid4(), nombre=f"{n} {sfx}", sic=f"{n}-{sfx}", codigo_co=codigo_co_unico())
                  for n in ("S1", "S2", "S4"))
    s3 = Sucursal(id=uuid.uuid4(), nombre=f"S3 {sfx}", sic=f"S3-{sfx}", codigo_co=codigo_co_unico(), principal_id=s1.id)
    db.add_all([prov, s1, s2, s4])
    await db.flush()
    db.add(s3)
    lineas = {"R1": "REPUESTOS", "R2": "REPUESTOS", "R3": "ACCESORIOS", "R4": "NO APLICA", "R5": "REPUESTOS",
              "R6": "LUBRICANTES", "R7": "BATERIAS", "R8": "GPS"}
    refs = {c: Referencia(id=uuid.uuid4(), codigo=f"{c}-{sfx}", nombre=f"Nombre {c}", proveedor_id=prov.id,
                          unidad_empaque=1, precio_normal=D(20) if c == "R5" else None, linea_comercial=ln)
            for c, ln in lineas.items()}
    db.add_all(refs.values())
    c_venta, c_vm = _carga("VENTAS", "APLICADO"), _carga("VENTAS", "APLICADO")
    c_vm_anul, c_inv, c_inv_anul = _carga("VENTAS", "ANULADO"), _carga("INVENTARIO", "APLICADO"), _carga("INVENTARIO", "ANULADO")
    c_dp, c_dp_anul, c_bot_anul = _carga("DEMANDA_PERDIDA", "APLICADO"), _carga("DEMANDA_PERDIDA", "ANULADO"), _bot("ANULADO")
    c_fac, c_ing = _carga("FACTURAS_PROVEEDOR", "APLICADO"), _carga("INGRESOS", "APLICADO")
    db.add_all([c_venta, c_vm, c_vm_anul, c_inv, c_inv_anul, c_dp, c_dp_anul, c_bot_anul, c_fac, c_ing])
    await db.flush()
    w.s, w.refs, w.sfx = {"S1": s1, "S2": s2, "S3": s3, "S4": s4}, refs, sfx
    if con_inventario:
        _inventario(db, w, c_inv, c_inv_anul)
    _ventas(db, w, c_venta, c_vm, c_vm_anul)
    _perdidas(db, w, c_dp, c_dp_anul, c_bot_anul)
    _transito(db, w, c_fac, c_ing)
    await db.flush()
    return w


def _inventario(db, w, c_inv, c_anulada):
    def inv_(carga, corte, tienda, ref, existencia, costo, bodega="B1"):
        db.add(InventarioDetalle(
            id=uuid.uuid4(), carga_id=carga.id, fecha_corte=corte, sucursal_id=w.s[tienda].id,
            referencia_id=w.refs[ref].id, bodega=bodega, existencia=D(existencia),
            costo_unitario=None if costo is None else D(costo)))
    inv_(c_inv, C_OCT, "S1", "R1", 10, 100), inv_(c_inv, C_OCT, "S1", "R1", 5, 300, "B2")
    inv_(c_inv, C_OCT, "S3", "R1", 5, 100)
    inv_(c_inv, C_OCT, "S1", "R2", 4, 50), inv_(c_inv, C_OCT, "S1", "R3", 2, 100), inv_(c_inv, C_OCT, "S1", "R4", 1, 10)
    inv_(c_inv, C_OCT, "S1", "R5", 6, None)  # valued with precio_normal 20
    inv_(c_inv, C_OCT, "S2", "R1", 20, 200), inv_(c_inv, C_OCT, "S2", "R7", 1, 50), inv_(c_inv, C_OCT, "S2", "R8", 2, 100)
    inv_(c_inv, C_OCT, "S2", "R2", 0, 50)
    inv_(c_inv, C_AGO, "S1", "R1", 8, 100), inv_(c_inv, C_AGO, "S2", "R1", 10, 200)
    inv_(c_inv, C_SEP, "S1", "R1", 9, 100), inv_(c_inv, C_SEP, "S2", "R1", 15, 200)
    inv_(c_anulada, C_ANUL, "S1", "R1", 1000, 9)  # annulled load: no corte


def _mensual(db, w, carga, tienda, ref, anio, mes, unidades, origen="MOSTRADOR"):
    db.add(VentaMensual(
        id=uuid.uuid4(), sucursal_id=w.s[tienda].id, referencia_id=w.refs[ref].id, anio=anio, mes=mes,
        origen=origen, unidades=D(unidades), carga_id=carga.id))


def _linea_costo(db, w, carga, tienda, ref, anio, mes, costo, cliente="Taller X"):
    db.add(VentaDetalle(
        id=uuid.uuid4(), carga_id=carga.id, fecha=F(anio, mes, 10), anio=anio, mes=mes, sucursal_id=w.s[tienda].id,
        referencia_id=w.refs[ref].id, origen="MOSTRADOR", cantidad=D(1), vendedor="V", vendedor_norm="V",
        valor_bruto=D(1000), valor_descuentos=D(0), cliente_factura=cliente, nro_documento=uuid.uuid4().hex[:10],
        costo=D(costo)))


def _ventas(db, w, c_venta, c_vm, c_vm_anul):
    for fila in [("S1", "R1", 2096, 10, 30), ("S3", "R1", 2096, 9, 4), ("S1", "R2", 2096, 3, 5), ("S1", "R4", 2096, 9, 3),
                 ("S2", "R1", 2096, 8, 12), ("S1", "R5", 2096, 2, 2), ("S1", "R6", 2096, 9, 20), ("S2", "R2", 2096, 10, 7),
                 ("S2", "R4", 2096, 9, -5), ("S2", "R7", 2095, 6, 1), ("S2", "R8", 2096, 5, 3)]:
        _mensual(db, w, c_vm, *fila)
    _mensual(db, w, c_vm_anul, "S1", "R3", 2096, 10, 9)  # annulled load: R3 stays "never sold"
    # Cost of sales (real cost on the line); July is outside the reference window.
    for fila in [("S1", "R1", 2096, 7, 5000), ("S1", "R1", 2096, 8, 600), ("S1", "R3", 2096, 9, 200),
                 ("S1", "R4", 2096, 10, 100), ("S3", "R1", 2096, 9, 300), ("S2", "R1", 2096, 8, 400),
                 ("S4", "R1", 2096, 9, 1000)]:
        _linea_costo(db, w, c_venta, *fila)
    _linea_costo(db, w, c_venta, "S2", "R1", 2096, 10, 400, cliente="900723988")  # HMCL: the cost keeps it


def _perdidas(db, w, c_dp, c_dp_anul, c_bot_anul):
    def perdida(carga, origen, tienda, ref, fecha, cantidad):
        db.add(DemandaPerdida(
            id=uuid.uuid4(), fecha=fecha, sucursal_id=w.s[tienda].id, referencia_id=w.refs[ref].id,
            cantidad_solicitada=D(cantidad), origen=origen, carga_id=carga.id))
    perdida(c_bot_anul, "BOT", "S1", "R6", F(2096, 10, 3), 6)      # BOT always counts, even in an annulled load
    perdida(c_dp_anul, "EXCEL", "S1", "R6", F(2096, 10, 4), 99)    # EXCEL in an annulled load does not
    perdida(c_dp, "EXCEL", "S1", "R6", F(2096, 7, 15), 50)         # before the window
    perdida(c_dp, "EXCEL", "S2", "R2", F(2096, 9, 9), 3)


def _transito(db, w, c_fac, c_ing):
    def factura(numero, tienda, cantidad, valor, dia):
        db.add(FacturaProveedorLinea(
            id=uuid.uuid4(), prefijo_rh="AA", numero_rh=numero, fecha_factura=F(2096, 10, dia),
            sucursal_id=w.s[tienda].id, referencia_id=w.refs["R6"].id, cantidad=D(cantidad), valor_total=D(valor),
            carga_id=c_fac.id))
    factura(1001, "S1", 60, 600, 2), factura(1002, "S3", 10, 100, 3)
    db.add(IngresoFactura(  # an unrelated ingreso: it makes the invoices "verifiable" (pending)
        id=uuid.uuid4(), prefijo_rh="BB", numero_rh=7, fecha_ingreso=F(2096, 1, 1), valor_neto=D(1), carga_id=c_ing.id))


async def _filtro(db, w, meses=("2096-10",), modo="incluir", tiendas=("S1", "S2", "S4")):
    return await q.cargar_filtro(db, list(meses), modo, [w.s[t_].id for t_ in tiendas])


# --- the payload ----------------------------------------------------------------------------------


async def test_the_cards_have_value_days_rotation_previous_month_and_transit(sesion):
    w = await _mundo(sesion)

    r = await k.calcular_kpis_inventario(sesion, await _filtro(sesion, w))

    assert (r["corte"], r["costo_desde"], r["costo_hasta"]) == ("2096-10-07", "2096-08-01", "2096-10-31")
    tarjetas = r["tarjetas"]
    assert tarjetas["valor"] == 7780.0 and tarjetas["valor_mes_anterior"] == 3900.0
    assert tarjetas["dias"] == pytest.approx(7780 * 92 / 2000)  # S4's 1000 stays out: no inventory
    assert tarjetas["rotacion"] == pytest.approx(365 / tarjetas["dias"]) and tarjetas["dias_meta"] == 60
    assert tarjetas["transito"] == {"valor": 700.0, "facturas": 2}
    assert r["cortes_color"] == {"verde_hasta": 60, "ambar_hasta": 90}
    assert {"meses", "hmcl", "sucursales", "reglas", "usando_resumen", "datos_actualizados_en"} <= set(r)


async def test_the_trend_has_one_point_per_corte_month_with_the_days_of_its_own_window(sesion):
    w = await _mundo(sesion)

    r = await k.calcular_kpis_inventario(sesion, await _filtro(sesion, w))

    assert [(p["mes"], p["corte"], p["valor"]) for p in r["tendencia"]] == [
        ("2096-08", "2096-08-31", 2800.0), ("2096-09", "2096-09-30", 3900.0), ("2096-10", "2096-10-07", 7780.0)]
    assert r["tendencia"][0]["dias"] == pytest.approx(2800 * 92 / 6000)   # Jun..Aug: S1 5000+600, S2 400
    assert r["tendencia"][1]["dias"] == pytest.approx(3900 * 92 / 6500)   # Jul..Sep, S4 out
    assert r["tendencia"][2]["dias"] == pytest.approx(r["tarjetas"]["dias"])


def _valores_de_bandas(punto):
    return [b["valor"] for b in punto["bandas"]]


async def test_each_month_bar_is_split_into_green_amber_violet_by_the_days_of_each_pair(sesion):
    w = await _mundo(sesion)

    r = await k.calcular_kpis_inventario(sesion, await _filtro(sesion, w))

    # Aug (Jun..Aug): S1 R1 800 / (5600 cost) = 13 days green; S2 R1 2000 / 400 = 460 days violet.
    # Sep (Jul..Sep): S1 R1 900 / 5900 cost green; S2 R1 3000 / 400 violet.
    # Oct (Aug..Oct): only S1 R4 10 (cost 100: 9 days) is green; S1 R1 3000 / 900 and S2 R1 4000 / 800 are slow,
    # R3 lasts 92 days (> 90), and the pairs without cost of sales (R2, R5, R7, R8) are not moving: violet.
    assert [_valores_de_bandas(p) for p in r["tendencia"]] == [
        [800.0, 0.0, 2000.0], [900.0, 0.0, 3000.0], [10.0, 0.0, 7770.0]]
    assert [b["banda"] for b in r["tendencia"][0]["bandas"]] == ["verde", "ambar", "violeta"]
    assert r["tendencia"][2]["bandas"][0]["pct"] == pytest.approx(10 / 7780)
    for punto in r["tendencia"]:
        assert sum(_valores_de_bandas(punto)) == pytest.approx(punto["valor"])
        assert sum(b["pct"] for b in punto["bandas"]) == pytest.approx(1.0)


async def test_the_bands_respect_the_store_filter(sesion):
    w = await _mundo(sesion)

    solo_s2 = await k.calcular_kpis_inventario(sesion, await _filtro(sesion, w, tiendas=("S2",)))

    assert [_valores_de_bandas(p) for p in solo_s2["tendencia"]] == [
        [0.0, 0.0, 2000.0], [0.0, 0.0, 3000.0], [0.0, 0.0, 4250.0]]


async def test_the_age_bands_and_the_idle_value_follow_the_last_month_with_sales(sesion):
    w = await _mundo(sesion)

    r = await k.calcular_kpis_inventario(sesion, await _filtro(sesion, w))

    assert r["antiguedad"]["historial_desde"] == "2095-06"
    assert [b["valor"] for b in r["antiguedad"]["bandas"]] == [7010.0, 200.0, 320.0, 250.0]
    assert sum(b["pct"] for b in r["antiguedad"]["bandas"]) == pytest.approx(1.0)
    assert r["tarjetas"]["sin_movimiento"]["valor"] == 570.0
    assert r["tarjetas"]["sin_movimiento"]["pct"] == pytest.approx(570 / 7780)
    top = r["sin_movimiento_top"]
    assert [(x["referencia"][:2], x["valor"], x["existencia"]) for x in top] == [
        ("R3", 200.0, 2), ("R2", 200.0, 4), ("R5", 120.0, 6), ("R7", 50.0, 1)]  # a value tie: the older first
    assert top[0]["dias_sin_venta"] == (C_OCT - F(2095, 6, 1)).days  # never sold: from the history start
    assert top[1]["dias_sin_venta"] == (C_OCT - F(2096, 3, 31)).days and top[1]["nombre"] == "Nombre R2"
    assert top[0]["sucursal_id"] == str(w.s["S1"].id) and top[0]["tienda"].startswith("S1")


async def test_the_lines_carry_value_days_and_availability(sesion):
    w = await _mundo(sesion)

    r = await k.calcular_kpis_inventario(sesion, await _filtro(sesion, w))

    lineas = {x["linea"]: x for x in r["lineas"]}
    assert list(lineas) == ["REPUESTOS", "ACCESORIOS", "LLANTAS", "LUBRICANTES", "BATERIAS", "GPS", "CASCOS", "Sin línea"]
    assert lineas["REPUESTOS"]["valor"] == 7320.0 and lineas["REPUESTOS"]["pct"] == pytest.approx(7320 / 7780)
    assert lineas["REPUESTOS"]["dias"] == pytest.approx(7320 * 92 / 1700)
    assert lineas["REPUESTOS"]["disponibilidad_pct"] == pytest.approx(2 / 3)
    assert lineas["ACCESORIOS"]["dias"] == pytest.approx(200 * 92 / 200)
    assert lineas["LUBRICANTES"]["valor"] == 0 and lineas["LUBRICANTES"]["disponibilidad_pct"] == 0.0
    assert lineas["Sin línea"]["valor"] == 10.0 and lineas["Sin línea"]["dias"] == pytest.approx(10 * 92 / 100)
    assert lineas["LLANTAS"]["dias"] is None and lineas["LLANTAS"]["disponibilidad_pct"] is None


async def test_the_stores_roll_up_associates_and_are_ordered_by_days(sesion):
    w = await _mundo(sesion)

    r = await k.calcular_kpis_inventario(sesion, await _filtro(sesion, w))

    assert [x["sucursal_id"] for x in r["tiendas"]] == [str(w.s["S2"].id), str(w.s["S1"].id)]  # S4/S3 do not show
    s2, s1 = r["tiendas"]
    assert (s2["valor"], s1["valor"]) == (4250.0, 3530.0)
    assert s2["dias"] == pytest.approx(4250 * 92 / 800) and s1["dias"] == pytest.approx(3530 * 92 / 1200)
    assert (s1["sin_movimiento_valor"], s2["sin_movimiento_valor"]) == (520.0, 50.0)
    assert (s1["disponibilidad_pct"], s1["agotadas"]) == (pytest.approx(2 / 3), 1)
    assert (s2["disponibilidad_pct"], s2["agotadas"]) == (0.5, 1)


async def test_the_stockouts_with_demand_show_sold_lost_and_the_transit_that_covers_them(sesion):
    w = await _mundo(sesion)

    r = await k.calcular_kpis_inventario(sesion, await _filtro(sesion, w))

    assert r["tarjetas"]["disponibilidad"] == {"pct": pytest.approx(0.6), "agotadas": 2}
    ag = r["agotadas"]
    assert (ag["total"], ag["en_transito"], ag["sin_pedir"]) == (2, 1, 1)
    assert [(a["referencia"][:2], a["nombre"], a["vendidas"], a["perdidas"], a["demanda"], a["transito"]) for a in ag["items"]] == [
        ("R6", "Nombre R6", 20, 6, 26, 70), ("R2", "Nombre R2", 7, 3, 10, 0)]
    assert ag["items"][0]["sucursal_id"] == str(w.s["S1"].id)
    assert [a["cobertura"] for a in ag["items"]] == [1.0, 0.0]


async def test_a_store_filter_limits_everything_to_that_store(sesion):
    w = await _mundo(sesion)

    r = await k.calcular_kpis_inventario(sesion, await _filtro(sesion, w, tiendas=("S2",)))

    assert r["tarjetas"]["valor"] == 4250.0 and r["tarjetas"]["valor_mes_anterior"] == 3000.0
    assert r["tarjetas"]["dias"] == pytest.approx(4250 * 92 / 800)
    assert r["tarjetas"]["transito"] == {"valor": 0.0, "facturas": 0}
    assert [x["sucursal_id"] for x in r["tiendas"]] == [str(w.s["S2"].id)]
    assert (r["agotadas"]["total"], r["agotadas"]["en_transito"]) == (1, 0)


async def test_selecting_a_principal_includes_its_associated_stores(sesion):
    w = await _mundo(sesion)

    r = await k.calcular_kpis_inventario(sesion, await _filtro(sesion, w, tiendas=("S1",)))

    assert r["tarjetas"]["valor"] == 3530.0 and r["tarjetas"]["transito"] == {"valor": 700.0, "facturas": 2}
    assert r["agotadas"]["items"][0]["transito"] == 70


async def test_an_earlier_reference_month_uses_the_corte_of_that_month(sesion):
    w = await _mundo(sesion)

    r = await k.calcular_kpis_inventario(sesion, await _filtro(sesion, w, meses=("2096-09",)))

    assert r["corte"] == "2096-09-30" and r["tarjetas"]["valor"] == 3900.0
    assert r["tarjetas"]["valor_mes_anterior"] == 2800.0
    assert [p["mes"] for p in r["tendencia"]] == ["2096-08", "2096-09"]
    assert (r["costo_desde"], r["costo_hasta"]) == ("2096-07-01", "2096-09-30")


async def test_a_month_before_every_corte_falls_back_to_the_latest_one(sesion):
    w = await _mundo(sesion)

    r = await k.calcular_kpis_inventario(sesion, await _filtro(sesion, w, meses=("2096-01",)))

    assert r["corte"] == "2096-10-07" and r["tarjetas"]["valor"] == 7780.0
    assert r["tendencia"] == [] and r["tarjetas"]["valor_mes_anterior"] is None


async def test_the_hmcl_mode_changes_neither_the_inventory_nor_the_cost_of_sales(sesion):
    w = await _mundo(sesion)

    base = await k.calcular_kpis_inventario(sesion, await _filtro(sesion, w))
    for modo in ("excluir", "solo"):
        otro = await k.calcular_kpis_inventario(sesion, await _filtro(sesion, w, modo=modo))
        assert otro["tarjetas"]["valor"] == base["tarjetas"]["valor"]
        assert otro["tarjetas"]["dias"] == pytest.approx(base["tarjetas"]["dias"])


async def test_the_configured_threshold_and_cuts_are_read_from_configuracion(sesion):
    from app.motored.models.parametro_metodologia import ParametroMetodologia
    w = await _mundo(sesion)
    for clave, valor in (("kpi_inventario_sin_movimiento_dias", 100), ("kpi_inventario_dias_meta", 45),
                         ("kpi_inventario_dias_cortes", {"verde_hasta": 30, "ambar_hasta": 50})):
        sesion.add(ParametroMetodologia(id=uuid.uuid4(), clave=clave, valor=valor, vigente_desde=F(2096, 10, 1)))
    await sesion.flush()

    r = await k.calcular_kpis_inventario(sesion, await _filtro(sesion, w))

    assert r["tarjetas"]["dias_meta"] == 45 and r["cortes_color"] == {"verde_hasta": 30, "ambar_hasta": 50}
    assert r["tarjetas"]["sin_movimiento"]["valor"] == 770.0  # the 570 plus S2 R8 (200, idle 129 days)


async def test_without_any_inventory_the_payload_is_empty_but_complete(sesion):
    w = await _mundo(sesion, con_inventario=False)

    r = await k.calcular_kpis_inventario(sesion, await _filtro(sesion, w))

    assert r["corte"] is None and r["tarjetas"]["valor"] == 0 and r["tarjetas"]["dias"] is None
    assert r["tendencia"] == [] and r["tiendas"] == [] and r["sin_movimiento_top"] == []
    assert r["agotadas"] == {"total": 0, "en_transito": 0, "sin_pedir": 0, "items": []}
    assert r["tarjetas"]["transito"] == {"valor": 700.0, "facturas": 2}


async def test_the_complete_result_for_the_excel_lists_every_idle_pair_and_every_stockout(sesion):
    w = await _mundo(sesion)

    resultado = await inv.calcular_inventario(sesion, await _filtro(sesion, w), completo=True)

    assert [x["referencia"][:2] for x in resultado.sin_movimiento] == ["R3", "R2", "R5", "R7"]
    assert [x["referencia"][:2] for x in resultado.agotadas] == ["R6", "R2"]
    assert all("referencia_id" not in x for x in resultado.sin_movimiento + resultado.agotadas)


async def test_one_request_runs_a_bounded_number_of_queries_whatever_the_number_of_stores(sesion):
    from sqlalchemy import event
    w = await _mundo(sesion)
    filtro = await _filtro(sesion, w)
    conteo = []
    motor = sesion.bind.sync_engine
    escuchar = lambda *args: conteo.append(args[2])  # noqa: E731
    event.listen(motor, "before_cursor_execute", escuchar)
    try:
        await k.calcular_kpis_inventario(sesion, filtro)
    finally:
        event.remove(motor, "before_cursor_execute", escuchar)

    assert len(conteo) <= 24, f"{len(conteo)} queries"


# --- end to end through the app -------------------------------------------------------------------


@pytest.fixture
async def http(sesion, monkeypatch):
    import httpx
    from app.config import settings
    from app.main import app
    from app.motored.deps import get_current_motored_user, get_motored_db
    from app.motored.services.auth import MotoredUser

    monkeypatch.setattr(settings, "MOTORED_ENABLED", True)
    monkeypatch.setattr(settings, "MOTORED_SECRET_KEY", "kpis-inv-api-pg")

    async def db():
        yield sesion

    async def usuario():
        return MotoredUser(user_id="00000000-0000-0000-0000-000000000001", role="GERENCIA")

    app.dependency_overrides[get_motored_db] = db
    app.dependency_overrides[get_current_motored_user] = usuario
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://motored") as cliente:
        yield cliente
    app.dependency_overrides.clear()


BASE = "/api/motored/tablero-asesores/kpis/inventario"


async def test_the_inventario_endpoint_returns_the_json_of_the_tab(http, sesion):
    w = await _mundo(sesion)
    tiendas = ",".join(str(w.s[n].id) for n in ("S1", "S2", "S4"))

    r = await http.get(BASE, params={"meses": "2096-10", "sucursales": tiendas})

    assert r.status_code == 200, r.text
    cuerpo = r.json()
    assert cuerpo["corte"] == "2096-10-07" and cuerpo["tarjetas"]["valor"] == 7780.0
    assert isinstance(cuerpo["tarjetas"]["dias"], float) and len(cuerpo["tendencia"]) == 3
    assert cuerpo["agotadas"]["items"][0]["referencia"].startswith("R6")


async def test_the_excel_endpoint_lists_all_idle_pairs_not_only_the_top(http, sesion):
    import io
    from openpyxl import load_workbook
    w = await _mundo(sesion)
    tiendas = ",".join(str(w.s[n].id) for n in ("S1", "S2", "S4"))

    r = await http.get(f"{BASE}/excel", params={"meses": "2096-10", "sucursales": tiendas})

    assert r.status_code == 200, r.text
    assert 'filename="inventario_2096-10-07.xlsx"' in r.headers["content-disposition"]
    libro = load_workbook(io.BytesIO(r.content))
    filas = [[c.value for c in f] for f in libro["Sin movimiento"].iter_rows()]
    referencias = [f[0][:2] for f in filas if f[0] and f[0][:1] == "R" and f[0][1:2].isdigit()]
    assert referencias == ["R3", "R2", "R5", "R7"]
    nombres = {f[0] for f in [[c.value for c in f] for f in libro["Tiendas"].iter_rows()]}
    assert {"Tiendas", "Tienda"} <= nombres


async def test_without_a_corte_the_configured_settings_are_still_read(sesion):
    from app.motored.models.parametro_metodologia import ParametroMetodologia
    from app.motored.services import tablero_inventario_excel as excel
    w = await _mundo(sesion, con_inventario=False)
    for clave, valor in (("kpi_inventario_sin_movimiento_dias", 100), ("kpi_inventario_dias_meta", 45),
                         ("kpi_inventario_dias_cortes", {"verde_hasta": 30, "ambar_hasta": 50})):
        sesion.add(ParametroMetodologia(id=uuid.uuid4(), clave=clave, valor=valor, vigente_desde=F(2096, 10, 1)))
    await sesion.flush()

    datos, sin_mov, agotadas = await k.calcular_kpis_inventario_completo(sesion, await _filtro(sesion, w))

    assert datos["corte"] is None and datos["tarjetas"]["dias_meta"] == 45
    assert datos["cortes_color"] == {"verde_hasta": 30, "ambar_hasta": 50}
    assert datos["sin_movimiento_umbral_dias"] == 100
    import io
    from openpyxl import load_workbook
    libro = load_workbook(io.BytesIO(excel.construir_libro(datos, sin_mov, agotadas, [])))
    textos = [str(c.value) for f in libro["Sin movimiento"].iter_rows() for c in f if c.value is not None]
    assert "más de 100 días" in textos


async def test_unknown_or_null_reference_ids_do_not_break_the_names_lookup(sesion):
    from app.motored.services import tablero_kpis_inventario_consultas as qi
    w = await _mundo(sesion)

    nombres = await qi.consultar_referencias(sesion, [None, "no-es-uuid", w.refs["R1"].id, str(w.refs["R2"].id)])

    assert set(nombres) == {w.refs["R1"].id, w.refs["R2"].id}
    filas = await inv._rotular(sesion, [{"referencia_id": None, "valor": 1.0}])
    assert filas == [{"referencia": "None", "nombre": None, "valor": 1.0}]
