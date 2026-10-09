"""
Inventory counts, close against a real Postgres (opt-in,
`MOTORED_TEST_PG_URL`, database migrated to head;
odd/motored-conteos-inventario, WU10; design §4.8, §5.1, §7).

Same savepoint-mode harness and scene as `test_conteos_reconteos_pg.py`:
five referencias (A, B, D, E in the snapshot; C a surplus outside it)
plus an unknown code, counted in two locations. Proves the result lines
and the KPI against hand-computed values, the principal-bodega
attribution on a two-bodega store, the forced close keeping round 1,
the double close, the missing principal bodega and the public link
dying at close.
"""
import io
import uuid
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.motored.models.bodega import Bodega
from app.motored.models.conteo import Conteo
from app.motored.models.conteo_reconteo import ConteoReconteo
from app.motored.models.conteo_resultado import ConteoResultado
from app.motored.models.conteo_sesion import ConteoSesion
from app.motored.models.referencia import Referencia
from app.motored.models.sucursal import Sucursal
from tests.motored.pg_real import test_conteos_reconteos_pg as rec
from tests.motored.pg_real import test_publico_conteos_pg as publico
from tests.motored.pg_real.test_conteos_api_pg import (  # noqa: F401
    pytestmark,
)

D = Decimal

# The shared harness and the scene, as fixtures here.
fabrica = rec.fabrica
_app_lista = rec._app_lista
abierto = rec.abierto
escena = rec.escena


async def _con_bodegas(fabrica, mundo):
    """The store gets a principal and a secondary bodega record."""
    sufijo = uuid.uuid4().hex[:6].upper()
    principal, secundaria = f"P{sufijo}", f"S{sufijo}"
    async with fabrica() as db:
        sucursal = await db.get(Sucursal, mundo.conteo.sucursal_id)
        sucursal.bodega_principal = principal
        filas = [
            Bodega(id=uuid.uuid4(), codigo=principal,
                   sucursal_id=sucursal.id),
            Bodega(id=uuid.uuid4(), codigo=secundaria,
                   sucursal_id=sucursal.id, bodega_principal=principal)]
        db.add_all(filas)
        await db.commit()
    mundo.principal, mundo.secundaria = filas
    return mundo


async def _precio(fabrica, codigo, precio):
    async with fabrica() as db:
        referencia = await db.scalar(
            select(Referencia).where(Referencia.codigo == codigo))
        referencia.precio_normal = D(precio)
        await db.commit()


async def _cerrar(mundo, **cuerpo):
    return await rec._lider(mundo, "POST", "/cerrar", json=cuerpo)


async def _resultado(fabrica, mundo):
    async with fabrica() as db:
        filas = (await db.scalars(select(ConteoResultado).where(
            ConteoResultado.conteo_id == mundo.conteo.id))).all()
    return {f.codigo: f for f in filas}


async def _cerrar_reconteos(fabrica, mundo):
    """Pair TRES recounts B = 4 (B1) and finds no E; the leader cancels
    C and the unknown code (they keep round 1)."""
    cod = mundo.codigos
    pendientes = await rec._reconteos(fabrica, mundo)
    tres = await rec._unirse(mundo, rec.TRES)
    for clave in ("B", "E"):
        r = await rec._asignar(mundo, pendientes[cod[clave]], tres)
        assert r.status_code == 200, r.text
    await rec._ubicar(mundo, tres, "UBI-B1")
    b = pendientes[cod["B"]]
    await rec._leer(mundo, tres, rec._item(cod["B"], "4", reconteo_id=b.id))
    for clave in ("B", "E"):
        r = await rec._pedir(
            mundo, tres, "POST",
            f"/reconteos/{pendientes[cod[clave]].id}/terminar")
        assert r.status_code == 200, r.text
    for codigo in (cod["C"], mundo.desconocido):
        r = await rec._lider(
            mundo, "POST", f"/reconteos/{pendientes[codigo].id}/cancelar")
        assert r.status_code == 200, r.text
    return tres


async def test_an_end_to_end_close_matches_hand_computed_values(
        escena, fabrica):
    await _con_bodegas(fabrica, escena)
    await _precio(fabrica, escena.codigos["C"], "7000")
    uno = await rec._ronda_uno(escena)
    await rec._terminar_ronda(escena)
    tres = await _cerrar_reconteos(fabrica, escena)
    cod = escena.codigos

    r = await _cerrar(escena)

    assert r.status_code == 200, r.text
    kpi = r.json()["kpi"]
    assert (kpi["refs_universo"], kpi["refs_exactas"]) == (6, 1)
    assert kpi["exactitud_pct"] == "16.67"
    assert (kpi["valor_sistema"], kpi["valor_diferencia_neta"],
            kpi["valor_diferencia_abs"]) == (
        "803000.00", "-426000.00", "454000.00")
    filas = await _resultado(fabrica, escena)
    assert set(filas) == {*cod.values(), escena.desconocido}
    esperado = {
        "A": ("10", "8", "-2", "-40000", False, None),
        "B": ("10", "4", "-6", "-300000", True, True),
        "C": ("0", "2", "2", "14000", False, None),
        "D": ("3", "3", "0", "0", False, None),
        "E": ("5", "0", "-5", "-100000", True, True),
    }
    for clave, (sis, cont, dif, valor, con_rec, conf) in esperado.items():
        f = filas[cod[clave]]
        assert (f.existencia_sistema, f.cantidad_contada, f.diferencia,
                f.valor_diferencia, f.con_reconteo, f.confirmada) == (
            D(sis), D(cont), D(dif), D(valor), con_rec, conf), clave
    assert filas[cod["C"]].costo_fuente == "PRECIO"
    assert filas[cod["A"]].ubicaciones == ["Estante A3", "Estante B1"]
    raro = filas[escena.desconocido]
    assert (raro.referencia_id, raro.costo_fuente,
            raro.valor_diferencia) == (None, "SIN_COSTO", None)
    assert filas[cod["E"]].critico is False
    await _link_is_dead(escena, fabrica, uno, tres)


async def _link_is_dead(mundo, fabrica, *parejas):
    async with fabrica() as db:
        conteo = await db.get(Conteo, mundo.conteo.id)
        estados = set((await db.scalars(select(ConteoSesion.estado).where(
            ConteoSesion.conteo_id == mundo.conteo.id))).all())
    assert (conteo.estado, conteo.codigo_hash) == ("CERRADO", None)
    assert conteo.cerrado_por == mundo.lider.id
    assert estados == {"CERRADA"}
    unirse = await publico._unirse(mundo)
    assert unirse.status_code == 401, unirse.text
    for pareja in parejas:
        r = await rec._pedir(mundo, pareja, "GET", "/sesion")
        assert r.status_code == 401, r.text


async def test_everything_goes_to_the_principal_bodega(escena, fabrica):
    await _con_bodegas(fabrica, escena)
    await rec._ronda_uno(escena)
    await rec._terminar_ronda(escena)

    r = await _cerrar(escena, forzar=True, motivo="Cierre de prueba")
    ajustes = await rec._lider(escena, "GET", "/ajustes.xlsx")

    assert r.status_code == 200, r.text
    filas = await _resultado(fabrica, escena)
    assert {f.bodega_ajuste_id for f in filas.values()} == {
        escena.principal.id}
    assert ajustes.status_code == 200, ajustes.text
    libro = load_workbook(io.BytesIO(ajustes.content))
    bodegas = {c.value for c in libro["Ajustes"]["C"][1:]}
    assert bodegas == {escena.principal.codigo}
    resultado = await rec._lider(escena, "GET", "/resultado")
    assert resultado.json()["bodega"] == escena.principal.codigo


async def test_a_forced_close_cancels_open_reconteos_and_keeps_round_one(
        escena, fabrica):
    await _con_bodegas(fabrica, escena)
    await rec._ronda_uno(escena)
    await rec._terminar_ronda(escena)
    cod = escena.codigos
    pendientes = await rec._reconteos(fabrica, escena)
    tres = await rec._unirse(escena, rec.TRES)
    await rec._asignar(escena, pendientes[cod["B"]], tres)
    await rec._ubicar(escena, tres, "UBI-B1")
    await rec._leer(escena, tres, rec._item(
        cod["B"], "9", reconteo_id=pendientes[cod["B"]].id))

    normal = await _cerrar(escena)
    sin_motivo = await _cerrar(escena, forzar=True)
    forzado = await _cerrar(escena, forzar=True, motivo=" Se fue la luz ")

    assert normal.status_code == 409, normal.text
    detalle = normal.json()["detail"]
    assert (detalle["code"], detalle["pendientes"],
            detalle["asignados"]) == ("RECONTEOS_ABIERTOS", 3, 1)
    assert sin_motivo.status_code == 422, sin_motivo.text
    assert forzado.status_code == 200, forzado.text
    assert forzado.json()["reconteos_cancelados"] == 4
    assert forzado.json()["motivo_cierre_forzado"] == "Se fue la luz"
    despues = await rec._reconteos(fabrica, escena)
    assert {r.estado for r in despues.values()} == {"CANCELADO"}
    b = (await _resultado(fabrica, escena))[cod["B"]]
    assert (b.cantidad_contada, b.con_reconteo, b.confirmada) == (
        D("1"), False, None)


async def test_a_double_close_is_a_409(escena, fabrica):
    await _con_bodegas(fabrica, escena)
    await rec._ronda_uno(escena)
    await rec._terminar_ronda(escena)

    primero = await _cerrar(escena, forzar=True, motivo="Fin")
    segundo = await _cerrar(escena, forzar=True, motivo="Fin")

    assert primero.status_code == 200, primero.text
    assert segundo.status_code == 409
    assert segundo.json()["detail"]["code"] == "ESTADO_INVALIDO"
    assert len(await _resultado(fabrica, escena)) == 6


async def test_closing_during_round_one_is_a_409(escena, fabrica):
    await _con_bodegas(fabrica, escena)

    r = await _cerrar(escena)

    assert r.status_code == 409
    assert r.json()["detail"]["code"] == "ESTADO_INVALIDO"


async def test_no_principal_bodega_refuses_the_close(escena, fabrica):
    await rec._ronda_uno(escena)
    await rec._terminar_ronda(escena)

    r = await _cerrar(escena, forzar=True, motivo="Fin")

    assert r.status_code == 409, r.text
    assert r.json()["detail"]["code"] == "SIN_BODEGA_PRINCIPAL"
    async with fabrica() as db:
        conteo = await db.get(Conteo, escena.conteo.id)
        abiertos = (await db.scalars(select(ConteoReconteo.estado).where(
            ConteoReconteo.conteo_id == escena.conteo.id))).all()
    assert conteo.estado == "EN_RECONTEO"
    assert set(abiertos) == {"PENDIENTE"}
    assert await _resultado(fabrica, escena) == {}


async def test_the_progress_excel_reads_the_live_differences(
        escena, fabrica):
    await _con_bodegas(fabrica, escena)
    await rec._ronda_uno(escena)

    r = await rec._lider(escena, "GET", "/avance.xlsx")

    assert r.status_code == 200, r.text
    libro = load_workbook(io.BytesIO(r.content))
    codigos = [c.value for c in libro["Ajustes"]["A"][1:]]
    cod = escena.codigos
    assert codigos == [cod["B"], cod["E"], cod["A"], cod["C"]]
    assert [c.value for c in libro["Sin costo"]["A"][1:]] == [
        cod["C"], escena.desconocido]


@pytest.mark.parametrize("rol", ["GERENCIA"])
async def test_gerencia_reads_the_result_but_cannot_close(
        escena, fabrica, rol):
    await _con_bodegas(fabrica, escena)
    await rec._ronda_uno(escena)
    await rec._terminar_ronda(escena)
    gerencia = (rol, uuid.uuid4())
    ruta = f"{rec.BASE}/{escena.conteo.id}"

    cerrar = await rec.llamar(
        "POST", f"{ruta}/cerrar", gerencia, json={})
    await _cerrar(escena, forzar=True, motivo="Fin")
    leer = await rec.llamar("GET", f"{ruta}/resultado", gerencia)

    assert cerrar.status_code == 403
    assert leer.status_code == 200, leer.text
    assert leer.json()["total"] == 6
