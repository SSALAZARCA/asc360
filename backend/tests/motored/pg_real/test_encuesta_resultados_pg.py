"""
Motored satisfaction survey: result queries against a real Postgres (opt-in,
`MOTORED_TEST_PG_URL`, database migrated to head). The fake session cannot
check the outer join or the Bogota-day boundaries of the date range.
"""
import io
import os
import uuid
from datetime import date, datetime

import openpyxl
import pytest
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine

from app.motored.models.caso_detractor import CasoDetractor
from app.motored.models.caso_detractor_accion import CasoDetractorAccion
from app.motored.models.encuesta_carga import EncuestaCarga
from app.motored.models.encuesta_registro import EncuestaRegistro
from app.motored.models.encuesta_respuesta import EncuestaRespuesta
from app.motored.models.usuario import MotoredRole, Usuario
from app.motored.services import encuesta_excel
from app.motored.services import encuesta_resultados as servicio

URL = os.environ.get("MOTORED_TEST_PG_URL")
pytestmark = [
    pytest.mark.pg_real,
    pytest.mark.skipif(not URL, reason="MOTORED_TEST_PG_URL no definida"),
]


@pytest.fixture
async def sesion():
    motor = create_async_engine(URL)
    async with AsyncSession(motor, expire_on_commit=False) as db:
        yield db
    await motor.dispose()


async def _carga(db, created_at, nombres, respuestas=None):
    """Seed a carga with one registro per name; `respuestas` maps name -> (score, created_at)."""
    suffix = uuid.uuid4().int % 10**8
    usuario = Usuario(
        id=uuid.uuid4(), nombre="Admin", role=MotoredRole.ADMIN, activo=True,
        email=f"admin{suffix}@test.co", hashed_password="x",
    )
    db.add(usuario)
    await db.flush()
    carga = EncuestaCarga(
        id=uuid.uuid4(), nombre_archivo=f"t{suffix}.xlsx", total_registros=len(nombres),
        usuario_id=usuario.id, created_at=created_at,
    )
    db.add(carga)
    await db.flush()
    for i, nombre in enumerate(nombres):
        registro = EncuestaRegistro(
            id=uuid.uuid4(), carga_id=carga.id, tipo="SERVICIO_TALLER", nombre=nombre,
            cedula=f"{suffix}{i}", celular="3001112233", placa=f"P{i}{suffix}"[:16],
        )
        db.add(registro)
        await db.flush()
        if respuestas and nombre in respuestas:
            score, at = respuestas[nombre]
            db.add(EncuestaRespuesta(
                id=uuid.uuid4(), registro_id=registro.id, satisfaccion_general=score,
                autoriza_datos=True, created_at=at,
            ))
    await db.commit()
    return carga


async def test_detail_joins_answers_and_keeps_unanswered(sesion):
    carga = await _carga(sesion, datetime(2030, 3, 10, 12), ["Ana", "Beto"], {"Ana": (2, datetime(2030, 3, 11, 12))})
    filas = [servicio.fila_de(r) for r in await servicio.filas_de_carga(sesion, carga.id)]
    por_nombre = {f["cliente"]: f for f in filas}
    assert por_nombre["Ana"]["estado"] == "RESPONDIDA" and por_nombre["Ana"]["categoria"] == "DETRACTOR"
    assert por_nombre["Beto"]["estado"] == "SIN_RESPONDER" and por_nombre["Beto"]["nota"] is None


async def test_range_by_send_date_uses_inclusive_bogota_days(sesion):
    # 2031-05-31 23:59 Bogota = 2031-06-01 04:59 UTC (inside May); 05:00 UTC is June 1 Bogota (outside).
    dentro = await _carga(sesion, datetime(2031, 6, 1, 4, 59), ["Dentro"])
    fuera = await _carga(sesion, datetime(2031, 6, 1, 5, 0), ["Fuera"])
    inicio = await _carga(sesion, datetime(2031, 5, 1, 5, 0), ["Inicio"])
    antes = await _carga(sesion, datetime(2031, 5, 1, 4, 59), ["Antes"])
    filas = await servicio.filas_de_rango(sesion, date(2031, 5, 1), date(2031, 5, 31), "envio")
    nombres = {r.nombre for r in filas}
    assert {"Dentro", "Inicio"} <= nombres
    assert not ({"Fuera", "Antes"} & nombres)
    assert dentro and fuera and inicio and antes


async def test_range_by_response_date_only_returns_answered(sesion):
    await _carga(
        sesion, datetime(2032, 1, 5, 12), ["Resp", "Sin", "Tarde"],
        {"Resp": (5, datetime(2032, 2, 10, 12)), "Tarde": (4, datetime(2032, 3, 1, 5, 0))},
    )
    filas = await servicio.filas_de_rango(sesion, date(2032, 2, 1), date(2032, 2, 29), "respuesta")
    assert {r.nombre for r in filas} == {"Resp"}


async def test_excel_links_survey_detractor_case_and_actions(sesion):
    carga = await _carga(
        sesion, datetime(2033, 4, 10, 12), ["Ana", "Beto", "Sin"],
        {"Ana": (5, datetime(2033, 4, 11, 12)), "Beto": (2, datetime(2033, 4, 12, 15, 0))},
    )
    filas = await servicio.filas_de_carga(sesion, carga.id)
    beto = next(r for r in filas if r.nombre == "Beto")
    gestor = Usuario(
        id=uuid.uuid4(), nombre="Gestora Pg", role=MotoredRole.ADMIN, activo=True,
        email=f"g{uuid.uuid4().int % 10**8}@test.co", hashed_password="x",
    )
    sesion.add(gestor)
    await sesion.flush()
    caso = CasoDetractor(id=uuid.uuid4(), respuesta_id=beto.respuesta_id, asignado_a=gestor.id,
                         estado="EN_GESTION", created_at=datetime(2033, 4, 12, 15, 0))
    sesion.add(caso)
    await sesion.flush()
    sesion.add(CasoDetractorAccion(caso_id=caso.id, usuario_id=None, tipo="APERTURA",
                                   descripcion="Caso abierto", created_at=datetime(2033, 4, 12, 15, 0)))
    sesion.add(CasoDetractorAccion(caso_id=caso.id, usuario_id=gestor.id, tipo="LLAMADA",
                                   descripcion="Cliente contesta", created_at=datetime(2033, 4, 13, 14, 0)))
    await sesion.commit()

    casos, acciones = await servicio.casos_y_acciones(sesion, filas)
    assert [c.respuesta_id for c in casos] == [beto.respuesta_id]
    assert [a.tipo for a in acciones] == ["APERTURA", "LLAMADA"]

    libro = openpyxl.load_workbook(io.BytesIO(encuesta_excel.construir_por_carga(filas, casos, acciones)))
    assert libro.sheetnames == ["Encuestas"]
    valores = [[c.value for c in r] for r in libro.active.iter_rows()]
    cab = valores[0]
    por_cliente = {r[cab.index("Cliente")]: dict(zip(cab, r)) for r in valores[1:]}
    fila = por_cliente["Beto"]
    assert fila["ID encuesta"] == str(beto.registro_id)
    assert fila["ID caso"] == str(caso.id) and fila["N.º caso"] == casos[0].numero
    assert fila["Responsable"] == "Gestora Pg" and fila["Estado del caso"] == "EN_GESTION"
    assert fila["P1. Satisfacción general (1-5)"] == 2 and fila["Categoría"] == "Detractor"
    assert fila["N.º acciones"] == 2
    assert fila["Última acción (usuario)"] == "Gestora Pg"
    assert fila["Historial de gestión"].split("\n") == [
        "12/04/2033 10:00 · Sistema · APERTURA: Caso abierto",
        "13/04/2033 09:00 · Gestora Pg · LLAMADA: Cliente contesta",
    ]
    assert por_cliente["Ana"]["ID caso"] is None and por_cliente["Sin"]["Estado"] == "Sin responder"
