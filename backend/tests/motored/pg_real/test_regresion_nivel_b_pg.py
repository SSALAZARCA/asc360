"""
Motored Pedidos F3 "Motor", S8b (sdd/motored-pedidos-motor, ADR-10): niveles B
y C contra un Postgres real (opt-in).

Corre sólo con `MOTORED_TEST_PG_URL` (`postgresql+asyncpg://...`) apuntando a
una base ya migrada con `alembic -c alembic_motored.ini upgrade head`. Cada
test trabaja dentro de una transacción que se revierte al final (las funciones
de las herramientas reciben `confirmar=False`). Recorre el camino completo del
pipeline con datos sintéticos: preflight, cargador, motor, persistencia, nivel
B contra un Excel sintético, una corrida escenario por switch (nivel C) y el
libro de corrida.
"""
import io
import os
from contextlib import asynccontextmanager
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine

from app.motored.herramientas.regresion import comandos
from app.motored.herramientas.regresion.corrida_db import (
    datos_corrida,
    leer_nivel_b,
)
from app.motored.herramientas.regresion.delta import (
    NOMBRE_BASE,
    Escenario,
    escenarios_estandar,
)
from app.motored.herramientas.regresion.delta_db import (
    crear_y_calcular,
    medir_con_base,
)
from app.motored.herramientas.regresion.extractor_excel import leer_libro
from app.motored.herramientas.regresion.informe_excel import (
    escribir_libro_corrida,
)
from app.motored.herramientas.regresion.nivel_b import ejecutar_nivel_b
from app.motored.models.corrida import Corrida
from app.motored.models.parametro_metodologia import ParametroMetodologia
from tests.motored.fixtures.regresion.libro_sintetico import escribir_libro
from tests.motored.fixtures.regresion.oraculo_excel import (
    CLASES,
    FilaSintetica,
    SucursalSintetica,
)
from tests.motored.pg_real.test_corrida_pg import CORTE, PATRON, _sembrar

URL = os.environ.get("MOTORED_TEST_PG_URL")
pytestmark = [
    pytest.mark.pg_real,
    pytest.mark.skipif(not URL, reason="MOTORED_TEST_PG_URL no definida"),
]


@pytest.fixture
async def sesion():
    motor = create_async_engine(URL)
    async with AsyncSession(motor, expire_on_commit=False) as db:
        yield db
        await db.rollback()
    await motor.dispose()


def _libro(tmp_path: Path, inventario=27) -> Path:
    """Excel sintético de la sucursal UNO, con las mismas entradas."""
    ruta = tmp_path / "libro.xlsx"
    fila = FilaSintetica(
        "94109-12000S", PATRON, precio=460.75, inventario=inventario,
        transito=70,
    )
    escribir_libro(ruta, [fila], SucursalSintetica(), etiquetas=CLASES)
    return ruta


def _lectura(tmp_path: Path, inventario=27):
    return leer_libro(_libro(tmp_path, inventario))


def _abrir(sesion):
    """Fábrica de sesiones que reutiliza la del test (se revierte al final)."""
    @asynccontextmanager
    async def abrir(url):
        yield sesion
    return abrir


async def _comando(sesion, *argv):
    salida = io.StringIO()
    codigo = await comandos.ejecutar(
        list(argv), {}, salida, abrir_sesion=_abrir(sesion), confirmar=False
    )
    return codigo, salida.getvalue()


def _hojas(ruta: Path) -> dict:
    libro = openpyxl.load_workbook(ruta, read_only=True)
    hojas = {n: list(libro[n].iter_rows(values_only=True))
             for n in libro.sheetnames}
    libro.close()
    return {
        nombre: [fila or (None,) for fila in filas]
        for nombre, filas in hojas.items()
    }


async def _baseline(db, datos):
    return await crear_y_calcular(
        db, CORTE, datos.uno.id, confirmar=False
    )


async def test_el_nivel_b_pasa_con_las_mismas_entradas(sesion, tmp_path):
    datos = await _sembrar(sesion)
    corrida_id = await _baseline(sesion, datos)

    leida = await leer_nivel_b(sesion, corrida_id, datos.uno.id)
    resultado = ejecutar_nivel_b(
        _lectura(tmp_path), leida.app,
        reproduccion_identica=leida.reproduccion_identica,
    )

    assert leida.reproduccion_identica is True
    assert resultado.entradas == ()
    assert resultado.filas_excel == resultado.filas_app == 1
    assert resultado.salidas.diferencias == ()
    assert resultado.paso


async def test_una_entrada_distinta_se_reporta_antes_que_las_salidas(
    sesion, tmp_path
):
    datos = await _sembrar(sesion)
    corrida_id = await _baseline(sesion, datos)
    leida = await leer_nivel_b(sesion, corrida_id, datos.uno.id)

    resultado = ejecutar_nivel_b(_lectura(tmp_path, inventario=30), leida.app)

    (diferencia,) = resultado.entradas
    assert (diferencia.codigo, diferencia.columna) == ("94109-12000S", "V")
    assert (diferencia.valor_excel, diferencia.valor_motor) == (
        "30.0000000000", "27.0000000000"
    )
    assert not resultado.paso
    aceptado = ejecutar_nivel_b(
        _lectura(tmp_path, inventario=30), leida.app,
        {"94109-12000S": "T7"},
    )
    assert aceptado.paso


async def test_una_sucursal_omitida_no_sirve_para_el_nivel_b(sesion):
    datos = await _sembrar(sesion)
    with pytest.raises(ValueError, match="OMITIDA"):
        await crear_y_calcular(
            sesion, CORTE, datos.tres.id, confirmar=False
        )


async def test_el_nivel_c_mide_cada_switch_con_una_corrida_escenario(
    sesion,
):
    datos = await _sembrar(sesion)
    antes = await sesion.scalar(
        select(func.count()).select_from(ParametroMetodologia)
    )

    medicion = await medir_con_base(
        sesion, CORTE, datos.uno.id, escenarios_estandar(),
        confirmar=False,
    )

    resultados = medicion.resultados
    assert medicion.codigo_base.startswith("PED-2026-S39-")
    assert medicion.sucursal == datos.uno.nombre
    por_nombre = {r.escenario.nombre: r for r in resultados}
    assert resultados[0].escenario.nombre == NOMBRE_BASE
    assert resultados[0].conjunto.lineas["94109-12000S"].pedido == 53
    dias = por_nombre["dias_entre_pedidos=7"].delta
    assert dias.unidades_base == 53 and dias.unidades_escenario < 53
    ponderado = por_nombre["mes_en_curso=PONDERADO"]
    assert ponderado.delta.unidades_escenario < 53
    assert "PONDERADO con d=14" in ponderado.nota
    assert por_nombre["perdida x1"].delta.lineas_cambiadas == 0
    assert por_nombre["combinado"].delta is not None
    despues = await sesion.scalar(
        select(func.count()).select_from(ParametroMetodologia)
    )
    assert despues == antes
    escenarios = await sesion.scalar(
        select(func.count()).select_from(Corrida).where(
            Corrida.es_escenario.is_(True)
        )
    )
    assert escenarios == 7


async def test_un_override_invalido_deja_el_escenario_sin_medir(sesion):
    datos = await _sembrar(sesion)
    malo = Escenario("malo", "dias inválidos", {"dias_entre_pedidos": 0})

    medicion = await medir_con_base(
        sesion, CORTE, datos.uno.id, [malo], confirmar=False
    )

    assert medicion.resultados[1].delta is None
    assert "E-CORRIDA-010" in medicion.resultados[1].motivo_no_medido


async def test_el_libro_de_corrida_sale_de_lo_guardado(sesion, tmp_path):
    datos = await _sembrar(sesion)
    corrida_id = await crear_y_calcular(
        sesion, CORTE, datos.uno.id, confirmar=False
    )

    contenido = await datos_corrida(sesion, corrida_id)
    ruta = escribir_libro_corrida(contenido, tmp_path / "corrida.xlsx")

    assert contenido.titulo.startswith("PED-2026-S39-")
    assert contenido.antiguedad[0].antiguedad_dias == 2
    libro = openpyxl.load_workbook(ruta, read_only=True)
    filas = list(libro[datos.uno.nombre].iter_rows(values_only=True))
    libro.close()
    assert filas[1][0] == "94109-12000S"
    assert filas[1][27] == 53


async def test_el_subcomando_nivel_b_crea_la_corrida_y_compara(
    sesion, tmp_path
):
    datos = await _sembrar(sesion)
    salida = tmp_path / "reportes"

    codigo, texto = await _comando(
        sesion, "nivel-b", "--excel", str(_libro(tmp_path)),
        "--salida", str(salida), "--db-url", "postgresql+asyncpg://x/x",
        "--sucursal-id", str(datos.uno.id), "--fecha-corte", str(CORTE),
    )

    assert codigo == 0
    assert "B-entradas: 0 diferencias; PASA" in texto
    assert "Reproducción de la corrida desde su snapshot: idéntica" in texto
    hojas = _hojas(salida / "comparacion_nivel_b.xlsx")
    assert list(hojas)[:7] == [
        "Resumen", "Antigüedad de datos", "Filas", "A1", "A2",
        "B-entradas", "B-salidas",
    ]
    inventario = next(
        f for f in hojas["Antigüedad de datos"] if f[0] == "Inventario"
    )
    assert inventario[1:4] == (datetime(2026, 9, 19), 2, 7)


async def test_el_subcomando_nivel_b_usa_una_corrida_ya_calculada(
    sesion, tmp_path
):
    datos = await _sembrar(sesion)
    corrida_id = await _baseline(sesion, datos)

    codigo, _ = await _comando(
        sesion, "nivel-b", "--excel", str(_libro(tmp_path, inventario=30)),
        "--salida", str(tmp_path / "r"),
        "--db-url", "postgresql+asyncpg://x/x",
        "--sucursal-id", str(datos.uno.id), "--corrida-id", str(corrida_id),
    )

    assert codigo == 1


async def test_el_subcomando_delta_con_base_escribe_el_efecto_de_cada_switch(
    sesion, tmp_path
):
    datos = await _sembrar(sesion)
    salida = tmp_path / "reportes"

    codigo, texto = await _comando(
        sesion, "delta", "--salida", str(salida),
        "--db-url", "postgresql+asyncpg://x/x",
        "--sucursal-id", str(datos.uno.id), "--fecha-corte", str(CORTE),
    )

    assert codigo == 0
    hojas = _hojas(salida / "efecto_switches.xlsx")
    assert "Resumen switches" in hojas
    assert "Delta dias_entre_pedidos=7" in hojas
    assert "Delta mes_en_curso=PONDERADO" in hojas
    assert "dias_entre_pedidos=7: MEDIDO" in texto


async def test_el_subcomando_exportar_corrida_lee_la_corrida_guardada(
    sesion, tmp_path
):
    datos = await _sembrar(sesion)
    corrida_id = await _baseline(sesion, datos)
    salida = tmp_path / "reportes"

    codigo, texto = await _comando(
        sesion, "exportar-corrida", "--salida", str(salida),
        "--db-url", "postgresql+asyncpg://x/x",
        "--corrida-id", str(corrida_id),
    )

    assert codigo == 0
    archivos = [p.name for p in salida.iterdir()]
    assert len(archivos) == 1 and archivos[0].startswith("corrida_PED-2026")
    assert "1 sucursales" in texto


async def test_un_preflight_rechazado_sale_con_codigo_2_y_el_codigo_e(
    sesion, tmp_path
):
    datos = await _sembrar(sesion)

    codigo, texto = await _comando(
        sesion, "nivel-b", "--excel", str(_libro(tmp_path)),
        "--salida", str(tmp_path / "r"),
        "--db-url", "postgresql+asyncpg://x/x",
        "--sucursal-id", str(datos.uno.id), "--fecha-corte", "2027-03-15",
    )

    assert codigo == 2
    assert texto.startswith("Error: E-CORRIDA-")
    assert not (tmp_path / "r" / "comparacion_nivel_b.xlsx").exists()
