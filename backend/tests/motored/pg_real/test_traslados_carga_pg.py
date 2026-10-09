"""
TRASLADOS load against a real Postgres (opt-in, `MOTORED_TEST_PG_URL`;
odd/tasks/motored-traslados-pendientes.md, T1).

Runs the real dry-run and `Aplicar`: padded text is trimmed, the receiving
store resolves through `bodega.codigo`, an unknown destination is a row error,
an unknown origin or reference is kept, and the real ERP sample file loads
end to end. Each test is rolled back.
"""
import datetime
import io
import os
import uuid
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine

from app.motored.models.bodega import Bodega
from app.motored.models.carga_archivo import CargaArchivo
from app.motored.models.carga_error import CargaError
from app.motored.models.proveedor import Proveedor
from app.motored.models.referencia import Referencia
from app.motored.models.sucursal import Sucursal
from app.motored.models.traslado import TrasladoLinea
from app.motored.services.ingesta import orquestador, traslados
from tests.motored.pg_real.codigos_co import codigo_co_unico

URL = os.environ.get("MOTORED_TEST_PG_URL")
pytestmark = [
    pytest.mark.pg_real,
    pytest.mark.skipif(not URL, reason="MOTORED_TEST_PG_URL no definida"),
]

MUESTRA = Path("/mnt/c/Users/ALEJO S/Documents/Motored/Traslados_09.xlsx")
ENCABEZADO = list(traslados.COLUMNAS_ESPERADAS + traslados.COLUMNAS_OPCIONALES)


@pytest.fixture
async def sesion(monkeypatch):
    motor = create_async_engine(URL)
    async with AsyncSession(motor, expire_on_commit=False) as db:
        monkeypatch.setattr(db, "commit", db.flush)
        yield db
        await db.rollback()
    await motor.dispose()


async def _tienda(db, codigo_bodega):
    sfx = uuid.uuid4().hex[:6]
    tienda = Sucursal(id=uuid.uuid4(), nombre=f"T {codigo_bodega} {sfx}",
                      codigo_co=codigo_co_unico())
    db.add(tienda)
    await db.flush()
    db.add(Bodega(id=uuid.uuid4(), codigo=codigo_bodega, sucursal_id=tienda.id))
    await db.flush()
    return tienda


async def _proveedor(db):
    prov = Proveedor(
        id=uuid.uuid4(), codigo=f"P-{uuid.uuid4().hex[:6]}", nombre="P",
        es_principal=True, dias_empaque_default=4, dias_transito_default=5,
        dias_seguridad_default=Decimal("3"))
    db.add(prov)
    await db.flush()
    return prov


async def _referencia(db, codigo):
    prov = await _proveedor(db)
    ref = Referencia(id=uuid.uuid4(), codigo=codigo, proveedor_id=prov.id,
                     unidad_empaque=1, precio_normal=None, linea_comercial="X")
    db.add(ref)
    await db.flush()
    return ref


def _xlsx(filas):
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(ENCABEZADO)
    for fila in filas:
        hoja.append(list(fila))
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


async def _cargar(db, monkeypatch, contenido, aplicar=True):
    monkeypatch.setattr(
        orquestador.storage, "descargar_archivo", lambda ruta: contenido)
    carga = CargaArchivo(
        id=uuid.uuid4(), tipo="TRASLADOS", origen="EXCEL",
        estado="PROCESANDO", nombre_archivo="t.xlsx",
        hash_sha256=uuid.uuid4().hex * 2, ruta_objeto="r",
        bytes=len(contenido))
    db.add(carga)
    await db.flush()
    await orquestador._dry_run(db, carga)
    await db.flush()
    if aplicar:
        assert carga.estado == "VALIDADO", carga.log
        await orquestador.ejecutar_aplicar(db, carga)
        await db.flush()
    return carga


def _fila(doc, bod_sal, bod_ent, ref, cant=2, fecha=datetime.date(2026, 9, 1)):
    return (doc, fecha, bod_sal, bod_ent, ref, cant, "  SALIDA  ", "  ENTRADA ",
            "  DESC  ", "x", "UND ")


async def test_carga_recorta_resuelve_y_conserva_origen_y_referencia_desconocidos(
        sesion, monkeypatch):
    sfx = uuid.uuid4().hex[:5].upper()
    destino = await _tienda(sesion, f"D{sfx}")
    origen = await _tienda(sesion, f"O{sfx}")
    ref = await _referencia(sesion, f"REF-{sfx}")
    filas = [
        _fila(f" 79-{sfx} ", f" O{sfx} ", f"D{sfx}  ", f"  REF-{sfx}  "),
        _fila(f"79-{sfx}B", "ZZZ-NOEXISTE", f"D{sfx}", "OTRA-REF-SIN-CATALOGO"),
        _fila(f"79-{sfx}C", f"O{sfx}", "ZZZ-DESTINO-NO", f"REF-{sfx}"),
        _fila(f"79-{sfx}D", f"O{sfx}", f"D{sfx}", f"REF-{sfx}", cant=0),
    ]

    carga = await _cargar(sesion, monkeypatch, _xlsx(filas))

    lineas = (await sesion.execute(
        select(TrasladoLinea).where(TrasladoLinea.carga_id == carga.id)
        .order_by(TrasladoLinea.nro_documento))).scalars().all()
    assert [l.nro_documento for l in lineas] == [f"79-{sfx}", f"79-{sfx}B"]
    primera, segunda = lineas
    assert primera.sucursal_entrada_id == destino.id
    assert primera.sucursal_salida_id == origen.id
    assert primera.referencia_id == ref.id
    assert primera.bodega_salida == f"O{sfx}"
    assert primera.cantidad == Decimal("2")
    assert segunda.sucursal_salida_id is None
    assert segunda.bodega_salida == "ZZZ-NOEXISTE"
    assert segunda.referencia_id is None
    assert segunda.referencia_codigo == "OTRA-REF-SIN-CATALOGO"
    codigos = (await sesion.execute(
        select(CargaError.codigo_error).where(CargaError.carga_id == carga.id)
        .order_by(CargaError.fila))).scalars().all()
    assert codigos == ["SUCURSAL_NO_ENCONTRADA", "CANTIDAD_INVALIDA"]
    assert carga.estado == "APLICADO"


async def test_un_archivo_sin_filas_validas_queda_con_errores(
        sesion, monkeypatch):
    await _proveedor(sesion)
    carga = await _cargar(
        sesion, monkeypatch,
        _xlsx([_fila("79-1", "A", "NOEXISTE-1", "R")]), aplicar=False)

    assert carga.estado == "CON_ERRORES"


@pytest.mark.skipif(not MUESTRA.exists(), reason="archivo de muestra ausente")
async def test_el_archivo_real_de_muestra_carga_de_punta_a_punta(
        sesion, monkeypatch):
    await _proveedor(sesion)
    libro = openpyxl.load_workbook(MUESTRA, data_only=True)
    filas = list(libro.active.iter_rows(values_only=True))[1:]
    codigos = {str(f[2]).strip() for f in filas} | {str(f[4]).strip() for f in filas}
    existentes = set((await sesion.execute(
        select(Bodega.codigo).where(Bodega.codigo.in_(codigos)))).scalars().all())
    for codigo in sorted(codigos - existentes):
        await _tienda(sesion, codigo)

    carga = await _cargar(sesion, monkeypatch, MUESTRA.read_bytes())

    total = (await sesion.execute(
        select(func.count()).select_from(TrasladoLinea)
        .where(TrasladoLinea.carga_id == carga.id))).scalar_one()
    documentos = (await sesion.execute(
        select(func.count(func.distinct(
            TrasladoLinea.nro_documento, TrasladoLinea.bodega_salida)))
        .where(TrasladoLinea.carga_id == carga.id))).scalar_one()
    assert (carga.filas_leidas, total, documentos) == (143, 143, 103)
    assert carga.filas_rechazadas == 0
    sin_espacios = (await sesion.execute(
        select(func.count()).select_from(TrasladoLinea).where(
            TrasladoLinea.carga_id == carga.id,
            TrasladoLinea.referencia_codigo != func.trim(
                TrasladoLinea.referencia_codigo)))).scalar_one()
    assert sin_espacios == 0
