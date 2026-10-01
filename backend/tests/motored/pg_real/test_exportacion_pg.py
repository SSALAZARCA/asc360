"""
Fase 4 "Pantallas del pedido" (sdd/motored-pedidos-ui, B4, ADR-7, spec
EX-01..EX-15, decisiones F4-1 y A5): exportar el pedido a HMCL y leer el
bloque `envio`, contra un Postgres real (opt-in).

Primero las reglas con sesiones sucesivas, sobre el SQL de verdad (los
filtros `pedido_final > 0` y `motivo_exclusion IS NULL`, el estado de cada
tienda, el SIC y el escenario) y con el libro abierto con openpyxl; luego las
carreras con DOS conexiones: una exportación en curso retiene la tienda
`FOR SHARE` y un reabrir espera, y un reabrir sin confirmar hace esperar a la
exportación, que después responde 055 en vez de exportar un pedido a medio
cambiar. Al final, el bloque `envio` del detalle y de la cabecera.

La última parte corre las RUTAS reales (la app completa, httpx sobre ASGI) con
sesiones reales de Postgres: lo que descarga un cliente de punta a punta.

Siembra, CONFIRMADO: la de `test_envio_pg.py` (un proveedor, un usuario,
dos tiendas OK con una línea de 50 en cada una de tres corridas BORRADOR).
"""
import asyncio
import io
import json
import uuid
import zipfile
from urllib.parse import unquote

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import update

from app.config import settings
from app.main import app
from app.motored.deps import get_current_motored_user, get_motored_db
from app.motored.models.corrida import Corrida
from app.motored.models.corrida_linea import CorridaLinea
from app.motored.models.corrida_sucursal import CorridaSucursal
from app.motored.models.sucursal import Sucursal
from app.motored.services.auth import MotoredUser
from app.motored.services.corridas import (
    codigos,
    consultas,
    envio,
    exportacion,
    lecturas_pedido,
    pedido_tienda,
)
from tests.motored.pg_real.test_corrida_pg import CORTE
from tests.motored.pg_real.test_envio_pg import (  # noqa: F401 (fixture)
    FECHA,
    OTRO_CORTE,
    URL,
    _cerrar,
    _enviar,
    _error,
    _esperar,
    _hacer,
    _reabrir,
    _terminar,
    escenario,
)

pytestmark = [
    pytest.mark.pg_real,
    pytest.mark.skipif(not URL, reason="MOTORED_TEST_PG_URL no definida"),
]


async def _preparar(esc, tienda, corrida=None):
    return await _hacer(esc, lambda db: exportacion.preparar_tienda(
        db, corrida or esc.c1, tienda))


async def _zip(esc, ids=None, corrida=None):
    return await _hacer(esc, lambda db: exportacion.preparar_corrida(
        db, corrida or esc.c1, ids))


async def _actualizar(esc, sentencia):
    async with esc.maker() as db:
        await db.execute(sentencia)
        await db.commit()


def _hoja(archivo):
    contenido = archivo.read()
    archivo.close()
    return load_workbook(io.BytesIO(contenido))["Pedido"]


# --- Una tienda -------------------------------------------------------------


async def test_a_closed_tienda_exports_its_header_and_its_line_ex_01(
        escenario):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)

    datos = await _preparar(esc, esc.a)
    archivo, _ = exportacion.construir_xlsx(datos)

    hoja = _hoja(archivo)
    assert [c.value for c in hoja["A"][:5]] == [
        "Tienda", "SIC", "Fecha", None, "Código"]
    assert hoja["B1"].value == esc.nombre_a
    assert hoja["B2"].value.startswith("S-A-")
    assert hoja["B3"].value.date() == CORTE
    assert [c.value for c in hoja[6]][1] == 50
    assert hoja["A6"].value.startswith("R1-")


async def test_a_borrador_tienda_is_055_and_nothing_is_exported_ex_06(
        escenario):
    esc = escenario

    error = await _error(_preparar(esc, esc.a))

    assert error.codigo == codigos.E_CORRIDA_EXPORTAR_NO_CERRADO
    assert esc.nombre_a in error.mensaje


async def test_a_sent_tienda_can_be_exported_again_ex_12_a5(escenario):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)
    await _enviar(esc, esc.c1, esc.a, "12345")

    primero = await _preparar(esc, esc.a)
    segundo = await _preparar(esc, esc.a)

    assert primero == segundo and primero.lineas[0][1] == 50


async def test_after_a_reopen_the_tienda_is_055_but_others_export_ex_11(
        escenario):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)
    await _cerrar(esc, esc.c1, esc.b)
    await _reabrir(esc, esc.c1, esc.a)

    error = await _error(_preparar(esc, esc.a))
    otra = await _preparar(esc, esc.b)

    assert error.codigo == codigos.E_CORRIDA_EXPORTAR_NO_CERRADO
    assert otra.nombre == esc.nombre_b


async def test_the_exported_value_is_the_edited_one_ex_03(escenario):
    esc = escenario
    await _actualizar(esc, update(CorridaLinea).where(
        CorridaLinea.corrida_id == esc.c1, CorridaLinea.sucursal_id == esc.a
    ).values(pedido_final=60))
    await _cerrar(esc, esc.c1, esc.a)

    datos = await _preparar(esc, esc.a)

    assert datos.lineas[0][1] == 60


async def test_a_pedido_at_zero_or_with_only_excluded_lines_is_056_ex_08(
        escenario):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)
    await _cerrar(esc, esc.c1, esc.b)
    linea = update(CorridaLinea).where(
        CorridaLinea.corrida_id == esc.c1, CorridaLinea.sucursal_id == esc.a)
    await _actualizar(esc, linea.values(pedido_final=0))
    await _actualizar(esc, update(CorridaLinea).where(
        CorridaLinea.corrida_id == esc.c1, CorridaLinea.sucursal_id == esc.b
    ).values(motivo_exclusion="SUSTITUIDA"))

    cero = await _error(_preparar(esc, esc.a))
    excluida = await _error(_preparar(esc, esc.b))

    assert cero.codigo == codigos.E_CORRIDA_NADA_QUE_ENVIAR
    assert excluida.codigo == codigos.E_CORRIDA_NADA_QUE_ENVIAR


@pytest.mark.parametrize("sic", [None, "", "  "])
async def test_a_tienda_without_sic_is_057_naming_it_ex_09(escenario, sic):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)
    await _actualizar(esc, update(Sucursal).where(
        Sucursal.id == esc.a).values(sic=sic))

    error = await _error(_preparar(esc, esc.a))

    assert error.codigo == codigos.E_CORRIDA_EXPORTAR_SIN_SIC
    assert esc.nombre_a in error.mensaje


async def test_a_failed_tienda_has_no_pedido_065_and_an_unknown_one_is_404(
        escenario):
    esc = escenario
    await _actualizar(esc, update(CorridaSucursal).where(
        CorridaSucursal.corrida_id == esc.c1,
        CorridaSucursal.sucursal_id == esc.a
    ).values(estado="FALLIDA", estado_pedido=None))

    sin_pedido = await _error(_preparar(esc, esc.a))

    assert sin_pedido.codigo == codigos.E_CORRIDA_SIN_PEDIDO
    with pytest.raises(LookupError):
        await _preparar(esc, esc.a, corrida=uuid.UUID(int=1))


async def test_a_scenario_is_042_before_looking_at_the_state_ex_07(escenario):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)
    await _actualizar(esc, update(Corrida).where(
        Corrida.id == esc.c1).values(es_escenario=True))

    error = await _error(_preparar(esc, esc.a))

    assert error.codigo == codigos.E_CORRIDA_ESCENARIO_NO_SE_CIERRA
    assert "exportar" in error.mensaje


# --- La corrida entera ------------------------------------------------------


async def test_the_zip_has_the_closed_tiendas_and_lists_the_borrador_one_ex_02(
        escenario):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)

    seleccion = await _zip(esc)

    assert [t.nombre for t in seleccion.tiendas] == [esc.nombre_a]
    assert [(o["nombre"], o["codigo"]) for o in seleccion.omitidas] == [
        (esc.nombre_b, "BORRADOR")]


async def test_closed_and_sent_tiendas_both_go_in_the_zip(escenario):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)
    await _cerrar(esc, esc.c1, esc.b)
    await _enviar(esc, esc.c1, esc.a, "12345")

    seleccion = await _zip(esc)
    archivo, tamano = exportacion.construir_zip(seleccion.tiendas)

    assert sorted(t.nombre for t in seleccion.tiendas) == sorted(
        [esc.nombre_a, esc.nombre_b])
    assert seleccion.omitidas == []
    assert tamano == len(archivo.read()) > 0
    archivo.close()


async def test_a_zip_with_every_tienda_in_borrador_is_055_ex_06(escenario):
    esc = escenario

    error = await _error(_zip(esc))

    assert error.codigo == codigos.E_CORRIDA_EXPORTAR_NO_CERRADO


async def test_a_listed_borrador_tienda_is_055_and_a_listed_closed_one_goes(
        escenario):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)

    error = await _error(_zip(esc, [esc.a, esc.b]))
    seleccion = await _zip(esc, [esc.a])

    assert error.codigo == codigos.E_CORRIDA_EXPORTAR_NO_CERRADO
    assert esc.nombre_b in error.mensaje
    assert [t.nombre for t in seleccion.tiendas] == [esc.nombre_a]
    assert seleccion.omitidas == []


async def test_a_closed_tienda_with_nothing_to_order_is_skipped_ex_10(
        escenario):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)
    await _cerrar(esc, esc.c1, esc.b)
    await _actualizar(esc, update(CorridaLinea).where(
        CorridaLinea.corrida_id == esc.c1, CorridaLinea.sucursal_id == esc.b
    ).values(pedido_final=0))

    seleccion = await _zip(esc)

    assert [t.nombre for t in seleccion.tiendas] == [esc.nombre_a]
    assert [o["codigo"] for o in seleccion.omitidas] == ["SIN_CANTIDAD"]


async def test_the_other_corrida_of_the_same_corte_exports_on_its_own(
        escenario):
    esc = escenario
    await _cerrar(esc, esc.c2, esc.a)

    seleccion = await _zip(esc, corrida=esc.c2)
    error = await _error(_zip(esc, corrida=esc.c1))

    assert [t.nombre for t in seleccion.tiendas] == [esc.nombre_a]
    assert error.codigo == codigos.E_CORRIDA_EXPORTAR_NO_CERRADO


# --- Carreras ---------------------------------------------------------------


async def test_a_reopen_waits_for_an_export_in_flight_and_the_file_is_whole(
        escenario):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)
    async with esc.maker() as exporta, esc.maker() as reabre:
        datos = await exportacion.preparar_tienda(exporta, esc.c1, esc.a)
        tarea = asyncio.create_task(pedido_tienda.reabrir_tienda(
            reabre, esc.c1, esc.a, esc.usuario_id, "ajuste"))
        await _esperar(tarea, "el reabrir debía esperar a la exportación")

        await exporta.commit()
        await _terminar(tarea)
        await reabre.commit()

    assert datos.lineas[0][1] == 50


async def test_an_export_waits_for_a_reopen_in_flight_and_is_refused_055(
        escenario):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)
    async with esc.maker() as reabre, esc.maker() as exporta:
        await pedido_tienda.reabrir_tienda(
            reabre, esc.c1, esc.a, esc.usuario_id, "ajuste")
        tarea = asyncio.create_task(exportacion.preparar_tienda(
            exporta, esc.c1, esc.a))
        await _esperar(tarea, "la exportación debía esperar al reabrir")

        await reabre.commit()
        with pytest.raises(codigos.ErrorCorrida) as error:
            await _terminar(tarea)
        await exporta.rollback()

    assert error.value.codigo == codigos.E_CORRIDA_EXPORTAR_NO_CERRADO


async def test_two_exports_do_not_wait_for_each_other(escenario):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)
    async with esc.maker() as una, esc.maker() as otra:
        primera = await exportacion.preparar_tienda(una, esc.c1, esc.a)
        segunda = await asyncio.wait_for(
            exportacion.preparar_tienda(otra, esc.c1, esc.a), timeout=5)
        await una.rollback()
        await otra.rollback()

    assert primera == segunda


# --- El bloque envio de las lecturas ----------------------------------------


async def test_the_envio_block_reaches_the_detail_and_the_header(escenario):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)
    await _cerrar(esc, esc.c1, esc.b)
    await _enviar(esc, esc.c1, esc.a, " 12345 ")

    async with esc.maker() as db:
        detalle = await consultas.detalle(db, esc.c1, None)
        cabecera_a = await lecturas_pedido.cabecera_tienda(
            db, esc.c1, esc.a, None)
        cabecera_b = await lecturas_pedido.cabecera_tienda(
            db, esc.c1, esc.b, None)

    filas = {f["sucursal_id"]: f for f in detalle["sucursales"]}
    bloque = filas[esc.a]["envio"]
    assert bloque["numero_orden"] == "12345"
    assert bloque["fecha_envio"] == FECHA
    assert bloque["enviado_por"] == "Compras"
    assert bloque["enviado_en"] is not None
    assert filas[esc.b]["envio"] is None
    assert cabecera_a["envio"] == bloque
    assert cabecera_b["envio"] is None
    assert filas[esc.a]["acciones"]["exportar"] is True
    assert filas[esc.b]["acciones"]["exportar"] is True


async def test_a_corrected_number_shows_in_the_envio_block(escenario):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)
    await _enviar(esc, esc.c1, esc.a, "11111")
    await _hacer(esc, lambda db: envio.corregir_numero(
        db, esc.c1, esc.a, "22222", esc.usuario_id))

    async with esc.maker() as db:
        envios = await lecturas_pedido.envios_de(db, esc.c1, None)

    assert envios[esc.a].numero == "22222"


async def test_the_envios_stay_inside_the_scope_and_the_corrida(escenario):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)
    await _cerrar(esc, esc.c1, esc.b)
    await _enviar(esc, esc.c1, esc.a, "1")
    await _enviar(esc, esc.c1, esc.b, "2")
    await _cerrar(esc, esc.c3, esc.a)
    await _enviar(esc, esc.c3, esc.a, "3", fecha=OTRO_CORTE)

    async with esc.maker() as db:
        todos = await lecturas_pedido.envios_de(db, esc.c1, None)
        acotado = await lecturas_pedido.envios_de(
            db, esc.c1, frozenset({esc.b}))
        una = await lecturas_pedido.envios_de(db, esc.c1, None, esc.a)

    assert {k: v.numero for k, v in todos.items()} == {
        esc.a: "1", esc.b: "2"}
    assert list(acotado) == [esc.b] and list(una) == [esc.a]


async def test_a_detail_with_nothing_sent_still_reads_with_no_envios(
        escenario):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)

    async with esc.maker() as db:
        detalle = await consultas.detalle(db, esc.c1, None)

    assert [s["envio"] for s in detalle["sucursales"]] == [None, None]


# --- De punta a punta: las rutas reales con Postgres real -------------------


@pytest.fixture
async def cliente(escenario, monkeypatch):
    esc = escenario
    monkeypatch.setattr(settings, "MOTORED_ENABLED", True)
    monkeypatch.setattr(settings, "MOTORED_SECRET_KEY", "export-pg")

    async def sesion():
        async with esc.maker() as db:
            yield db

    async def usuario():
        return MotoredUser(user_id=str(esc.usuario_id), role="COMPRAS")

    app.dependency_overrides[get_motored_db] = sesion
    app.dependency_overrides[get_current_motored_user] = usuario
    transporte = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(
            transport=transporte, base_url="http://motored") as http:
        yield http
    app.dependency_overrides.clear()


async def test_the_real_route_downloads_the_xlsx_of_a_closed_tienda(
        escenario, cliente):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)

    respuesta = await cliente.get(
        f"/api/motored/corridas/{esc.c1}/sucursales/{esc.a}/exportar")

    assert respuesta.status_code == 200, respuesta.text
    assert respuesta.headers["content-disposition"].startswith(
        'attachment; filename="Pedido_SIC')
    assert int(respuesta.headers["content-length"]) == len(respuesta.content)
    hoja = load_workbook(io.BytesIO(respuesta.content))["Pedido"]
    assert hoja["B1"].value == esc.nombre_a
    assert [c.value for c in hoja[5]] == ["Código", "Cantidad"]
    assert hoja["B6"].value == 50


async def test_the_real_route_downloads_a_zip_and_lists_the_skipped_tiendas(
        escenario, cliente):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)

    respuesta = await cliente.get(f"/api/motored/corridas/{esc.c1}/exportar")

    assert respuesta.status_code == 200, respuesta.text
    assert respuesta.headers["content-type"] == "application/zip"
    paquete = zipfile.ZipFile(io.BytesIO(respuesta.content))
    assert len(paquete.namelist()) == 1
    omitidas = json.loads(unquote(respuesta.headers["x-tiendas-omitidas"]))
    assert [(o["nombre"], o["codigo"]) for o in omitidas] == [
        (esc.nombre_b, "BORRADOR")]


async def test_the_real_route_refuses_a_borrador_tienda_with_a_coded_409(
        escenario, cliente):
    esc = escenario

    respuesta = await cliente.get(
        f"/api/motored/corridas/{esc.c1}/sucursales/{esc.a}/exportar")

    assert respuesta.status_code == 409
    assert respuesta.json()["detail"]["code"] == "E-CORRIDA-055"
    assert respuesta.headers["content-type"] == "application/json"


async def test_the_real_detail_and_header_serve_the_envio_block(
        escenario, cliente):
    esc = escenario
    await _cerrar(esc, esc.c1, esc.a)
    await _enviar(esc, esc.c1, esc.a, "12345")

    detalle = await cliente.get(f"/api/motored/corridas/{esc.c1}")
    cabecera = await cliente.get(
        f"/api/motored/corridas/{esc.c1}/sucursales/{esc.a}")

    filas = {f["sucursal_id"]: f for f in detalle.json()["sucursales"]}
    bloque = filas[str(esc.a)]["envio"]
    assert bloque["numero_orden"] == "12345"
    assert bloque["fecha_envio"] == FECHA.isoformat()
    assert bloque["enviado_por"] == "Compras"
    assert filas[str(esc.b)]["envio"] is None
    assert cabecera.json()["envio"] == bloque
