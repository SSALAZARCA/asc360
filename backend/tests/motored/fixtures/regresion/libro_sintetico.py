"""
Motored Pedidos F3 "Motor" (S8a) — libro de Excel sintético con la
distribución de la hoja `Pedido` real.

`openpyxl` no puede escribir una fórmula junto con su valor cacheado, así que
el libro se arma con fórmulas (texto) y luego se inyecta el valor cacheado en
el XML de cada hoja. El extractor lee los dos pasos (`data_only` False/True)
exactamente como con el archivo real. Sólo datos inventados.
"""
import zipfile
from pathlib import Path
from typing import Mapping, Optional, Sequence
from xml.etree import ElementTree

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from tests.motored.fixtures.regresion.oraculo_excel import (
    CLASES,
    ETIQUETAS_EXCEL,
    FilaSintetica,
    ResultadoHoja,
    SucursalSintetica,
    calcular_hoja,
)

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
PRIMERA_FILA = 14
COLUMNAS_FORMULA = (
    "K", "N", "O", "P", "Q", "R", "T", "U", "V", "W", "X", "Y",
    "AA", "AB", "AC", "AD",
)
COLUMNAS_VENTAS = ("E", "F", "G", "H", "I", "J")
ColumnasCache = dict[str, dict[str, object]]


def _formula_fila(columna: str, fila: int, divisor: int) -> str:
    if columna == "N":
        return (f"=+(E{fila}*1+F{fila}*2+G{fila}*3+H{fila}*4+I{fila}*5"
                f"+J{fila}*6)/{divisor}")
    return f"=+SUMIFS(origen!{columna}:{columna},origen!A:A,A{fila})"


def _escribir_fila(hoja, cache: dict, fila: int, datos: FilaSintetica,
                   valores: Mapping, s_formula: bool) -> None:
    hoja[f"A{fila}"] = datos.codigo
    for columna, venta in zip(COLUMNAS_VENTAS, datos.ventas):
        hoja[f"{columna}{fila}"] = venta
    hoja[f"Z{fila}"] = datos.ajuste
    hoja[f"S{fila}"] = f"=+CONCATENATE(Q{fila},R{fila})" if s_formula \
        else valores["S"]
    if s_formula:
        cache[f"S{fila}"] = valores["S"]
    for columna in COLUMNAS_FORMULA:
        hoja[f"{columna}{fila}"] = _formula_fila(columna, fila, datos.divisor)
        cache[f"{columna}{fila}"] = valores[columna]


def _escribir_cabecera(hoja, cache: dict, sucursal: SucursalSintetica,
                       calculo: ResultadoHoja) -> None:
    hoja["D4"] = f"{sucursal.nombre}          "
    hoja["E5"] = "=VLOOKUP(D4,cobertura!A:C,3,0)"
    cache["E5"] = sucursal.bodega
    hoja["E6"] = "=VLOOKUP(D4,cobertura!A:C,2,0)"
    cache["E6"] = sucursal.sic
    for posicion, clase in enumerate(CLASES):
        columna = get_column_letter(5 + posicion)
        hoja[f"{columna}2"] = clase
        hoja[f"{columna}3"] = "=VLOOKUP($D$4,cobertura!A5:T138,10,0)"
        cache[f"{columna}3"] = calculo.coberturas[clase]


def _escribir_resumen(hoja, cache: dict, calculo: ResultadoHoja) -> None:
    for columna, titulo in zip(
        ("AA", "AB", "AC", "AD", "AE"),
        ("Clas", "Pedido uns", "Referencias", "Vlr Ped", "% Peso."),
    ):
        hoja[f"{columna}1"] = titulo
    bloques = list(calculo.resumen) + [calculo.total]
    for fila, bloque in enumerate(bloques, start=2):
        hoja[f"AA{fila}"] = bloque[0]
        for columna, valor in zip(("AB", "AC", "AD", "AE"), bloque[1:]):
            hoja[f"{columna}{fila}"] = f"=SUMIF($S$14:$S$99,AA{fila},X)"
            cache[f"{columna}{fila}"] = valor


def _escribir_cobertura(libro: Workbook, cache: dict,
                        sucursal: SucursalSintetica) -> None:
    hoja = libro.create_sheet("Meses de cobertura")
    for columna, titulo in zip(
        "ABCFGHI",
        ("Sucursal", "SIC", "Bodega Principal", "dias empaque",
         "dias transito", "total", "total+ dias de seguridad"),
    ):
        hoja[f"{columna}5"] = titulo
    hoja["A6"] = "OTRA SUCURSAL     "
    hoja["F6"], hoja["G6"] = 9, 9
    hoja["I6"] = "=H6+2.5"
    cache["I6"] = 20.5
    hoja["A7"] = f"{sucursal.nombre}      "
    hoja["B7"], hoja["C7"] = sucursal.sic, sucursal.bodega
    hoja["F7"], hoja["G7"] = sucursal.dias_empaque, sucursal.dias_transito
    hoja["H7"], hoja["I7"] = "=+G7+F7", "=+H7+2.5"
    cache["H7"] = sucursal.dias_empaque + sucursal.dias_transito
    cache["I7"] = sucursal.intervalo


def _inyectar_cache(ruta: Path, hojas: Sequence[str],
                    cache: Mapping[str, Mapping[str, object]]) -> None:
    """Escribe el valor cacheado de cada celda con fórmula en el XML."""
    ElementTree.register_namespace("", NS)
    with zipfile.ZipFile(ruta) as origen:
        miembros = {n: origen.read(n) for n in origen.namelist()}
    for indice, nombre in enumerate(hojas, start=1):
        clave = f"xl/worksheets/sheet{indice}.xml"
        raiz = ElementTree.fromstring(miembros[clave])
        for celda in raiz.iter(f"{{{NS}}}c"):
            valor = cache.get(nombre, {}).get(celda.get("r"))
            if celda.find(f"{{{NS}}}f") is not None and valor is not None:
                _fijar_valor(celda, valor)
        miembros[clave] = ElementTree.tostring(raiz)
    with zipfile.ZipFile(ruta, "w", zipfile.ZIP_DEFLATED) as destino:
        for nombre, contenido in miembros.items():
            destino.writestr(nombre, contenido)


def _fijar_valor(celda, valor) -> None:
    nodo = celda.find(f"{{{NS}}}v")
    if isinstance(valor, str):
        celda.set("t", "e" if valor.startswith("#") else "str")
        nodo.text = valor
    else:
        nodo.text = repr(valor)


def escribir_libro(
    ruta: Path, filas: Sequence[FilaSintetica],
    sucursal: Optional[SucursalSintetica] = None, *,
    etiquetas: Sequence[str] = ETIQUETAS_EXCEL,
    sobrescribe: Optional[Mapping[str, Mapping]] = None,
    s_formula: bool = False,
) -> ResultadoHoja:
    """Escribe el libro en `ruta` y devuelve los valores cacheados."""
    sucursal = sucursal or SucursalSintetica()
    calculo = calcular_hoja(
        filas, sucursal, etiquetas=etiquetas, sobrescribe=sobrescribe
    )
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Pedido"
    cache: ColumnasCache = {"Pedido": {}, "Meses de cobertura": {}}
    _escribir_cabecera(hoja, cache["Pedido"], sucursal, calculo)
    _escribir_resumen(hoja, cache["Pedido"], calculo)
    for fila, (datos, valores) in enumerate(
        zip(filas, calculo.filas), start=PRIMERA_FILA
    ):
        _escribir_fila(
            hoja, cache["Pedido"], fila, datos, valores, s_formula
        )
    ultima = PRIMERA_FILA + len(filas)
    hoja[f"N{ultima}"] = f"=+(E{ultima}*1)/21"
    cache["Pedido"][f"N{ultima}"] = 0.0
    _escribir_cobertura(libro, cache["Meses de cobertura"], sucursal)
    libro.save(ruta)
    _inyectar_cache(ruta, libro.sheetnames, cache)
    return calculo
