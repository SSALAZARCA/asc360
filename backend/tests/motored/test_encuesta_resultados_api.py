"""
Motored satisfaction survey: per-carga detail and Excel exports
(`/api/motored/encuesta/cargas/{id}/detalle`, `/{id}/excel`, `/resultados/excel`).
Same conventions as `test_encuesta_carga_api.py`: fake session, one queued
result per query (the first queued list is the readiness probe).
"""
import io
import uuid
from datetime import date, datetime
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.config import settings
from app.main import app
from app.motored.services import encuesta_excel, encuesta_resultados
from app.motored.services.auth import MotoredUser
from tests.motored.conftest import FakeAsyncSession, override_motored_db, override_motored_user

BASE = "/api/motored/encuesta/cargas"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
USER_ID = str(uuid.uuid4())
CARGA_ID = uuid.uuid4()
# 2026-09-30 23:30 UTC is 18:30 in Bogota; 2026-10-01 03:00 UTC is still 2026-09-30 22:00 there.
CARGA_AT = datetime(2026, 10, 1, 3, 0, 0)


@pytest.fixture(autouse=True)
def _motored_ready(monkeypatch):
    monkeypatch.setattr(settings, "MOTORED_ENABLED", True)
    monkeypatch.setattr(settings, "MOTORED_SECRET_KEY", "encuesta-res-test-motored-secret")
    monkeypatch.setattr(settings, "SECRET_KEY", "encuesta-res-test-asc360-secret")
    override_motored_user(MotoredUser(user_id=USER_ID, role="ADMIN"))
    yield
    app.dependency_overrides.clear()


def _row(nombre="Ana Perez", cedula="0012345", celular="3001112233", centro="Cali Norte",
         nota=None, obs=None, respondida_at=None, archivo="mayo.xlsx", carga_at=CARGA_AT,
         matriz=(4, 5, None, 3, 2, 1), autoriza=True):
    respondida = nota is not None
    p = matriz if respondida else (None,) * 6
    return SimpleNamespace(
        registro_id=uuid.uuid4(), respuesta_id=uuid.uuid4() if respondida else None,
        placa="ABC123", linea="Boxer", sic="S1",
        nombre_archivo=archivo, carga_created_at=carga_at, nombre=nombre, cedula=cedula,
        celular=celular, centro_servicio=centro, satisfaccion_general=nota,
        p_explicacion_tecnica=p[0], p_confianza_reparacion=p[1], p_servicio_taller=p[2],
        p_calidad_mecanicos=p[3], p_claridad_cobros=p[4], p_originalidad_repuestos=p[5],
        observaciones=obs, autoriza_datos=autoriza if respondida else None,
        respuesta_created_at=respondida_at,
    )


ROWS = [
    _row("Ana", nota=5, obs="Excelente", respondida_at=datetime(2026, 10, 2, 15, 0, 0)),
    _row("Beto", nota=2, obs="Mal servicio", respondida_at=datetime(2026, 10, 3, 4, 0, 0)),
    _row("Carlos", nota=3),
    _row("Dora"),
]
ROWS[2].satisfaccion_general = 3
ROWS[2].respuesta_created_at = datetime(2026, 10, 3, 5, 0, 0)


def _get(url, rows, casos=(), acciones=()):
    session = FakeAsyncSession(execute_queue=[[], rows, list(casos), list(acciones)])
    override_motored_db(session)
    with TestClient(app) as client:
        return client.get(url), session


# --- classification and estado ---------------------------------------------

@pytest.mark.parametrize("nota,esperada", [(1, "DETRACTOR"), (3, "DETRACTOR"), (4, "SATISFECHO"), (5, "SATISFECHO"), (None, None)])
def test_categoria_reuses_the_existing_detractor_threshold(nota, esperada):
    assert encuesta_resultados.clasificar(nota) == esperada


def test_detail_derives_estado_categoria_and_utc_date():
    response, session = _get(f"{BASE}/{CARGA_ID}/detalle", ROWS)
    assert response.status_code == 200
    filas = {f["cliente"]: f for f in response.json()}
    assert filas["Ana"] == {
        "cliente": "Ana", "cedula": "0012345", "telefono": "3001112233", "tienda": "Cali Norte",
        "estado": "RESPONDIDA", "nota": 5, "categoria": "SATISFECHO", "comentario": "Excelente",
        "fecha_respuesta": "2026-10-02T15:00:00+00:00",
    }
    assert filas["Beto"]["categoria"] == "DETRACTOR"
    assert filas["Carlos"]["categoria"] == "DETRACTOR"
    assert filas["Dora"]["estado"] == "SIN_RESPONDER"
    assert filas["Dora"]["nota"] is None and filas["Dora"]["categoria"] is None
    assert filas["Dora"]["fecha_respuesta"] is None
    assert "encuesta_registro.carga_id" in str(session.executed_statements[1])


def test_detail_unknown_carga_is_404():
    response, _ = _get(f"{BASE}/{CARGA_ID}/detalle", [])
    assert response.status_code == 404


@pytest.mark.parametrize("filtro,esperados", [
    ("todas", {"Ana", "Beto", "Carlos", "Dora"}),
    ("respondidas", {"Ana", "Beto", "Carlos"}),
    ("sin_responder", {"Dora"}),
    ("detractores", {"Beto", "Carlos"}),
])
def test_detail_filters(filtro, esperados):
    response, _ = _get(f"{BASE}/{CARGA_ID}/detalle?filtro={filtro}", ROWS)
    assert {f["cliente"] for f in response.json()} == esperados


def test_detail_rejects_unknown_filter():
    response, _ = _get(f"{BASE}/{CARGA_ID}/detalle?filtro=otra", ROWS)
    assert response.status_code == 422


# --- roles -----------------------------------------------------------------

@pytest.mark.parametrize("role", ["ADMIN", "SERVICIO_CLIENTE"])
def test_allowed_roles(role):
    override_motored_user(MotoredUser(user_id=USER_ID, role=role))
    response, _ = _get(f"{BASE}/{CARGA_ID}/detalle", ROWS)
    assert response.status_code == 200


@pytest.mark.parametrize("role", ["COMPRAS", "CONSULTA", "SUCURSAL"])
def test_other_roles_are_forbidden(role):
    override_motored_user(MotoredUser(user_id=USER_ID, role=role))
    override_motored_db(FakeAsyncSession(execute_queue=[[], [], [], []]))
    with TestClient(app) as client:
        assert client.get(f"{BASE}/{CARGA_ID}/detalle").status_code == 403
        assert client.get(f"{BASE}/{CARGA_ID}/excel").status_code == 403
        assert client.get(f"{BASE}/resultados/excel?desde=2026-10-01&hasta=2026-10-31").status_code == 403


# --- Excel per carga -------------------------------------------------------

def _sheet(response):
    return openpyxl.load_workbook(io.BytesIO(response.content)).active


def _values(sheet):
    return [[c.value for c in row] for row in sheet.iter_rows()]


def _table(response):
    """(header, rows-as-dicts) of the single sheet."""
    values = _values(_sheet(response))
    i = next(i for i, r in enumerate(values) if r and r[0] == "ID encuesta")
    header = values[i]
    return header, [dict(zip(header, r)) for r in values[i + 1:] if any(v is not None for v in r)]


def _caso(row, numero=7, estado="CERRADO", resultado="RECUPERADO", asignado="Laura Gestora"):
    return SimpleNamespace(
        caso_id=uuid.uuid4(), numero=numero, respuesta_id=row.respuesta_id, estado=estado,
        resultado=resultado, asignado_nombre=asignado,
        created_at=datetime(2026, 10, 3, 15, 0, 0), updated_at=datetime(2026, 10, 5, 16, 0, 0),
        cerrado_at=datetime(2026, 10, 5, 16, 0, 0) if estado == "CERRADO" else None,
    )


def _accion(caso, tipo, desc, at, usuario="Laura Gestora", ant=None, nuevo=None):
    return SimpleNamespace(
        accion_id=uuid.uuid4(), caso_id=caso.caso_id, usuario_nombre=usuario, tipo=tipo,
        descripcion=desc, estado_anterior=ant, estado_nuevo=nuevo, created_at=at,
    )


QUESTIONS = [
    "P1. Satisfacción general (1-5)",
    "P2.1 Explicación y asesoría técnica",
    "P2.2 Confianza en la reparación",
    "P2.3 Servicio en el taller",
    "P2.4 Calidad del trabajo de los mecánicos",
    "P2.5 Claridad de los cobros",
    "P2.6 Procedencia y originalidad de los repuestos",
    "P3. Observaciones",
    "P4. Autoriza uso de datos",
]


def test_single_sheet_named_encuestas_with_one_column_per_question_in_order():
    response, _ = _get(f"{BASE}/{CARGA_ID}/excel", ROWS)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(XLSX_MIME)
    assert 'filename="encuesta_mayo_2026-09-30.xlsx"' in response.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Encuestas"]
    header, _rows = _table(response)
    first = header.index(QUESTIONS[0])
    assert header[first:first + 9] == QUESTIONS
    assert header[:first][:6] == ["ID encuesta", "Cliente", "Cédula", "Teléfono", "Tienda", "Placa"]
    for name in ("Archivo de carga", "Fecha de envío", "Fecha de respuesta", "Estado", "Categoría"):
        assert name in header
    assert header[header.index("Categoría"):] == [
        "Categoría", "N.º caso", "ID caso", "Estado del caso", "Responsable", "Fecha de apertura",
        "Fecha de cierre", "Resultado", "Última actualización del caso", "N.º acciones",
        "Última acción (fecha)", "Última acción (usuario)", "Historial de gestión"]


def test_answers_text_cells_and_bogota_dates():
    response, _ = _get(f"{BASE}/{CARGA_ID}/excel", ROWS)
    _, rows = _table(response)
    ana = next(r for r in rows if r["Cliente"] == "Ana")
    assert ana["Cédula"] == "0012345" and ana["Teléfono"] == "3001112233"
    assert ana["Estado"] == "Respondida" and ana["Categoría"] == "Satisfecho"
    assert [ana[q] for q in QUESTIONS] == [
        5, 4, 5, "NS/NR", 3, 2, 1, "Excelente", "Sí"]
    assert ana["Archivo de carga"] == "mayo.xlsx"
    # real Excel dates in Bogota time (UTC-5)
    assert ana["Fecha de envío"] == datetime(2026, 9, 30, 22, 0)
    assert ana["Fecha de respuesta"] == datetime(2026, 10, 2, 10, 0)
    ws = _sheet(response)
    cells = {c.value: c for row in ws.iter_rows() for c in row if isinstance(c.value, str)}
    assert cells["0012345"].data_type == "s" and cells["0012345"].number_format == "@"
    assert cells["3001112233"].data_type == "s" and cells["3001112233"].number_format == "@"
    date_cell = next(c for row in ws.iter_rows() for c in row if c.value == datetime(2026, 9, 30, 22, 0))
    assert "yyyy" in date_cell.number_format
    dora = next(r for r in rows if r["Cliente"] == "Dora")
    assert dora["Estado"] == "Sin responder" and dora["Fecha de respuesta"] is None
    assert all(dora[q] is None for q in QUESTIONS)
    assert dora["N.º caso"] is None and dora["Historial de gestión"] is None


def test_detractor_case_and_action_summary_on_the_same_row():
    beto = ROWS[1]
    caso = _caso(beto)
    acciones = [
        _accion(caso, "APERTURA", "Caso abierto", datetime(2026, 10, 3, 15, 0, 0), usuario=None),
        _accion(caso, "LLAMADA", "Cliente contesta y acepta visita", datetime(2026, 10, 4, 14, 30, 0)),
        _accion(caso, "CAMBIO_ESTADO", "Cierre", datetime(2026, 10, 5, 16, 0, 0),
                ant="EN_GESTION", nuevo="CERRADO"),
    ]
    response, _ = _get(f"{BASE}/{CARGA_ID}/excel", ROWS, [caso], acciones)
    header, rows = _table(response)
    beto_row = next(r for r in rows if r["Cliente"] == "Beto")
    assert beto_row["Categoría"] == "Detractor"
    assert beto_row["N.º caso"] == 7 and beto_row["ID caso"] == str(caso.caso_id)
    assert beto_row["Estado del caso"] == "CERRADO" and beto_row["Resultado"] == "RECUPERADO"
    assert beto_row["Responsable"] == "Laura Gestora"
    assert beto_row["Fecha de apertura"] == datetime(2026, 10, 3, 10, 0)
    assert beto_row["Fecha de cierre"] == datetime(2026, 10, 5, 11, 0)
    assert beto_row["N.º acciones"] == 3
    assert beto_row["Última acción (fecha)"] == datetime(2026, 10, 5, 11, 0)
    assert beto_row["Última acción (usuario)"] == "Laura Gestora"
    lines = beto_row["Historial de gestión"].split("\n")
    assert lines == [
        "03/10/2026 10:00 · Sistema · APERTURA: Caso abierto",
        "04/10/2026 09:30 · Laura Gestora · LLAMADA: Cliente contesta y acepta visita",
        "05/10/2026 11:00 · Laura Gestora · CAMBIO_ESTADO: Cierre (EN_GESTION → CERRADO)",
    ]
    ws = _sheet(response)
    col = header.index("Historial de gestión") + 1
    historial = [c for row in ws.iter_rows() for c in row
                 if c.column == col and isinstance(c.value, str) and "\n" in c.value]
    assert historial and historial[0].alignment.wrap_text
    assert ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width >= 50
    # non-detractors have empty case columns
    ana = next(r for r in rows if r["Cliente"] == "Ana")
    assert ana["ID caso"] is None and ana["N.º acciones"] is None


def test_case_without_actions_has_zero_actions_and_empty_history():
    caso = _caso(ROWS[1], estado="ABIERTO", resultado=None)
    response, _ = _get(f"{BASE}/{CARGA_ID}/excel", ROWS, [caso], [])
    _, rows = _table(response)
    beto = next(r for r in rows if r["Cliente"] == "Beto")
    assert beto["N.º acciones"] == 0 and beto["Historial de gestión"] is None
    assert beto["Fecha de cierre"] is None and beto["Resultado"] is None


def test_header_row_is_frozen_below_the_header():
    response, _ = _get(f"{RANGO}?desde=2026-09-01&hasta=2026-09-30", ROWS)
    values = _values(_sheet(response))
    i = next(i for i, r in enumerate(values) if r and r[0] == "ID encuesta")
    assert _sheet(response).freeze_panes == f"A{i + 2}"


def test_excel_per_carga_filename_is_sanitized():
    rows = [_row(archivo="Base Mayo / 2026 (final).xlsx")]
    response, _ = _get(f"{BASE}/{CARGA_ID}/excel", rows)
    assert 'filename="encuesta_Base_Mayo_2026_final_2026-09-30.xlsx"' in response.headers["content-disposition"]


def test_excel_per_carga_unknown_is_404():
    response, _ = _get(f"{BASE}/{CARGA_ID}/excel", [])
    assert response.status_code == 404


# --- Excel by range --------------------------------------------------------

RANGO = f"{BASE}/resultados/excel"


def test_range_excel_has_headers_and_filename():
    response, session = _get(f"{RANGO}?desde=2026-09-01&hasta=2026-09-30", ROWS)
    assert response.status_code == 200
    assert 'filename="encuestas_envio_2026-09-01_2026-09-30.xlsx"' in response.headers["content-disposition"]
    values = _values(_sheet(response))
    flat = [v for r in values for v in r if v is not None]
    assert any("2026-09-01" in str(v) and "2026-09-30" in str(v) for v in flat)
    assert any("fecha de envío" in str(v).lower() for v in flat)
    assert any(r and r[0] == "ID encuesta" for r in values)
    assert sum(1 for r in values if len(r) > 1 and r[1] in {"Ana", "Beto", "Carlos", "Dora"}) == 4
    assert "encuesta_carga.created_at" in str(session.executed_statements[1])


def test_range_bounds_are_inclusive_in_bogota_time():
    desde, hasta = date(2026, 9, 1), date(2026, 9, 30)
    inicio, fin = encuesta_resultados.limites_utc(desde, hasta)
    # Bogota midnight is 05:00 UTC; the whole last day is included, the next day is not.
    assert inicio == datetime(2026, 9, 1, 5, 0, 0)
    assert fin == datetime(2026, 10, 1, 5, 0, 0)
    assert inicio <= datetime(2026, 9, 1, 5, 0, 0) < fin
    assert inicio <= datetime(2026, 10, 1, 4, 59, 59) < fin
    assert not (inicio <= datetime(2026, 10, 1, 5, 0, 0) < fin)


def test_range_query_uses_those_bounds():
    _, session = _get(f"{RANGO}?desde=2026-09-01&hasta=2026-09-30", ROWS)
    params = session.executed_statements[1].compile().params
    assert datetime(2026, 9, 1, 5, 0, 0) in params.values()
    assert datetime(2026, 10, 1, 5, 0, 0) in params.values()


def test_range_rejects_desde_after_hasta():
    response, _ = _get(f"{RANGO}?desde=2026-10-02&hasta=2026-10-01", [])
    assert response.status_code == 422


def test_range_rejects_more_than_366_days_but_accepts_exactly_366():
    response, _ = _get(f"{RANGO}?desde=2026-01-01&hasta=2027-01-02", [])
    assert response.status_code == 422
    response, _ = _get(f"{RANGO}?desde=2026-01-01&hasta=2027-01-01", [])
    assert response.status_code == 200


def test_range_requires_both_dates():
    response, _ = _get(f"{RANGO}?desde=2026-10-01", [])
    assert response.status_code == 422


def test_range_with_no_rows_still_returns_a_workbook():
    response, _ = _get(f"{RANGO}?desde=2026-09-01&hasta=2026-09-30", [])
    assert response.status_code == 200
    assert any(r and r[0] == "ID encuesta" for r in _values(_sheet(response)))


# --- range filtered by response date ---------------------------------------

def test_range_by_response_date_filename_header_and_query():
    response, session = _get(f"{RANGO}?desde=2026-10-01&hasta=2026-10-31&por=respuesta", ROWS[:3])
    assert response.status_code == 200
    assert 'filename="encuestas_respuesta_2026-10-01_2026-10-31.xlsx"' in response.headers["content-disposition"]
    flat = [v for r in _values(_sheet(response)) for v in r if v is not None]
    assert any("fecha de respuesta" in str(v).lower() for v in flat)
    statement = session.executed_statements[1]
    sql = str(statement)
    assert "encuesta_respuesta.created_at >=" in sql and "encuesta_carga.created_at >=" not in sql
    params = statement.compile().params
    assert datetime(2026, 10, 1, 5, 0, 0) in params.values()
    assert datetime(2026, 11, 1, 5, 0, 0) in params.values()


def test_range_by_send_date_filters_on_the_carga_not_the_response():
    _, session = _get(f"{RANGO}?desde=2026-09-01&hasta=2026-09-30&por=envio", ROWS)
    sql = str(session.executed_statements[1])
    assert "encuesta_carga.created_at >=" in sql and "encuesta_respuesta.created_at >=" not in sql


def test_range_rejects_unknown_por():
    response, _ = _get(f"{RANGO}?desde=2026-09-01&hasta=2026-09-30&por=otra", [])
    assert response.status_code == 422
