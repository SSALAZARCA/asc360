"""
KPI's, Comisiones: the Excel of the settled month (`GET /kpis/comisiones/excel`). The workbook is built
from the same payload as the tab, so these tests read it back with openpyxl.
"""
import io
import uuid

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import settings
from app.main import app
from app.motored.services import tablero_asesores as t
from app.motored.services import tablero_asesores_consultas as consultas
from app.motored.services import tablero_comisiones_excel as excel
from app.motored.services import tablero_kpis as kpis
from app.motored.services.auth import MotoredUser
from tests.motored.conftest import FakeAsyncSession, override_motored_db, override_motored_user

BASE = "/api/motored/tablero-asesores/kpis/comisiones/excel"
CABECERA = [
    "Cédula", "Asesor", "Tienda", "Cargo", "Presupuesto", "Venta con HMCL (base de cumplimiento)",
    "Cumplimiento %", "Tramo", "% comisión", "Venta sin HMCL (base de pago)", "Comisión",
    "Bono Cascos (≥ 6%)", "Bono Tecnired (clientes) (≥ 6,5%)", "Bono total", "Total a pagar"]
LINEAS = [
    {"linea": "CASCOS", "etiqueta": "Cascos", "pct_meta": 6.0, "bono": 30000, "activo": True},
    {"linea": "TECNIRED", "etiqueta": "Tecnired (clientes)", "pct_meta": 6.5, "bono": 25000, "activo": False}]


def _bonos(cascos, tecnired, paga_cascos):
    return [
        {"linea": "CASCOS", "etiqueta": "Cascos", "venta": 1.0, "pct_real": 0.1, "pct_meta": 6.0, "bono": 30000,
         "cumple": cascos, "paga": paga_cascos, "activo": True, "bono_pagado": 30000 if paga_cascos else 0},
        {"linea": "TECNIRED", "etiqueta": "Tecnired (clientes)", "venta": 1.0, "pct_real": 0.1, "pct_meta": 6.5,
         "bono": 25000, "cumple": tecnired, "paga": False, "activo": False, "bono_pagado": 0}]


def _asesor(cedula, nombre, presupuesto, vcump, vcom, tramo, tasa, comision, bonos=None):
    bonos = bonos or _bonos(False, False, False)
    bono_total = sum(b["bono_pagado"] for b in bonos)
    return {
        "bonos": bonos, "bono_total": bono_total, "total_a_pagar": comision + bono_total,
        "gate": {"umbral": 95.0, "cumple": vcump / presupuesto >= 0.95},
        "cedula": cedula, "nombre": nombre, "tienda": "Norte", "sucursal_id": "s1", "cargo": "ASESOR DE REPUESTOS",
        "presupuesto": presupuesto, "venta_cumplimiento": vcump, "venta_comision": vcom,
        "cumplimiento_pct": vcump / presupuesto, "tramo": tramo, "tasa_pct": tasa, "comision": comision, "sig": None}


DATOS = {
    "meses": ["2026-06", "2026-07"], "mes_liquidado": "2026-07", "sucursales": [],
    "reglas": {
        "cumplimiento_base": "con_hmcl", "comision_base_pago": "sin_hmcl",
        "comision_tramos": [{"nombre": "BASE", "desde_pct": 0, "tasa_pct": 1.0}, {"nombre": "PRO", "desde_pct": 90, "tasa_pct": 1.5}],
        "comision_cargos_asesor": ["ASESOR DE REPUESTOS"], "comision_lineas": LINEAS,
        "comision_bono_umbral_pct": 95.0},
    "asesores": [
        _asesor("0012345678", "Ana", 1000000, 950000, 900000, "PRO", 1.5, 13500.0, _bonos(True, True, True)),
        _asesor("987", "Beto", 500000, 100000, 90000, "BASE", 1.0, 900.0)],
    "advertencias": {"sin_presupuesto": [{"cedula": "0555", "nombre": "Eli", "venta": 50000.0}]},
}


def _libro(datos=DATOS, tiendas=()):
    return load_workbook(io.BytesIO(excel.construir_libro(datos, list(tiendas))))


def _filas(hoja):
    return [[c.value for c in fila] for fila in hoja.iter_rows()]


def _tabla(hoja):
    filas = _filas(hoja)
    inicio = next(i for i, f in enumerate(filas) if f[0] == "Cédula")
    return inicio, filas[inicio:]


def test_the_columns_are_the_requested_ones_in_order():
    _, tabla = _tabla(_libro()["Comisiones"])

    assert tabla[0] == CABECERA


def test_the_cedula_is_text_with_its_leading_zeros():
    hoja = _libro()["Comisiones"]
    inicio, tabla = _tabla(hoja)

    assert [f[0] for f in tabla[1:3]] == ["0012345678", "987"]
    celda = hoja.cell(row=inicio + 2, column=1)
    assert celda.data_type == "s" and celda.number_format == "@"


def test_amounts_are_numbers_and_percentages_are_percent_formatted():
    hoja = _libro()["Comisiones"]
    inicio, tabla = _tabla(hoja)

    ana = tabla[1]
    assert ana[1:11] == [
        "Ana", "Norte", "ASESOR DE REPUESTOS", 1000000, 950000, pytest.approx(0.95), "PRO", pytest.approx(0.015),
        900000, 13500.0]
    fila = inicio + 2
    for columna in (5, 6, 10, 11):
        assert hoja.cell(row=fila, column=columna).number_format == "#,##0"
    for columna in (7, 9):
        assert hoja.cell(row=fila, column=columna).number_format == "0.0%"


def test_the_footer_totals_the_money_columns():
    _, tabla = _tabla(_libro()["Comisiones"])

    total = tabla[-1]
    assert total[0] == "Total" and total[4:7] == [1500000, 1050000, None] and total[9:11] == [990000, 14400.0]
    assert total[11:] == [30000, 0, 30000, 44400.0]


def test_each_line_column_shows_met_or_not_with_the_bonus_and_the_totals_follow():
    hoja = _libro()["Comisiones"]
    inicio, tabla = _tabla(hoja)

    ana, beto = tabla[1], tabla[2]
    assert ana[11:] == ["✓ $30.000", "✓ no se paga (línea apagada)", 30000, 43500.0]
    assert beto[11:] == ["✗", "✗", 0, 900.0]
    for columna in (14, 15):
        assert hoja.cell(row=inicio + 2, column=columna).number_format == "#,##0"


def test_a_met_line_behind_the_gate_is_marked_as_not_paid():
    activa = [{**LINEAS[0]}]
    bonos = [{**_bonos(True, False, False)[0]}]
    datos = {**DATOS, "reglas": {**DATOS["reglas"], "comision_lineas": activa},
             "asesores": [_asesor("1", "Zoe", 1000000, 500000, 500000, "BASE", 1.0, 5000.0, bonos)]}

    _, tabla = _tabla(_libro(datos)["Comisiones"])

    assert tabla[1][11:] == ["✓ no se paga (sin cumplir el 95%)", 0, 5000.0]


def test_the_rules_rows_say_the_threshold_and_the_bonus_lines():
    texto = "\n".join(str(v) for f in _filas(_libro()["Comisiones"]) for v in f if v is not None)

    assert "Bonos por línea" in texto and "95%" in texto and "Cascos" in texto


def test_the_header_rows_say_the_month_the_stores_and_the_rules():
    plano = [str(v) for f in _filas(_libro(tiendas=["Norte", "Sur"])["Comisiones"]) for v in f if v is not None]
    texto = "\n".join(plano)

    assert "2026-07" in texto and "Norte, Sur" in texto
    assert "BASE desde 0% · 1%" in texto and "PRO desde 90% · 1,5%" in texto
    assert "Todas las tiendas" in "\n".join(str(v) for f in _filas(_libro()["Comisiones"]) for v in f if v)


def test_the_second_sheet_lists_the_asesores_without_budget_with_text_cedula():
    hoja = _libro()["Sin presupuesto"]

    assert _filas(hoja) == [["Cédula", "Nombre", "Venta"], ["0555", "Eli", 50000.0]]
    assert hoja["A2"].data_type == "s" and hoja["A2"].number_format == "@" and hoja["C2"].number_format == "#,##0"


def test_the_filename_carries_the_month():
    assert excel.nombre_de_archivo(DATOS) == "comisiones_2026-07.xlsx"


# --- HTTP ---------------------------------------------------------------------------------------


@pytest.fixture
def _ready(monkeypatch):
    monkeypatch.setattr(settings, "MOTORED_ENABLED", True)
    monkeypatch.setattr(settings, "MOTORED_SECRET_KEY", "kpis-test-motored-secret")
    monkeypatch.setattr(settings, "SECRET_KEY", "kpis-test-asc360-secret")

    async def filtro(db, meses, modo, ids=None):
        return t.filtro_de_meses(t.validar_meses(meses), modo, ids)

    async def calculo(db, filtro_):
        return {**DATOS, "sucursales": []}

    async def sucursales(db, ids):
        return {str(i): (f"Tienda {i}", None) for i in ids}

    monkeypatch.setattr(consultas, "cargar_filtro", filtro)
    monkeypatch.setattr(kpis, "calcular_kpis_comisiones", calculo)
    monkeypatch.setattr(consultas, "consultar_sucursales", sucursales)
    yield
    app.dependency_overrides.clear()


def _get(rol, **params):
    override_motored_user(MotoredUser(user_id=str(uuid.uuid4()), role=rol))
    override_motored_db(FakeAsyncSession(execute_queue=[[]]))
    with TestClient(app) as client:
        return client.get(BASE, params=params)


@pytest.mark.parametrize("rol", ["ADMIN", "COMPRAS", "GERENCIA"])
def test_allowed_roles_download_the_workbook(_ready, rol):
    r = _get(rol, meses="2026-06,2026-07")

    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert 'filename="comisiones_2026-07.xlsx"' in r.headers["content-disposition"]
    assert "Comisiones" in load_workbook(io.BytesIO(r.content)).sheetnames


@pytest.mark.parametrize("rol", ["CONSULTA", "SUCURSAL", "SERVICIO_CLIENTE"])
def test_other_roles_are_forbidden(_ready, rol):
    assert _get(rol, meses="2026-07").status_code == 403


def test_the_filter_is_validated_like_the_tab(_ready):
    assert _get("ADMIN").status_code == 422
