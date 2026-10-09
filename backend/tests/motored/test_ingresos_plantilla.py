"""ERP "Entradas x Compra" template (odd/tasks/motored-ingresos-responsable-plantilla.md, T2).

The workbook is built by a pure function and the endpoint is exercised with
the data service stubbed (the SQL is proven in the pg_real file).
"""
import io
import uuid
from datetime import date
from decimal import Decimal as D
from unittest.mock import AsyncMock

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import settings
from app.main import app
from app.motored.services import ingresos_pendientes as ip
from app.motored.services import ingresos_plantilla as pl
from app.motored.services import sucursal_grupo
from app.motored.services.auth import MotoredUser
from tests.motored.conftest import (
    FakeAsyncSession, override_motored_db, override_motored_user,
)

BASE = "/api/motored/gestion-repuestos/ingresos-facturas/plantilla"
TIENDA = uuid.uuid4()
AJUSTES = {
    "tipo_documento": 16, "descuento_global": 5, "proveedor": 900723988,
    "sucursal_proveedor": "001", "comprador": 1151943311,
    "descuento_item": 0, "unidad_negocio": "003",
}


def _datos(lineas=None, **extra):
    base = dict(
        prefijo_rh="RH", numero_rh=208629, codigo_co="F02",
        bodega="BF021", fecha=date(2026, 10, 9), ajustes=AJUSTES,
        lineas=lineas if lineas is not None else [
            pl.LineaPlantilla("11393-KWK-900S", D(10), D("3962")),
            pl.LineaPlantilla("15439-AAH-00099S", D(10), D("1893")),
            pl.LineaPlantilla("61330-KVN-900S", D(1), D("2789.5")),
        ])
    base.update(extra)
    return pl.DatosPlantilla(**base)


def _hoja(datos):
    libro = load_workbook(io.BytesIO(pl.construir_libro(datos)))
    assert libro.sheetnames == ["Entrada Compra"]
    return libro["Entrada Compra"]


def test_header_block_matches_the_template():
    ws = _hoja(_datos())

    assert [ws[c].value for c in ("A1", "B1", "C1", "E1", "H1", "I1")] == [
        "Centro de Operación", "F02", "B", "AAAAMMDD", "MODO", 2]
    assert ws["D1"].value.date() == date(2026, 10, 9)
    assert ws["D1"].is_date
    assert [ws[c].value for c in ("A2", "B2", "C2", "D2")] == [
        "Tipo Documento:", 16, "Descuento Global", 5]
    assert [ws[c].value for c in ("A3", "B3", "C3", "D3")] == [
        "Proveedor", 900723988, "Suc. Proveedor", "001"]
    assert [ws[c].value for c in ("A4", "B4", "C4", "D4")] == [
        "Comprador", 1151943311, "Descuento x Item", 0]
    assert [ws[c].value for c in ("A5", "B5", "C5", "D5")] == [
        "Notas Documento", "ENTRADA POR COMPRA RH208629",
        "Doc. Referencia", "RH208629"]
    assert all(c.value is None for c in ws[6])


def test_table_headers_are_on_row_7():
    ws = _hoja(_datos())

    assert [c.value for c in ws[7]][:10] == [
        "Consecutivo Entrada", "Referencia", "Bodega", "Precio Unit",
        "Cantidad", "Vr Bruto", "Dcto Item", "Notas del Mvto",
        "Centro Operación", "U.N."]


def test_one_row_per_line_with_formulas_and_whole_numbers_as_ints():
    ws = _hoja(_datos())

    assert [c.value for c in ws[8]][:10] == [
        1, "11393-KWK-900S", "BF021", 3962, 10, "=D8*E8", "=+F8*$D$4%",
        "Mvto 1", "=+B1", "003"]
    assert [c.value for c in ws[9]][:10] == [
        1, "15439-AAH-00099S", "BF021", 1893, 10, "=D9*E9", "=+F9*$D$4%",
        "Mvto 2", "F02", "003"]
    assert ws["D10"].value == 2789.5 and ws["E10"].value == 1
    assert isinstance(ws["E10"].value, int)
    assert ws["I10"].value == "F02" and ws["H10"].value == "Mvto 3"
    assert ws.max_row == 10


def test_fixed_values_come_from_the_settings():
    ajustes = {**AJUSTES, "tipo_documento": 18, "descuento_global": 7.5,
               "sucursal_proveedor": "002", "unidad_negocio": "007",
               "descuento_item": 2}
    ws = _hoja(_datos(ajustes=ajustes))

    assert (ws["B2"].value, ws["D2"].value, ws["D3"].value,
            ws["D4"].value, ws["J8"].value) == (18, 7.5, "002", 2, "007")


def test_a_price_missing_from_the_invoice_falls_back_to_total_over_quantity():
    ws = _hoja(_datos(lineas=[
        pl.LineaPlantilla("A", D(3), None, valor_total=D(10)),
        pl.LineaPlantilla("B", D(4), None, valor_total=D(1000)),
    ]))

    assert ws["D8"].value == 3.33
    assert ws["D9"].value == 250 and isinstance(ws["D9"].value, int)


def test_a_line_without_any_price_leaves_the_cell_empty():
    ws = _hoja(_datos(lineas=[pl.LineaPlantilla("A", D(3), None)]))

    assert ws["D8"].value is None and ws["E8"].value == 3


def test_column_widths_and_number_formats_follow_the_template():
    ws = _hoja(_datos())

    assert ws.column_dimensions["B"].width == pytest.approx(41.29, abs=0.1)
    assert ws["D1"].number_format == "mm-dd-yy"
    assert ws["F8"].number_format == "#,##0.00"
    assert ws["A1"].font.b and ws["A7"].font.b and not ws["B1"].font.b


# --- endpoint ---------------------------------------------------------------


@pytest.fixture(autouse=True)
def _listo(monkeypatch):
    monkeypatch.setattr(settings, "MOTORED_ENABLED", True)
    monkeypatch.setattr(settings, "MOTORED_SECRET_KEY", "plantilla-motored")
    monkeypatch.setattr(settings, "SECRET_KEY", "plantilla-asc360")
    yield
    app.dependency_overrides.clear()


def _como(rol):
    override_motored_user(MotoredUser(user_id=str(uuid.uuid4()), role=rol))
    override_motored_db(FakeAsyncSession(execute_queue=[[]] * 6))


def _pedir(**params):
    params.setdefault("factura", "RH 208629")
    params.setdefault("sucursal", str(TIENDA))
    with TestClient(app) as client:
        return client.get(BASE, params=params)


@pytest.mark.parametrize("rol", ["ADMIN", "ANALISTA_ADMINISTRATIVO"])
def test_admin_and_analista_download_the_workbook(monkeypatch, rol):
    monkeypatch.setattr(pl, "preparar", AsyncMock(return_value=_datos()))
    _como(rol)

    r = _pedir()

    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert 'filename="Entrada_compra_RH208629.xlsx"' in (
        r.headers["content-disposition"])
    ws = load_workbook(io.BytesIO(r.content))["Entrada Compra"]
    assert ws["B5"].value == "ENTRADA POR COMPRA RH208629"


@pytest.mark.parametrize("rol", [
    "COORDINADOR_REPUESTOS", "GERENCIA", "COMPRAS", "SERVICIO_CLIENTE",
    "SUCURSAL", "LIDER_INVENTARIOS"])
def test_every_other_role_is_forbidden(monkeypatch, rol):
    preparar = AsyncMock(return_value=_datos())
    monkeypatch.setattr(pl, "preparar", preparar)
    _como(rol)

    assert _pedir().status_code == 403
    preparar.assert_not_awaited()


@pytest.mark.parametrize("status, mensaje", [
    (404, "La factura no está pendiente."),
    (409, "Esta factura la ingresa el asesor de la tienda."),
    (409, "Confirma primero que la factura llegó."),
    (409, "La tienda no tiene bodega principal."),
])
def test_service_errors_become_http_errors(monkeypatch, status, mensaje):
    monkeypatch.setattr(pl, "preparar", AsyncMock(
        side_effect=pl.PlantillaError(status, mensaje)))
    _como("ANALISTA_ADMINISTRATIVO")

    r = _pedir()

    assert r.status_code == status and r.json()["detail"] == mensaje


def test_a_malformed_invoice_is_a_422(monkeypatch):
    monkeypatch.setattr(pl, "preparar", AsyncMock(return_value=_datos()))
    _como("ADMIN")

    r = _pedir(factura="xyz")

    assert r.status_code == 422 and r.json()["detail"] == ip.MSG_FACTURA


# --- the data service (queries stubbed one level down) ----------------------


def _item(**extra):
    base = {"prefijo_rh": "RH", "numero_rh": 208629, "sucursal_id": TIENDA,
            "responsable": "ANALISTA", "estado": "LLEGO",
            "puede_descargar_plantilla": True}
    base.update(extra)
    return base


@pytest.fixture
def mundo(monkeypatch):
    estado = {
        "items": [_item()],
        "sucursal": type("S", (), {
            "codigo_co": "F02", "bodega_principal": "BF021"})(),
        "lineas": [pl.LineaPlantilla("A", D(1), D(5))],
    }
    monkeypatch.setattr(
        ip, "pendientes", AsyncMock(side_effect=lambda *a, **k: estado["items"]))
    monkeypatch.setattr(
        sucursal_grupo, "principal_de",
        AsyncMock(return_value={TIENDA: TIENDA}))
    monkeypatch.setattr(
        pl, "_sucursal", AsyncMock(side_effect=lambda *a: estado["sucursal"]))
    monkeypatch.setattr(
        pl, "_lineas", AsyncMock(side_effect=lambda *a: estado["lineas"]))
    monkeypatch.setattr(
        pl, "_ajustes", AsyncMock(return_value=dict(AJUSTES)))
    return estado


async def _preparar():
    return await pl.preparar(object(), ("RH", 208629), TIENDA, date(2026, 10, 9))


async def test_prepares_the_data_of_a_downloadable_invoice(mundo):
    datos = await _preparar()

    assert (datos.codigo_co, datos.bodega, datos.numero_rh) == (
        "F02", "BF021", 208629)
    assert datos.fecha == date(2026, 10, 9) and datos.ajustes == AJUSTES
    assert [l.referencia for l in datos.lineas] == ["A"]


async def test_unknown_or_already_entered_invoice_is_404(mundo):
    mundo["items"] = [_item(numero_rh=1)]

    with pytest.raises(pl.PlantillaError) as e:
        await _preparar()
    assert e.value.status_code == 404


async def test_asesor_invoice_is_409(mundo):
    mundo["items"] = [_item(responsable="ASESOR")]

    with pytest.raises(pl.PlantillaError) as e:
        await _preparar()
    assert e.value.status_code == 409 and "asesor" in e.value.detail


@pytest.mark.parametrize("estado", ["SIN_CONFIRMAR", "NO_HA_LLEGADO"])
async def test_invoice_not_confirmed_as_arrived_is_409(mundo, estado):
    mundo["items"] = [_item(estado=estado)]

    with pytest.raises(pl.PlantillaError) as e:
        await _preparar()
    assert e.value.status_code == 409 and "llegó" in e.value.detail


@pytest.mark.parametrize("bodega", [None, "", "   "])
async def test_store_without_a_main_warehouse_is_409(mundo, bodega):
    mundo["sucursal"].bodega_principal = bodega

    with pytest.raises(pl.PlantillaError) as e:
        await _preparar()
    assert e.value.status_code == 409
    assert "Configuración" in e.value.detail and "Tiendas" in e.value.detail


async def test_invoice_without_positive_lines_is_409(mundo):
    mundo["lineas"] = []

    with pytest.raises(pl.PlantillaError) as e:
        await _preparar()
    assert e.value.status_code == 409
