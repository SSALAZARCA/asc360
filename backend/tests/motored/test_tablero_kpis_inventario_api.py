"""
KPI's, Inventario: the HTTP layer (`GET /kpis/inventario` and `/kpis/inventario/excel`). The service is replaced
by a double to check roles, parameter parsing, the Spanish 422s, the passthrough and the file headers; the real
queries run in `pg_real/test_tablero_kpis_inventario_pg.py`.
"""
import io
import uuid

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import settings
from app.main import app
from app.motored.services import tablero_asesores as t
from app.motored.services import tablero_asesores_consultas as consultas
from app.motored.services import tablero_kpis as kpis
from app.motored.services.auth import MotoredUser
from tests.motored.conftest import FakeAsyncSession, override_motored_db, override_motored_user

BASE = "/api/motored/tablero-asesores/kpis/inventario"
UUID_A, UUID_B = "11111111-1111-1111-1111-111111111111", "22222222-2222-2222-2222-222222222222"
DATOS = {
    "meses": ["2026-10"], "hmcl": "incluir", "sucursales": [UUID_A], "corte": "2026-10-07",
    "costo_desde": "2026-08-01", "costo_hasta": "2026-10-31", "sin_movimiento_umbral_dias": 180,
    "tarjetas": {"valor": 100.5}, "tendencia": [], "lineas": [], "tiendas": [],
    "sin_movimiento_top": [], "agotadas": {"total": 0, "en_transito": 0, "sin_pedir": 0, "items": []}}


@pytest.fixture
def llamadas(monkeypatch):
    monkeypatch.setattr(settings, "MOTORED_ENABLED", True)
    monkeypatch.setattr(settings, "MOTORED_SECRET_KEY", "inv-test-motored-secret")
    monkeypatch.setattr(settings, "SECRET_KEY", "inv-test-asc360-secret")
    registro = []

    async def falso_filtro(db, meses, modo_hmcl, sucursal_ids=None):
        registro.append(("filtro", meses, modo_hmcl, sucursal_ids))
        return t.filtro_de_meses(t.validar_meses(meses), modo_hmcl, sucursal_ids)

    async def falso_calculo(db, filtro):
        registro.append(("inventario", filtro.meses, filtro.modo_hmcl, filtro.sucursal_ids))
        return {**DATOS, "monto": 12.5}

    async def falso_completo(db, filtro):
        registro.append(("excel", filtro.meses, filtro.modo_hmcl, filtro.sucursal_ids))
        return DATOS, [{"referencia": "R2", "nombre": "N", "tienda": "Norte", "existencia": 4, "valor": 200.0,
                        "dias_sin_venta": 190}], [{"referencia": "R6", "nombre": None, "tienda": "Sur", "vendidas": 20,
                                                   "perdidas": 6, "demanda": 26, "transito": 70, "cobertura": 1.0}]

    async def falsas_tiendas(db, ids):
        registro.append(("tiendas", list(ids)))
        return {UUID_A: ("Norte", None)}

    monkeypatch.setattr(consultas, "cargar_filtro", falso_filtro)
    monkeypatch.setattr(consultas, "consultar_sucursales", falsas_tiendas)
    monkeypatch.setattr(kpis, "calcular_kpis_inventario", falso_calculo)
    monkeypatch.setattr(kpis, "calcular_kpis_inventario_completo", falso_completo)
    yield registro
    app.dependency_overrides.clear()


def _como(rol):
    override_motored_user(MotoredUser(user_id=str(uuid.uuid4()), role=rol))
    override_motored_db(FakeAsyncSession(execute_queue=[[]]))


def _get(ruta="", **params):
    with TestClient(app) as client:
        return client.get(f"{BASE}{ruta}", params=params)


@pytest.mark.parametrize("ruta", ["", "/excel"])
@pytest.mark.parametrize("rol", ["CONSULTA", "SUCURSAL", "SERVICIO_CLIENTE"])
def test_other_roles_are_forbidden(llamadas, ruta, rol):
    _como(rol)

    assert _get(ruta, meses="2026-10").status_code == 403
    assert llamadas == []


@pytest.mark.parametrize("ruta", ["", "/excel"])
def test_unauthenticated_is_401(llamadas, ruta):
    override_motored_db(FakeAsyncSession(execute_queue=[[]]))

    assert _get(ruta, meses="2026-10").status_code == 401


@pytest.mark.parametrize("rol", ["ADMIN", "COMPRAS", "GERENCIA", "COORDINADOR_REPUESTOS"])
def test_allowed_roles_get_the_service_payload(llamadas, rol):
    _como(rol)

    r = _get(meses="2026-10,2026-08")

    assert r.status_code == 200 and r.json()["monto"] == 12.5 and r.json()["corte"] == "2026-10-07"
    assert llamadas[-1] == ("inventario", ("2026-08", "2026-10"), "incluir", None)


def test_the_filter_params_reach_the_service(llamadas):
    _como("ADMIN")

    r = _get(meses="2026-10", sucursales=f"{UUID_A},{UUID_B},{UUID_A}", hmcl="solo")

    assert r.status_code == 200
    assert llamadas[0][2] == "solo" and [str(i) for i in llamadas[0][3]] == [UUID_A, UUID_B]
    assert llamadas[1][2:] == ("solo", frozenset(llamadas[0][3]))


@pytest.mark.parametrize("ruta", ["", "/excel"])
@pytest.mark.parametrize("params, motivo", [
    ({}, "meses"),
    ({"meses": "2026-13"}, "inválido"),
    ({"meses": "2026-10", "hmcl": "todos"}, "hmcl"),
    ({"meses": "2026-10", "sucursales": "no-es-uuid"}, "Sucursal"),
])
def test_invalid_params_are_422_in_spanish(llamadas, ruta, params, motivo):
    _como("COMPRAS")

    r = _get(ruta, **params)

    assert r.status_code == 422 and motivo in r.json()["detail"]
    assert not [x for x in llamadas if x[0] in ("inventario", "excel")]


def test_the_excel_downloads_the_workbook_of_the_corte(llamadas):
    _como("GERENCIA")

    r = _get("/excel", meses="2026-10", sucursales=UUID_A)

    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert 'filename="inventario_2026-10-07.xlsx"' in r.headers["content-disposition"]
    libro = load_workbook(io.BytesIO(r.content))
    assert libro.sheetnames == ["Tiendas", "Líneas", "Sin movimiento", "Agotadas"]
    filas = [[c.value for c in f] for f in libro["Sin movimiento"].iter_rows()]
    assert ["R2", "N", "Norte", 4, 200.0, 190] in filas
    assert any("Norte" == v for f in [[c.value for c in f] for f in libro["Tiendas"].iter_rows()] for v in f)
    assert ("tiendas", [UUID_A]) in llamadas


def test_the_routes_are_documented_paths():
    rutas = sorted(p for p in app.openapi()["paths"] if p.startswith(BASE))
    assert rutas == [BASE, f"{BASE}/excel"]
