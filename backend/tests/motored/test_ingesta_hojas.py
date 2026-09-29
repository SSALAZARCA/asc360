"""
Sheet selection for movement uploads: the whole production workbook (many
sheets, active sheet unrelated) must be accepted. Reading (`lector`) and
type verification (`deteccion`, used by `POST /cargas`) pick the sheet whose
header row best matches the declared type (highest match ratio, at least 60%,
first on ties), never just `workbook.active`.
Synthetic fixtures only, same header shapes as the real workbook.
"""
import io
import logging

import openpyxl
import pytest
from fastapi import HTTPException

from app.motored.api import cargas as cargas_api
from app.motored.services.ingesta import deteccion, lector

ENCABEZADOS = {
    "VENTAS": (
        "Estado", "Módulo", "Fecha", "Cantidad inv.", "Tipo inventario", "Desc.bodega",
        "Bodega", "Referencia",
    ),
    "INVENTARIO": ("Referencia", "Bodega", "Desc.bodega", "Existencia"),
    "BACKORDER": (
        "SIC", "Sucursal", "Número del pedido", "Estado del pedido", "Referencia Parte",
        "Cantidad Pendiente",
    ),
    "FACTURAS_PEDIDOS": (
        "SIIC", "Sucursal", "Nota crédito", "Factura", "Fecha", "Parte", "Cantidad",
        "Vlr. Total Neto",
    ),
    "INGRESOS_FACTURAS": ("Nrodocumento", "Fecha", "Estado", "Dct.referencia", "Valornetolocal"),
    "DEMANDA_PERDIDA": ("sucursal", "sucursal Drive", "Referencia", "Cantidad Solicitada"),
}
HOJA_POR_TIPO = {
    "VENTAS": "BD ventas ultimos 6 meses",
    "INVENTARIO": "inventario actual",
    "BACKORDER": "backorder",
    "FACTURAS_PEDIDOS": "facturas pedidos",
    "INGRESOS_FACTURAS": "ingreso facturas ultimo 45 dias",
    "DEMANDA_PERDIDA": "Ventas perdidas",
}


def _fila_de_datos(tipo):
    return tuple(f"{tipo}-dato{i}" for i in range(len(ENCABEZADOS[tipo])))


def _hoja_con_encabezado(workbook, titulo, tipo, filas_titulo=0):
    hoja = workbook.create_sheet(titulo)
    for i in range(filas_titulo):
        hoja.append([f"titulo {i}"])
    hoja.append(list(ENCABEZADOS[tipo]))
    hoja.append(list(_fila_de_datos(tipo)))
    return hoja


def _workbook_completo(activa="Lista HMCL Sep 2026") -> openpyxl.Workbook:
    """Un libro multi-hoja: la hoja activa es una lista de precios sin
    relación con ningún tipo, como en el archivo real de producción."""
    workbook = openpyxl.Workbook()
    lista = workbook.active
    lista.title = activa
    lista.append(["Código", "Precio lista", "Descripción"])
    lista.append(["A1", 100, "algo"])
    for tipo, titulo in HOJA_POR_TIPO.items():
        filas_titulo = 5 if tipo == "FACTURAS_PEDIDOS" else 0
        _hoja_con_encabezado(workbook, titulo, tipo, filas_titulo=filas_titulo)
    workbook.active = 0
    return workbook


def _bytes(workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def _leer_todo(file_bytes, columnas):
    lotes = lector.leer_lotes(file_bytes, columnas_esperadas=columnas)
    return [fila async for lote in lotes for fila in lote]


@pytest.mark.parametrize("tipo", ENCABEZADOS)
async def test_leer_lotes_lee_la_hoja_del_tipo_aunque_la_activa_sea_otra(tipo):
    filas = await _leer_todo(_bytes(_workbook_completo()), ENCABEZADOS[tipo])

    assert tuple(ENCABEZADOS[tipo]) in [tuple(f) for f in filas]
    assert _fila_de_datos(tipo) in [tuple(f) for f in filas]
    assert ("Código", "Precio lista", "Descripción") not in [tuple(f) for f in filas]


async def test_leer_lotes_sin_columnas_esperadas_conserva_la_hoja_activa():
    filas = await _leer_todo(_bytes(_workbook_completo()), None)

    assert filas[0] == ("Código", "Precio lista", "Descripción")


async def test_leer_lotes_sin_hoja_que_coincida_cae_a_la_hoja_activa():
    workbook = openpyxl.Workbook()
    workbook.active.append(["nada", "reconocible"])
    workbook.create_sheet("otra").append(["tampoco", "aqui"])

    filas = await _leer_todo(_bytes(workbook), ENCABEZADOS["INVENTARIO"])

    assert filas == [("nada", "reconocible")]


async def test_leer_lotes_con_varias_hojas_que_coinciden_toma_la_primera_y_lo_registra(caplog):
    workbook = openpyxl.Workbook()
    workbook.active.title = "activa"
    _hoja_con_encabezado(workbook, "primera", "INVENTARIO")
    segunda = _hoja_con_encabezado(workbook, "segunda", "INVENTARIO")
    segunda["A2"] = "SEGUNDA-HOJA"

    with caplog.at_level(logging.INFO, logger="motored.ingesta.lector"):
        filas = await _leer_todo(_bytes(workbook), ENCABEZADOS["INVENTARIO"])

    assert _fila_de_datos("INVENTARIO") in [tuple(f) for f in filas]
    assert not any("SEGUNDA-HOJA" in tuple(f) for f in filas)
    assert "primera" in caplog.text and "segunda" in caplog.text


@pytest.mark.parametrize("tipo", ENCABEZADOS)
def test_verificar_tipo_acepta_el_libro_completo_con_hoja_activa_ajena(tipo):
    file_bytes = _bytes(_workbook_completo())

    muestra = deteccion.extraer_filas_muestra(file_bytes, columnas_esperadas=ENCABEZADOS[tipo])

    deteccion.verificar_tipo(tipo, muestra)


def test_extraer_filas_muestra_sin_columnas_lee_la_hoja_activa_como_antes():
    muestra = deteccion.extraer_filas_muestra(_bytes(_workbook_completo()))

    assert muestra[0] == ("Código", "Precio lista", "Descripción")


def test_verificar_tipo_o_400_de_la_api_acepta_el_libro_completo():
    cargas_api._verificar_tipo_o_400("VENTAS", _bytes(_workbook_completo()))


def test_verificar_tipo_o_400_de_la_api_sigue_rechazando_si_ninguna_hoja_coincide():
    workbook = openpyxl.Workbook()
    workbook.active.append(["Columna A", "Columna B"])

    with pytest.raises(HTTPException) as excinfo:
        cargas_api._verificar_tipo_o_400("VENTAS", _bytes(workbook))

    assert excinfo.value.status_code == 400
    assert excinfo.value.detail["tipo_declarado"] == "VENTAS"
