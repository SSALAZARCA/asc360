"""
Fase 4 "Pantallas del pedido" (sdd/motored-pedidos-ui, B4, ADR-7, decisión
F4-1, spec EX-01, EX-03, EX-04, EX-05, EX-14, R4.5 y R4.6 de las tareas):
los constructores PUROS del archivo HMCL.

Nada de base de datos ni de HTTP: el libro se escribe en memoria y se
vuelve a LEER con openpyxl, así lo que se prueba es lo que verá quien lo
abra: la cabecera (tienda, SIC, fecha), las columnas `Código | Cantidad`,
los códigos como TEXTO (ceros a la izquierda; un código que empieza con `=`
nunca es una fórmula), sólo las cantidades mayores que 0 y en orden de
código. También el nombre del archivo, el zip y las cabeceras de la
respuesta.
"""
import datetime
import io
import json
import zipfile
from decimal import Decimal
from urllib.parse import unquote

from openpyxl import load_workbook

from app.motored.services.corridas import exportacion_hmcl as hmcl

D = Decimal
CORTE = datetime.date(2026, 10, 1)


def _datos(lineas, nombre="Manizales", sic="1234"):
    return hmcl.DatosTienda(nombre, sic, CORTE, lineas)


def _libro(datos):
    destino = io.BytesIO()
    hmcl.escribir_libro(datos, destino)
    return load_workbook(io.BytesIO(destino.getvalue()))


def _filas(hoja):
    return [[celda.value for celda in fila] for fila in hoja.iter_rows()]


# --- nombre del archivo (EX-01, EX-14) --------------------------------------


def test_the_file_name_follows_the_f4_1_format_ex_01():
    assert hmcl.nombre_archivo("1234", "Manizales", CORTE) == (
        "Pedido_SIC1234_Manizales_2026-10-01.xlsx")


def test_accents_go_and_spaces_become_underscores_ex_14():
    assert hmcl.nombre_archivo("77", "Medellín Poblado", CORTE) == (
        "Pedido_SIC77_Medellin_Poblado_2026-10-01.xlsx")
    assert hmcl.nombre_archivo("77", "Peña Ñandú", CORTE) == (
        "Pedido_SIC77_Pena_Nandu_2026-10-01.xlsx")


def test_the_same_name_in_nfc_and_nfd_gives_the_same_file():
    compuesto = "Medellín"
    descompuesto = "Medellín"
    assert hmcl.nombre_archivo("1", compuesto, CORTE) == (
        hmcl.nombre_archivo("1", descompuesto, CORTE))


def test_forbidden_and_odd_characters_are_dropped():
    sucia = 'A\\B:C*D?E"F<G>H|I\x00J\x07K'
    nombre = hmcl.nombre_archivo("12/3", sucia, CORTE)

    assert nombre == "Pedido_SIC123_ABCDEFGHIJK_2026-10-01.xlsx"


def test_runs_of_spaces_make_one_underscore_and_edges_are_trimmed():
    assert hmcl.nombre_archivo("1", "  Cali   Norte  ", CORTE) == (
        "Pedido_SIC1_Cali_Norte_2026-10-01.xlsx")


def test_a_name_with_nothing_usable_still_gives_a_file_name():
    assert hmcl.nombre_archivo("1", "???", CORTE) == (
        "Pedido_SIC1_Tienda_2026-10-01.xlsx")


def test_hyphens_and_digits_survive():
    assert hmcl.nombre_archivo("A-9", "Zona 2-B", CORTE) == (
        "Pedido_SICA-9_Zona_2-B_2026-10-01.xlsx")


def test_the_zip_name_carries_the_corrida_code_and_the_corte():
    assert hmcl.nombre_zip("PED-2026-S40-001", CORTE) == (
        "Pedidos_PED-2026-S40-001_2026-10-01.zip")


# --- el libro (EX-01, EX-03, EX-04, EX-05) -----------------------------------


def test_the_sheet_has_the_header_block_and_the_two_columns_ex_01():
    libro = _libro(_datos([("00123-AB", D("12")), ("94109-12000S", D("50"))]))

    hoja = libro["Pedido"]

    assert libro.sheetnames == ["Pedido"]
    assert _filas(hoja) == [
        ["Tienda", "Manizales"],
        ["SIC", "1234"],
        ["Fecha", datetime.datetime(2026, 10, 1)],
        [None, None],
        ["Código", "Cantidad"],
        ["00123-AB", 12],
        ["94109-12000S", 50]]


def test_the_header_date_is_the_corte_as_a_real_date_cell():
    hoja = _libro(_datos([("A", D("1"))]))["Pedido"]

    assert hoja["B3"].is_date is True
    assert hoja["B3"].number_format == "yyyy-mm-dd"


def test_there_are_no_sugerido_price_or_value_columns_ex_04():
    hoja = _libro(_datos([("A", D("1")), ("B", D("2"))]))["Pedido"]

    assert hoja.max_column == 2
    assert [c.value for c in hoja[5]] == ["Código", "Cantidad"]


def test_only_quantities_above_zero_are_written_ex_01():
    hoja = _libro(_datos([
        ("00123-AB", D("12")), ("55512-A", D("0")), ("77777", D("-3")),
        ("94109-12000S", D("50"))]))["Pedido"]

    assert _filas(hoja)[5:] == [["00123-AB", 12], ["94109-12000S", 50]]


def test_the_lines_are_ordered_by_code_whatever_the_input_order():
    hoja = _libro(_datos([
        ("B-2", D("1")), ("A-10", D("2")), ("A-1", D("3"))]))["Pedido"]

    assert [f[0] for f in _filas(hoja)[5:]] == ["A-1", "A-10", "B-2"]


def test_quantities_are_numbers_and_whole_ones_are_integers():
    hoja = _libro(_datos([("A", D("12.00")), ("B", D("7.50"))]))["Pedido"]

    assert hoja["B6"].value == 12 and isinstance(hoja["B6"].value, int)
    assert hoja["B6"].data_type == "n"
    assert hoja["B7"].value == 7.5


def test_leading_zeros_survive_as_text_ex_05():
    hoja = _libro(_datos([("000123", D("5"))]))["Pedido"]

    assert hoja["A6"].value == "000123"
    assert hoja["A6"].data_type == "s"


def test_a_code_that_starts_with_an_equals_sign_stays_text():
    peligro = '=HYPERLINK("http://malo.example","x")'
    destino = io.BytesIO()

    hmcl.escribir_libro(_datos([(peligro, D("5"))]), destino)

    hoja = load_workbook(io.BytesIO(destino.getvalue()))["Pedido"]
    assert hoja["A6"].value == peligro and hoja["A6"].data_type == "s"
    crudo = zipfile.ZipFile(io.BytesIO(destino.getvalue())).read(
        "xl/worksheets/sheet1.xml").decode()
    assert "<f>" not in crudo and "<f " not in crudo


def test_other_formula_openers_and_the_tienda_name_stay_text_too():
    datos = _datos(
        [("+1+1", D("1")), ("-2+3", D("2")), ("@SUM(A1)", D("3"))],
        nombre="=cmd|' /C calc'!A0", sic="=1+1")

    hoja = _libro(datos)["Pedido"]

    assert hoja["B1"].value == "=cmd|' /C calc'!A0"
    assert hoja["B1"].data_type == "s"
    assert hoja["B2"].value == "=1+1" and hoja["B2"].data_type == "s"
    assert [f[0] for f in _filas(hoja)[5:]] == ["+1+1", "-2+3", "@SUM(A1)"]
    assert {hoja.cell(r, 1).data_type for r in range(6, 9)} == {"s"}


def test_control_characters_are_stripped_instead_of_breaking_the_file():
    hoja = _libro(_datos([("AB\x00\x1fC", D("1"))]))["Pedido"]

    assert hoja["A6"].value == "ABC"


def test_a_pedido_with_no_lines_still_writes_a_valid_header_and_titles():
    hoja = _libro(_datos([]))["Pedido"]

    assert _filas(hoja)[4] == ["Código", "Cantidad"]
    assert hoja.max_row == 5


def test_the_book_is_written_into_a_file_like_with_a_zip_inside():
    destino = io.BytesIO()

    hmcl.escribir_libro(_datos([("A", D("1"))]), destino)

    assert zipfile.is_zipfile(io.BytesIO(destino.getvalue())) is True


# --- el zip (EX-02) ----------------------------------------------------------


def _zip(archivos):
    destino = io.BytesIO()
    hmcl.empaquetar(archivos, destino)
    return zipfile.ZipFile(io.BytesIO(destino.getvalue()))


def test_the_zip_holds_every_file_with_its_exact_content():
    paquete = _zip([("uno.xlsx", b"AAA"), ("dos.xlsx", b"BBBB")])

    assert paquete.namelist() == ["uno.xlsx", "dos.xlsx"]
    assert paquete.read("uno.xlsx") == b"AAA"
    assert paquete.read("dos.xlsx") == b"BBBB"
    assert paquete.testzip() is None


def test_zip_member_names_are_utf_8():
    paquete = _zip([("Pedido_Ñandú.xlsx", b"x")])

    assert paquete.namelist() == ["Pedido_Ñandú.xlsx"]
    assert paquete.infolist()[0].flag_bits & 0x800


def test_the_zip_is_compressed():
    paquete = _zip([("a.xlsx", b"0" * 5000)])

    assert paquete.infolist()[0].compress_type == zipfile.ZIP_DEFLATED
    assert paquete.infolist()[0].compress_size < 5000


def test_two_files_with_the_same_name_do_not_overwrite_each_other():
    paquete = _zip([("a.xlsx", b"1"), ("a.xlsx", b"2"), ("a.xlsx", b"3")])

    assert paquete.namelist() == ["a.xlsx", "a_2.xlsx", "a_3.xlsx"]
    assert [paquete.read(n) for n in paquete.namelist()] == [
        b"1", b"2", b"3"]


def test_an_empty_zip_is_still_a_valid_zip():
    assert _zip([]).namelist() == []


# --- cabeceras de la respuesta -----------------------------------------------


def test_content_disposition_carries_an_ascii_name_and_the_rfc_5987_one():
    valor = hmcl.content_disposition("Pedido_SIC1_Medellin_2026-10-01.xlsx")

    assert valor == (
        'attachment; filename="Pedido_SIC1_Medellin_2026-10-01.xlsx"; '
        "filename*=UTF-8''Pedido_SIC1_Medellin_2026-10-01.xlsx")


def test_a_non_ascii_name_gets_a_folded_fallback_and_a_percent_encoded_one():
    valor = hmcl.content_disposition("Pedido_Ñandú.xlsx")

    assert 'filename="Pedido_Nandu.xlsx"' in valor
    assert "filename*=UTF-8''Pedido_%C3%91and%C3%BA.xlsx" in valor
    assert valor.isascii()


def test_quotes_and_line_breaks_cannot_break_out_of_the_header():
    valor = hmcl.content_disposition('a"b\r\nSet-Cookie: x=1.xlsx')

    assert "\r" not in valor and "\n" not in valor
    assert valor.count('"') == 2
    assert "%22" in valor and "%0D%0A" in valor


def test_the_skipped_tiendas_header_is_ascii_percent_encoded_json():
    omitidas = [
        {"sucursal_id": "1", "nombre": "Medellín", "motivo": "Pedido en"
         " BORRADOR", "codigo": "BORRADOR"}]

    valor = hmcl.cabecera_omitidas(omitidas)

    assert valor.isascii() and "Medell" in valor and "í" not in valor
    assert json.loads(unquote(valor)) == omitidas


def test_no_skipped_tiendas_is_an_empty_json_list():
    assert json.loads(unquote(hmcl.cabecera_omitidas([]))) == []
