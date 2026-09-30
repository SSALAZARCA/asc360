"""
Motored Pedidos F3 "Motor" (S8b-1, ADR-10, decisión #11 y #16) — libros de
Excel para el dueño: primitivas de escritura, sección "Antigüedad de datos"
y libro de corrida simple (una hoja por sucursal en el orden de columnas A..AD
de `Pedido`, resumen por clase, excluidas y advertencias).

Todo con datos sintéticos y la fila patrón §10.1. Los libros se escriben en
`tmp_path` (fuera de cualquier árbol git) y se leen de vuelta con openpyxl.
"""
from datetime import date, datetime
from decimal import Decimal
from fractions import Fraction
from pathlib import Path

import openpyxl
import pytest

from app.motored.herramientas.regresion.antiguedad import (
    ENCABEZADO_ANTIGUEDAD,
    FilaAntiguedad,
    como_filas,
    desde_seleccion,
    sin_datos,
)
from app.motored.herramientas.regresion.informe_excel import (
    ENCABEZADO_PEDIDO,
    HOJA_ANTIGUEDAD,
    DatosCorrida,
    SucursalCorrida,
    escribir_libro_corrida,
)
from app.motored.herramientas.regresion.libro_excel import (
    LibroSalida,
    titulo_hoja,
)
from app.motored.herramientas.regresion.rutas import RutaInsegura
from app.motored.services.motor.tipos import Advertencia
from tests.motored.fixtures.regresion.resultados import (
    CORTE,
    resultado_con_sustitucion,
    resultado_omitida,
    resultado_patron,
)

SELECCION = {
    "antiguedad": {
        "inventario": {
            "carga_id": "c1", "fecha_usada": "2026-09-19",
            "antiguedad_dias": 2, "limite_dias": 7, "fuente_limite": "DEFAULT",
        },
        "backorder": {
            "carga_id": "c2", "fecha_usada": "2026-09-14",
            "antiguedad_dias": 7, "limite_dias": 7, "fuente_limite": "GLOBAL",
        },
        "facturas": {
            "carga_id": "c3", "fecha_usada": "2026-09-10",
            "antiguedad_dias": 11, "limite_dias": 15,
            "fuente_limite": "GLOBAL",
        },
        "ingresos": {
            "carga_id": "c4", "fecha_usada": "2026-09-20",
            "antiguedad_dias": 1, "limite_dias": 7, "fuente_limite": "DEFAULT",
        },
    }
}


def _rellenar(filas: list) -> list:
    """Filas del mismo ancho: openpyxl recorta los `None` del final."""
    ancho = max((len(fila) for fila in filas), default=0)
    return [fila + (None,) * (ancho - len(fila)) for fila in filas]


def _leer(ruta: Path) -> dict:
    """`{hoja: [filas]}` de un libro, con los valores tal como quedaron."""
    libro = openpyxl.load_workbook(ruta, read_only=True)
    hojas = {
        nombre: _rellenar(list(libro[nombre].iter_rows(values_only=True)))
        for nombre in libro.sheetnames
    }
    libro.close()
    return hojas


def _datos(*sucursales, antiguedad=None, advertencias=()):
    return DatosCorrida(
        titulo="PED-2026-S39-001",
        fecha_corte=CORTE,
        sucursales=tuple(SucursalCorrida(a, r) for a, r in sucursales),
        antiguedad=antiguedad or sin_datos("prueba"),
        advertencias=tuple(advertencias),
    )


class TestTituloHoja:
    def test_reemplaza_los_caracteres_prohibidos_por_excel(self):
        assert titulo_hoja("A/B:C*D?E[F]G\\H") == "A-B-C-D-E-F-G-H"

    def test_corta_a_31_caracteres_y_quita_espacios_de_los_bordes(self):
        titulo = titulo_hoja("  " + "X" * 40 + "  ")
        assert titulo == "X" * 31

    def test_un_titulo_vacio_se_nombra_hoja(self):
        assert titulo_hoja("   ") == "Hoja"


class TestLibroSalida:
    def test_usa_el_modo_write_only_de_openpyxl(self, tmp_path):
        assert LibroSalida(tmp_path / "a.xlsx").libro.write_only is True

    def test_escribe_encabezado_filas_y_convierte_fracciones(
        self, tmp_path
    ):
        libro = LibroSalida(tmp_path / "salida" / "libro.xlsx")
        libro.hoja(
            "Datos", ("Codigo", "N", "Vacio"),
            [("A", Fraction(598, 7), None), ("B", Decimal("2.50"), "x")],
        )
        filas = _leer(libro.guardar())["Datos"]
        assert filas[0] == ("Codigo", "N", "Vacio")
        assert filas[1] == ("A", 85.428571, None)
        assert filas[2] == ("B", 2.5, "x")

    def test_las_filas_previas_van_antes_del_encabezado(self, tmp_path):
        libro = LibroSalida(tmp_path / "b.xlsx")
        libro.hoja(
            "Bloque", ("Col",), [("fila",)],
            previas=[("Clave", "Valor"), ()],
        )
        filas = _leer(libro.guardar())["Bloque"]
        assert filas[0] == ("Clave", "Valor")
        assert filas[2] == ("Col", None)
        assert filas[3] == ("fila", None)

    def test_dos_hojas_con_el_mismo_titulo_reciben_un_sufijo(self, tmp_path):
        libro = LibroSalida(tmp_path / "c.xlsx")
        primero = libro.hoja("Suc", ("a",), [])
        segundo = libro.hoja("Suc", ("a",), [])
        assert (primero, segundo) == ("Suc", "Suc (2)")
        assert list(_leer(libro.guardar())) == ["Suc", "Suc (2)"]

    def test_rechaza_una_ruta_dentro_del_repositorio_sin_escribir(self):
        dentro = Path(__file__).parent / "no_debe_existir.xlsx"
        with pytest.raises(RutaInsegura):
            LibroSalida(dentro)
        assert not dentro.exists()


class TestAntiguedad:
    def test_lee_las_cuatro_entradas_en_orden(self):
        filas = desde_seleccion(SELECCION)
        assert [f.tipo for f in filas] == [
            "Inventario", "Backorder", "Facturas de pedidos",
            "Ingresos de facturas",
        ]
        assert filas[0] == FilaAntiguedad(
            "Inventario", date(2026, 9, 19), 2, 7, "DEFAULT"
        )
        assert filas[2] == FilaAntiguedad(
            "Facturas de pedidos", date(2026, 9, 10), 11, 15, "GLOBAL"
        )

    def test_un_tipo_ausente_se_informa_sin_dato(self):
        parcial = {"antiguedad": {"inventario": SELECCION[
            "antiguedad"]["inventario"]}}
        filas = desde_seleccion(parcial)
        assert filas[0].antiguedad_dias == 2
        assert filas[1] == FilaAntiguedad(
            "Backorder", None, None, None, "Sin dato"
        )

    def test_sin_datos_explica_el_motivo_en_la_fuente(self):
        filas = sin_datos("libro de Excel")
        assert len(filas) == 4
        assert {f.fuente for f in filas} == {"libro de Excel"}
        assert all(f.fecha_usada is None for f in filas)

    def test_como_filas_sigue_el_encabezado(self):
        filas = como_filas(desde_seleccion(SELECCION))
        assert len(ENCABEZADO_ANTIGUEDAD) == len(filas[0]) == 5
        assert filas[1] == ("Backorder", date(2026, 9, 14), 7, 7, "GLOBAL")


class TestLibroCorrida:
    def test_las_hojas_van_en_el_orden_del_informe(self, tmp_path):
        ruta = escribir_libro_corrida(
            _datos(resultado_patron()), tmp_path / "corrida.xlsx"
        )
        assert list(_leer(ruta)) == [
            "Resumen clases", HOJA_ANTIGUEDAD, "MANIZALES AV SANTANDER",
            "Excluidas", "Advertencias",
        ]

    def test_la_hoja_de_sucursal_sigue_las_columnas_a_a_ad(self, tmp_path):
        ruta = escribir_libro_corrida(
            _datos(resultado_patron()), tmp_path / "corrida.xlsx"
        )
        filas = _leer(ruta)["MANIZALES AV SANTANDER"]
        assert len(ENCABEZADO_PEDIDO) == 30
        assert filas[0] == ENCABEZADO_PEDIDO
        assert ENCABEZADO_PEDIDO[0] == "Referencia"
        assert ENCABEZADO_PEDIDO[29] == "Cobertura Final"
        assert len(filas) == 2

    def test_la_fila_patron_sale_con_los_ocho_valores_de_la_especificacion(
        self, tmp_path
    ):
        ruta = escribir_libro_corrida(
            _datos(resultado_patron()), tmp_path / "corrida.xlsx"
        )
        fila = _leer(ruta)["MANIZALES AV SANTANDER"][1]
        assert fila[0] == "94109-12000S"
        assert fila[4:10] == (102, 112, 108, 105, 74, 59)
        assert fila[13] == 85.428571
        assert fila[18] == "CF"
        assert fila[19] == 460.75
        assert fila[21:26] == (27, 70, 0, 97, -3)
        assert fila[26] == 149.5
        assert fila[27] == 50
        assert fila[28] == 23037.5
        assert fila[29] == 1.720736

    def test_el_resumen_por_clase_incluye_totales_de_sucursal_y_de_red(
        self, tmp_path
    ):
        ruta = escribir_libro_corrida(
            _datos(resultado_patron(), resultado_con_sustitucion()),
            tmp_path / "corrida.xlsx",
        )
        filas = _leer(ruta)["Resumen clases"]
        assert filas[0][:2] == ("Sucursal", "Clase")
        manizales = [f for f in filas if f[0] == "MANIZALES AV SANTANDER"]
        clases = {f[1]: f for f in manizales}
        assert clases["CF"][2:5] == (50, 1, 23037.5)
        assert clases["TOTAL"][2:5] == (50, 1, 23037.5)
        assert clases["TOTAL"][5] == 1
        totales = {
            f[0]: f for f in filas if f[1] == "TOTAL"
        }
        assert totales["SUCURSAL CONSOLIDADA"][2] > 0
        assert totales["TODA LA RED"][2] == (
            totales["MANIZALES AV SANTANDER"][2]
            + totales["SUCURSAL CONSOLIDADA"][2]
        )
        assert totales["TODA LA RED"][5] == 1

    def test_con_una_sola_sucursal_no_hay_filas_de_red(self, tmp_path):
        ruta = escribir_libro_corrida(
            _datos(resultado_patron()), tmp_path / "corrida.xlsx"
        )
        filas = _leer(ruta)["Resumen clases"]
        assert {f[0] for f in filas[1:]} == {"MANIZALES AV SANTANDER"}

    def test_la_seccion_de_antiguedad_lleva_fecha_edad_limite_y_fuente(
        self, tmp_path
    ):
        datos = _datos(
            resultado_patron(), antiguedad=desde_seleccion(SELECCION)
        )
        ruta = escribir_libro_corrida(datos, tmp_path / "corrida.xlsx")
        filas = _leer(ruta)[HOJA_ANTIGUEDAD]
        assert filas[0][0] == "Corrida" and filas[0][1] == "PED-2026-S39-001"
        assert filas[1][:2] == ("Fecha de corte", datetime(2026, 9, 21))
        encabezado = next(f for f in filas if f[0] == "Tipo de dato")
        assert encabezado == ENCABEZADO_ANTIGUEDAD
        datos_hoja = {f[0]: f for f in filas if f[0] in (
            "Inventario", "Backorder", "Facturas de pedidos",
            "Ingresos de facturas",
        )}
        assert datos_hoja["Facturas de pedidos"][1:] == (
            datetime(2026, 9, 10), 11, 15, "GLOBAL"
        )
        assert len(datos_hoja) == 4

    def test_las_excluidas_listan_motivo_y_sustituta(self, tmp_path):
        ruta = escribir_libro_corrida(
            _datos(resultado_con_sustitucion()), tmp_path / "corrida.xlsx"
        )
        filas = _leer(ruta)["Excluidas"]
        assert filas[0][:3] == ("Sucursal", "Referencia", "Motivo")
        assert filas[1][:4] == (
            "SUCURSAL CONSOLIDADA", "REF-A", "SUSTITUIDA", "REF-B"
        )

    def test_una_sucursal_omitida_deja_su_hoja_vacia_y_su_advertencia(
        self, tmp_path
    ):
        datos = _datos(
            resultado_patron(), resultado_omitida(),
            advertencias=[Advertencia("A-CORRIDA-101", "Sin demanda perdida")],
        )
        ruta = escribir_libro_corrida(datos, tmp_path / "corrida.xlsx")
        hojas = _leer(ruta)
        assert len(hojas["SUCURSAL NUEVA"]) == 1
        advertencias = hojas["Advertencias"]
        assert advertencias[0] == ("Sucursal", "Código", "Mensaje")
        assert advertencias[1][:2] == ("TODA LA CORRIDA", "A-CORRIDA-101")
        assert any(
            f[0] == "SUCURSAL NUEVA" and f[1] == "A-CORRIDA-102"
            for f in advertencias
        )

    def test_rechaza_escribir_dentro_del_repositorio(self):
        dentro = Path(__file__).parent / "corrida_prohibida.xlsx"
        with pytest.raises(RutaInsegura):
            escribir_libro_corrida(_datos(resultado_patron()), dentro)
        assert not dentro.exists()

    def test_dos_sucursales_con_el_mismo_nombre_no_chocan(self, tmp_path):
        sucursal, resultado = resultado_patron()
        repetidas = [(sucursal, resultado)] * 2
        ruta = escribir_libro_corrida(_datos(*repetidas), tmp_path / "g.xlsx")
        nombres = list(_leer(ruta))
        assert nombres.count("MANIZALES AV SANTANDER") == 1
        assert "MANIZALES AV SANTANDER (2)" in nombres
