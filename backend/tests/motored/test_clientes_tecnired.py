"""
CLIENTES TECNIRED (feature motored-tablero-asesores, T2).

Lista de NIT de clientes Tecnired que alimenta el tablero de asesores. Se
sube desde Maestros con el patron todo-o-nada de los 4 maestros, pero con una
semantica distinta: cada carga REEMPLAZA la lista completa (borra todo e
inserta lo nuevo en UNA transaccion, solo si el archivo entero valida).

Esta lista NO es un tipo de movimiento ni alimenta al motor de corridas.
"""
import io
import uuid

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.config import settings
from app.main import app
from app.motored.models.cliente_tecnired import ClienteTecnired
from app.motored.services import carga, carga_excel, validators
from app.motored.services import kpi_resumen
from app.motored.services.auth import MotoredUser
from tests.motored.conftest import FakeAsyncSession, override_motored_db, override_motored_user


@pytest.fixture(autouse=True)
def _sin_marca_de_sucio(monkeypatch):
    """The KPI summaries' dirty flag (R7a) needs a real session (`pg_real/test_kpi_resumen_sucio_pg.py`)."""
    async def marcar(db):
        return True

    monkeypatch.setattr(kpi_resumen, "marcar_sucio_si_construido", marcar)


# --- Normalizacion del NIT --------------------------------------------------


@pytest.mark.parametrize(
    "crudo, esperado",
    [
        ("900123456", "900123456"),
        ("  900123456  ", "900123456"),
        ("900123456.", "900123456"),
        (" 900123456. ", "900123456"),
        ("900 123 456", "900123456"),
        ("900123456.0", "900123456"),
        ("900.123.456-7", "900.123.456-7"),
    ],
)
def test_normalizar_nit_copia_la_regla_del_excel(crudo, esperado):
    from app.motored.schemas.cliente_tecnired import normalizar_nit

    assert normalizar_nit(crudo) == esperado


# --- Validacion --------------------------------------------------------------


def test_validate_rows_exige_nit():
    _validas, errores = validators.validate_rows("cliente_tecnired", [{"nit": "  "}, {"razon_social": "X"}])

    assert [e["fila"] for e in errores] == [1, 2]


def test_nit_que_queda_vacio_tras_normalizar_es_error_de_fila():
    _validas, errores = validators.validate_rows("cliente_tecnired", [{"nit": "."}])

    assert len(errores) == 1 and errores[0]["fila"] == 1


def test_archivo_sin_filas_se_rechaza_para_no_borrar_la_lista():
    _validas, errores = validators.validate_rows("cliente_tecnired", [])

    assert errores and errores[0]["fila"] == 0
    assert "no se reemplaza" in errores[0]["motivo"].lower()


def test_validate_rows_normaliza_el_nit_de_las_filas_validas():
    validas, errores = validators.validate_rows(
        "cliente_tecnired", [{"nit": " 900123456. ", "razon_social": " ACME SAS "}]
    )

    assert errores == []
    assert validas[0]["nit"] == "900123456"


# --- procesar_carga: reemplazo atomico ---------------------------------------


async def test_archivo_invalido_no_toca_la_base():
    db = FakeAsyncSession()  # cualquier execute() revienta: la cola esta vacia

    resultado = await carga.procesar_carga(
        db, "cliente_tecnired", [{"nit": "900"}, {"nit": ""}]
    )

    assert resultado.ok is False
    assert db.executed_statements == []
    assert db.committed is False


async def test_archivo_valido_borra_todo_inserta_y_confirma_una_vez():
    anteriores = [[ClienteTecnired(id=uuid.uuid4(), nit=f"1{i}") for i in range(3)]]
    db = FakeAsyncSession(execute_queue=[anteriores[0]])  # DELETE (rowcount=3)

    resultado = await carga.procesar_carga(
        db, "cliente_tecnired",
        [{"nit": "900123456", "razon_social": "ACME"}, {"nit": "800111222."}],
        usuario_id=uuid.uuid4(),
    )

    assert resultado.ok is True
    assert resultado.insertados == 2
    assert resultado.eliminados == 3
    assert db.committed is True
    sentencias = [str(s).split()[0].upper() for s in db.executed_statements]
    assert sentencias == ["DELETE"]
    assert sorted(c.nit for c in db.added_of_type(ClienteTecnired)) == ["800111222", "900123456"]


async def test_nit_repetido_en_el_archivo_se_deduplica_con_advertencia():
    db = FakeAsyncSession(execute_queue=[[]])

    resultado = await carga.procesar_carga(
        db, "cliente_tecnired",
        [{"nit": "900", "razon_social": "PRIMERA"}, {"nit": "900.", "razon_social": "SEGUNDA"}],
    )

    assert resultado.ok is True
    assert resultado.insertados == 1
    assert len(resultado.advertencias) == 1
    assert resultado.advertencias[0]["fila"] == 2
    [conservado] = db.added_of_type(ClienteTecnired)
    assert conservado.razon_social == "PRIMERA"


# --- Excel: columnas y alias --------------------------------------------------


def test_plantilla_tiene_nit_y_razon_social():
    assert carga_excel.column_labels("cliente_tecnired") == ["NIT", "Razón social"]


def _xlsx(encabezados, filas) -> bytes:
    libro = openpyxl.Workbook()
    libro.active.append(encabezados)
    for fila in filas:
        libro.active.append(fila)
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize("alias", ["NIT", "NIT / CEDULA", "Cliente factura", "Cedula", "Cédula"])
def test_alias_del_nit_en_el_excel(alias):
    contenido = _xlsx([alias, "Razón social"], [[900123456, "ACME"]])

    filas = carga_excel.parse_excel_rows("cliente_tecnired", "c.xlsx", contenido)

    assert filas == [{"nit": "900123456", "razon_social": "ACME"}]


def test_razon_social_es_opcional_en_el_excel():
    contenido = _xlsx(["NIT"], [["900123456"]])

    filas = carga_excel.parse_excel_rows("cliente_tecnired", "c.xlsx", contenido)

    assert filas == [{"nit": "900123456"}]


def test_el_excel_sin_columna_nit_se_rechaza():
    contenido = _xlsx(["Razón social"], [["ACME"]])

    with pytest.raises(carga_excel.ColumnaObligatoriaFaltanteError):
        carga_excel.parse_excel_rows("cliente_tecnired", "c.xlsx", contenido)


# --- No es movimiento ni lo lee el motor --------------------------------------


def test_clientes_tecnired_no_es_un_tipo_de_movimiento():
    from app.motored.services.ingesta import orquestador

    assert not any("TECNIRED" in t.upper() for t in orquestador.TIPOS_MOVIMIENTO)


def test_vigencia_no_lista_clientes_tecnired():
    from app.motored.services.corridas import vigencia

    assert "TECNIRED" not in repr(vigencia._TIPOS).upper()
    assert "TECNIRED" not in repr(vigencia._TIPOS_PREFLIGHT).upper()


# --- Migracion ------------------------------------------------------------------


def _migracion():
    import importlib.util
    from pathlib import Path

    archivo = (
        Path(__file__).resolve().parents[2] / "alembic_motored" / "versions"
        / "d3b7f19a4c26_cliente_tecnired.py")
    spec = importlib.util.spec_from_file_location("tablero_t2_cliente_tecnired", archivo)
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    return modulo


def test_la_migracion_encadena_sobre_inventario_detalle_y_es_aditiva():
    from unittest.mock import patch

    m = _migracion()
    assert m.down_revision == "c4a9e7d1b852"
    with patch.object(m, "op") as op_mock:
        m.upgrade()
    assert [c.args[0] for c in op_mock.create_table.call_args_list] == ["cliente_tecnired"]
    assert not op_mock.drop_table.called and not op_mock.alter_column.called


# --- HTTP ---------------------------------------------------------------------

CARGA_URL = "/api/motored/maestros/cliente_tecnired/carga"
LISTA_URL = "/api/motored/clientes-tecnired"


@pytest.fixture
def _motored_ready(monkeypatch):
    monkeypatch.setattr(settings, "MOTORED_ENABLED", True)
    monkeypatch.setattr(settings, "MOTORED_SECRET_KEY", "tecnired-test-motored-secret")
    monkeypatch.setattr(settings, "SECRET_KEY", "tecnired-test-asc360-secret")
    yield
    app.dependency_overrides.clear()


def _como(rol):
    override_motored_user(MotoredUser(user_id=str(uuid.uuid4()), role=rol))


def test_carga_http_reemplaza_y_responde_conteos(_motored_ready):
    _como("COMPRAS")
    db = FakeAsyncSession(execute_queue=[[], [1, 2]])  # probe, DELETE(2)
    override_motored_db(db)

    with TestClient(app) as client:
        r = client.post(CARGA_URL, json={"filas": [{"nit": "900123456"}]})

    assert r.status_code == 200
    cuerpo = r.json()
    assert cuerpo["ok"] is True and cuerpo["insertados"] == 1 and cuerpo["eliminados"] == 2
    assert db.committed is True


def test_carga_http_exige_admin_o_compras(_motored_ready):
    _como("CONSULTA")
    override_motored_db(FakeAsyncSession(execute_queue=[[]]))

    with TestClient(app) as client:
        r = client.post(CARGA_URL, json={"filas": [{"nit": "900123456"}]})

    assert r.status_code == 403


def test_plantilla_http(_motored_ready):
    _como("ADMIN")

    with TestClient(app) as client:
        r = client.get("/api/motored/maestros/cliente_tecnired/plantilla.xlsx")

    assert r.status_code == 200
    libro = openpyxl.load_workbook(io.BytesIO(r.content))
    assert [c.value for c in next(libro.active.iter_rows(max_row=1))] == ["NIT", "Razón social"]


def test_lista_http_pagina_y_exige_rol(_motored_ready):
    _como("SUCURSAL")
    override_motored_db(FakeAsyncSession(execute_queue=[[]]))
    with TestClient(app) as client:
        assert client.get(LISTA_URL).status_code == 403

    _como("ADMIN")
    filas = [ClienteTecnired(id=uuid.uuid4(), nit="900123456", razon_social="ACME")]
    override_motored_db(FakeAsyncSession(execute_queue=[[], [1], filas]))  # probe, COUNT, SELECT
    with TestClient(app) as client:
        r = client.get(LISTA_URL, params={"q": "900"})

    assert r.status_code == 200
    cuerpo = r.json()
    assert cuerpo["total"] == 1
    assert cuerpo["items"][0]["nit"] == "900123456"


# --- Advertencias de la revision de T2 ------------------------------------------------------


@pytest.mark.parametrize("crudo", ["123.0.", "123.", " 1 23.0 ", "900123456.0", "900.123.456-7", "1.0.0", "..", "12.00."])
def test_normalizar_nit_es_idempotente(crudo):
    from app.motored.schemas.cliente_tecnired import normalizar_nit

    una_vez = normalizar_nit(crudo)

    assert normalizar_nit(una_vez) == una_vez


def test_carga_concurrente_que_choca_en_el_indice_unico_es_409_y_hace_rollback(_motored_ready):
    _como("COMPRAS")
    # Otro reemplazo confirmo primero: el COMMIT de este choca con uq_cliente_tecnired_nit.
    db = FakeAsyncSession(execute_queue=[[], [1]], raise_integrity_error=True)
    override_motored_db(db)

    with TestClient(app) as client:
        r = client.post(CARGA_URL, json={"filas": [{"nit": "900123456"}]})

    assert r.status_code == 409
    assert "otra carga" in r.json()["detail"].lower()
    assert db.rolled_back is True and db.committed is False
